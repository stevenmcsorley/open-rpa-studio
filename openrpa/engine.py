from __future__ import annotations

import csv
import json
import datetime
import email.utils
import math
import shutil
import time
import tempfile
import os
import io
import copy as copy_module
import zipfile
import sqlite3
import subprocess
import base64
import re
import unicodedata
from difflib import SequenceMatcher
from dataclasses import dataclass, field
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urljoin, urlsplit, urlunsplit, parse_qsl, urlencode, unquote

from .expressions import evaluate, interpolate
from .model import WorkflowError, validate_workflow, validate_workflow_bindings
from .network import request as restricted_http_request, EgressDenied, assert_public_url
from .credentials import get_safe_keyring


class RunCancelled(Exception):
    pass


def _json_pointer_get(value: Any, pointer: str) -> Any:
    if pointer == "":
        return value
    current = value
    for raw in pointer[1:].split("/"):
        token = raw.replace("~1", "/").replace("~0", "~")
        if isinstance(current, dict) and token in current:
            current = current[token]
        elif isinstance(current, list) and token.isdigit() and str(int(token)) == token and int(token) < len(current):
            current = current[int(token)]
        else:
            raise KeyError("JSON Pointer did not resolve")
    return current


def _retry_after_seconds(headers: Any, *, now: datetime.datetime | None = None) -> float | None:
    """Parse a bounded Retry-After value without exposing the raw header."""
    if not isinstance(headers, dict):
        return None
    raw = next((value for key, value in headers.items()
                if isinstance(key, str) and key.lower() == "retry-after"), None)
    if not isinstance(raw, str) or len(raw) > 128:
        return None
    value = raw.strip()
    if value.isascii() and value.isdigit():
        try:
            return min(float(int(value)), 30.0)
        except (ValueError, OverflowError):
            return None
    try:
        retry_at = email.utils.parsedate_to_datetime(value)
        if retry_at is None:
            return None
        if retry_at.tzinfo is None:
            retry_at = retry_at.replace(tzinfo=datetime.timezone.utc)
        current = now or datetime.datetime.now(datetime.timezone.utc)
        if current.tzinfo is None:
            current = current.replace(tzinfo=datetime.timezone.utc)
        seconds = (retry_at - current).total_seconds()
        if not math.isfinite(seconds):
            return None
        return min(max(0.0, seconds), 30.0)
    except (TypeError, ValueError, OverflowError):
        return None


def _inspect_json_shape(value: Any, *, max_depth: int = 6, max_properties: int = 50,
                        max_array_samples: int = 3, max_nodes: int = 500) -> dict[str, Any]:
    """Return a bounded, value-free JSON shape and object-array pointer candidates."""
    array_pointers: list[str] = []
    nodes = 0
    truncated = False

    def escape(token: str) -> str:
        return token.replace("~", "~0").replace("/", "~1")

    def visit(item: Any, pointer: str, depth: int) -> dict[str, Any]:
        nonlocal nodes, truncated
        nodes += 1
        if nodes > max_nodes:
            truncated = True
            return {"type": "truncated"}
        if item is None:
            return {"type": "null"}
        if isinstance(item, bool):
            return {"type": "boolean"}
        if isinstance(item, str):
            return {"type": "string"}
        if isinstance(item, (int, float)):
            return {"type": "number"}
        if depth >= max_depth:
            if isinstance(item, (dict, list)) and item:
                truncated = True
                return {"type": "object" if isinstance(item, dict) else "array", "truncated": True}
        if isinstance(item, dict):
            keys = sorted(item, key=lambda key: str(key))
            if len(keys) > max_properties:
                truncated = True
                keys = keys[:max_properties]
            return {"type": "object", "properties": {
                str(key): visit(item[key], pointer + "/" + escape(str(key)), depth + 1)
                for key in keys
            }}
        if isinstance(item, list):
            object_items = [entry for entry in item[:max_array_samples] if isinstance(entry, dict)]
            if object_items and pointer not in array_pointers:
                array_pointers.append(pointer)
            if len(item) > max_array_samples:
                truncated = True
            samples = item[:max_array_samples]
            if not samples:
                return {"type": "array", "items": {"type": "unknown"}}
            shapes = [visit(entry, pointer + "/" + str(index), depth + 1) for index, entry in enumerate(samples)]
            first = merge_shapes(shapes)
            if first.get("type") == "object":
                first["presentInAllSamples"] = sorted(set.intersection(*(set(shape.get("properties", {})) for shape in shapes))) if shapes else []
            return {"type": "array", "items": first, "sampledItems": len(samples)}
        return {"type": "unknown"}

    def merge_shapes(shapes: list[dict[str, Any]]) -> dict[str, Any]:
        nonlocal truncated
        if not shapes:
            return {"type": "unknown"}
        kinds={shape.get("type","unknown") for shape in shapes}
        if kinds=={"object"}:
            keys=sorted({key for shape in shapes for key in shape.get("properties",{})})
            if len(keys)>max_properties:
                truncated=True
                keys=keys[:max_properties]
            return {"type":"object","properties":{
                key:merge_shapes([shape["properties"][key] for shape in shapes if key in shape.get("properties",{})])
                for key in keys
            }}
        if kinds=={"array"}:
            return {"type":"array","items":merge_shapes([shape.get("items",{"type":"unknown"}) for shape in shapes]),"sampledItems":sum(shape.get("sampledItems",0) for shape in shapes)}
        if len(kinds)==1:
            return shapes[0]
        unique={json.dumps(shape,ensure_ascii=False,sort_keys=True,separators=(",",":")):shape for shape in shapes}
        return {"type":"mixed","alternatives":[unique[key] for key in sorted(unique)]}

    schema = visit(value, "", 0)
    return {"schema": schema, "objectArrayPointers": array_pointers, "truncated": truncated}


def _normalize_accessible_name(value: str) -> str:
    normalized=unicodedata.normalize("NFKC",value).casefold()
    return " ".join(re.findall(r"[\w]+",normalized,flags=re.UNICODE))


def _accessible_role_name(snapshot: str, role: str) -> str:
    if not isinstance(snapshot,str):
        return ""
    match=re.match(r"^\s*-\s+"+re.escape(role)+r'\s+"((?:[^"\\]|\\.)*)"(?:\s*:|\s*$)',snapshot)
    if not match:
        return ""
    try:
        value=json.loads('"'+match.group(1)+'"')
    except (TypeError,ValueError,json.JSONDecodeError):
        return ""
    return value if isinstance(value,str) else ""


@dataclass
class RunOptions:
    allow_network: bool = False
    allow_desktop: bool = False
    allow_destructive: bool = False
    extra_file_roots: tuple[Path, ...] = ()
    cancel: Callable[[], bool] = lambda: False
    event: Callable[[dict[str, Any]], None] = lambda _e: None
    max_log_events: int = 1000
    max_log_event_chars: int = 16384
    audit: Callable[[str, dict[str, Any]], None] = lambda _action, _details: None
    max_duration_seconds: int = 3600
    max_activity_executions: int = 100_000
    debugger: Any | None = None
    api_connections_path: Path | None = None


@dataclass
class RunResult:
    status: str
    variables: dict[str, Any]
    logs: list[dict[str, Any]]
    activities: dict[str, dict[str, Any]]
    error: dict[str, Any] | None = None


class WorkflowEngine:
    def __init__(self, project: dict[str, Any], root: Path, options: RunOptions | None = None):
        self.project = project
        self.root = root.resolve()
        self.options = options or RunOptions()
        if self.options.api_connections_path is None:
            from .api_connections import default_store_path
            self.options.api_connections_path=default_store_path()
        limits=project.get("limits",{}) if isinstance(project.get("limits",{}),dict) else {}
        self._max_run_seconds=min(max(int(limits.get("maxRunSeconds",self.options.max_duration_seconds)),1),86_400)
        self._max_activity_executions=min(max(int(limits.get("maxActivityExecutions",self.options.max_activity_executions)),1),1_000_000)
        self._max_loop_items=min(max(int(limits.get("maxLoopItems",100_000)),1),1_000_000)
        self._max_file_bytes=min(max(int(limits.get("maxFileBytes",20_000_000)),1),100_000_000)
        self._activity_executions=0
        self._loop_items_executed=0
        self._run_started=time.monotonic()
        self._browser_runtime: Any = None
        self._browser: Any = None
        self._browser_context: Any = None
        self._browser_persistent_profile = False
        self._browser_page: Any = None
        self._browser_tabs: dict[str, Any] = {}
        self._browser_unclaimed_popups: list[Any] = []
        self._browser_popup_counter = 0
        self._browser_responses: dict[Any, list[dict[str, Any]]] = {}
        self._secret_values: set[str] = set()
        self._api_last_request_at: dict[str, float] = {}
        self._api_tokens: dict[str, tuple[str,float]] = {}
        self.workflows: dict[str, dict[str, Any]] = {}
        for key, rel in project["workflows"].items():
            path = (self.root / rel).resolve()
            self._assert_path(path, allow_file=True)
            try:
                self.workflows[key] = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise WorkflowError(f"Cannot load workflow {key}: {exc}", code="workflow-read-error") from exc
        for key, workflow in self.workflows.items():
            problems = validate_workflow(workflow, set(self.workflows))
            if problems:
                first = problems[0]
                raise WorkflowError(f"{key}: {first['message']}", first.get("activityId"), first["code"])
        self._validate_call_graph()

    def _validate_call_graph(self) -> None:
        calls: dict[str,set[str]]={name:set() for name in self.workflows}
        def visit(name: str, activities: list[dict[str,Any]]) -> None:
            for activity in activities:
                if activity.get("type")=="invoke_workflow" and activity.get("workflow") in calls:
                    calls[name].add(activity["workflow"])
                for branch in ("then","else","body","catch","finally"):
                    if isinstance(activity.get(branch),list): visit(name,activity[branch])
                if activity.get("type")=="switch" and isinstance(activity.get("cases"),list):
                    for case in activity["cases"]:
                        if isinstance(case,dict) and isinstance(case.get("activities"),list):
                            visit(name,case["activities"])
        for name,workflow in self.workflows.items(): visit(name,workflow["activities"])
        for name,workflow in self.workflows.items():
            if workflow.get("executionModel")=="flowchart":
                for node in workflow["flowchart"]["nodes"]:
                    if node["kind"]=="action": visit(name,[node["activity"]])
            if workflow.get("executionModel")=="state_machine":
                for state in workflow["stateMachine"]["states"]:
                    visit(name,state.get("activities",[]))
        active:set[str]=set();complete:set[str]=set()
        def dfs(name:str)->None:
            if name in active: raise WorkflowError(f"Workflow invocation cycle includes '{name}'",code="workflow-cycle")
            if name in complete:return
            active.add(name)
            for child in calls[name]:dfs(child)
            active.remove(name);complete.add(name)
        for name in calls:dfs(name)

    def _assert_path(self, target: Path, allow_file: bool = False) -> None:
        roots = (self.root, *(p.expanduser().resolve() for p in self.options.extra_file_roots))
        for root in roots:
            try:
                target.relative_to(root)
                return
            except ValueError:
                continue
        raise WorkflowError(f"File path escapes approved roots: {target}", code="file-permission-denied")

    def _audit_path(self, target: Path) -> str:
        roots=(self.root,*(item.expanduser().resolve() for item in self.options.extra_file_roots))
        for index,root in enumerate(roots):
            try:relative=target.relative_to(root)
            except ValueError:continue
            return str(relative) if index==0 else str(Path(f"approved-root-{index}")/relative)
        return "[outside-approved-roots]"

    def _path(self, raw: Any, variables: dict[str, Any], for_write: bool = False) -> Path:
        value = interpolate(raw, variables)
        path = Path(str(value)).expanduser()
        if not path.is_absolute():
            path = self.root / path
        path = path.resolve()
        self._assert_path(path)
        if for_write:
            path.parent.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def _atomic_write(path: Path, data: bytes) -> None:
        fd,temporary=tempfile.mkstemp(prefix=".openrpa-write-",dir=path.parent)
        try:
            with os.fdopen(fd,"wb") as stream:
                stream.write(data);stream.flush();os.fsync(stream.fileno())
            if path.exists():
                try: os.chmod(temporary,path.stat().st_mode)
                except OSError: pass
            os.replace(temporary,path)
        finally:
            if os.path.exists(temporary): os.unlink(temporary)

    def run(self, workflow_name: str, inputs: dict[str, Any] | None = None) -> RunResult:
        logs: list[dict[str, Any]] = []
        states: dict[str, dict[str, Any]] = {}
        variables: dict[str, Any] = {}
        try:
            binding_problems=validate_workflow_bindings(self.workflows)
            if binding_problems:
                first=binding_problems[0]
                raise WorkflowError(first["message"],first.get("activityId"),first["code"])
            if workflow_name not in self.workflows:
                raise WorkflowError(f"Unknown workflow '{workflow_name}'", code="workflow-not-found")
            if inputs is not None and not isinstance(inputs,dict):
                raise WorkflowError("Workflow inputs must be an object",code="argument-inputs")
            if len(json.dumps(inputs or {},ensure_ascii=False,default=str))>2_000_000:
                raise WorkflowError("Workflow input data exceeds the 2 MB limit",code="input-size-limit")
            wf = self.workflows[workflow_name]
            variables = {v["name"]: v.get("value") for v in wf.get("variables", [])}
            for spec in wf.get("variables", []):
                if not _matches_type(variables[spec["name"]],spec.get("type","any")):
                    raise WorkflowError(f"Variable '{spec['name']}' does not match declared type {spec.get('type')}",code="variable-type")
            args = {a["name"]: a for a in wf.get("arguments", [])}
            unknown=set(inputs or {})-set(args)
            if unknown:
                raise WorkflowError(f"Unknown input argument(s): {', '.join(sorted(map(str,unknown)))}",code="argument-unknown")
            for name, spec in args.items():
                if name in (inputs or {}):
                    variables[name] = (inputs or {})[name]
                elif "default" in spec:
                    variables[name] = spec["default"]
                elif spec.get("direction", "in") in {"in", "inout"} and spec.get("required", True):
                    raise WorkflowError(f"Missing required input argument '{name}'", code="argument-missing")
                if name in variables and not _matches_type(variables[name],spec.get("type","any")):
                    raise WorkflowError(f"Argument '{name}' does not match declared type {spec.get('type')}",code="argument-type")
            if wf.get("executionModel","sequence")=="flowchart":
                self._flowchart(wf["flowchart"],variables,logs,states,workflow_name)
            elif wf.get("executionModel")=="state_machine":
                self._state_machine(wf["stateMachine"],variables,logs,states,workflow_name)
            else:
                self._sequence(wf["activities"], variables, logs, states, workflow_name)
            return RunResult("succeeded", self._redact_value(variables), logs, self._redact_value(states))
        except RunCancelled:
            return RunResult("cancelled", self._redact_value(variables), logs, self._redact_value(states))
        except Exception as exc:
            aid = getattr(exc, "activity_id", None)
            frames=getattr(exc,"workflow_stack",[])
            safe_frames=[{"workflow":self._redact_text(str(frame.get("workflow","")))[:200],"activityId":self._redact_text(str(frame.get("activityId","")))[:128],"targetWorkflow":self._redact_text(str(frame.get("targetWorkflow","")))[:200]}
                for frame in frames[:40] if isinstance(frame,dict)] if isinstance(frames,list) else []
            message=self._redact_text(str(exc))
            if safe_frames:
                path=" → ".join(f"{frame['workflow']}/{frame['activityId']} → {frame['targetWorkflow']}" for frame in safe_frames)
                message=f"{message} · Call path: {path}"
            error = {"code": getattr(exc, "code", "activity-failed"), "message": message, "activityId": aid,
                "workflow":getattr(exc,"workflow",workflow_name),"workflowStack":safe_frames,"type": type(exc).__name__}
            return RunResult("failed", self._redact_value(variables), logs, self._redact_value(states), error)
        finally:
            self._close_browser()

    def _check_cancel(self) -> None:
        if self.options.cancel():
            raise RunCancelled()
        if time.monotonic()-self._run_started>self._max_run_seconds:
            raise WorkflowError(f"Run exceeded its {self._max_run_seconds} second execution quota",code="run-timeout")

    def _emit(self, event: dict[str, Any], logs: list[dict[str, Any]]) -> None:
        event = self._redact_value(event)
        if len(logs) >= max(1, min(int(self.options.max_log_events), 50000)):
            return
        encoded = json.dumps(event, ensure_ascii=False, default=str)
        if len(encoded) > max(256, min(int(self.options.max_log_event_chars), 65536)):
            event = {"event": str(event.get("event", "event")), "activityId": event.get("activityId"),
                     "truncated": True, "preview": encoded[: max(128, self.options.max_log_event_chars - 128)]}
        logs.append(event)
        self.options.event(event)

    def _redact_text(self, text: str) -> str:
        for secret in self._secret_values:
            if secret:
                text=text.replace(secret,"[REDACTED]")
        return text

    def _redact_value(self, value: Any) -> Any:
        if isinstance(value,str): return self._redact_text(value)
        if isinstance(value,list): return [self._redact_value(item) for item in value]
        if isinstance(value,tuple): return [self._redact_value(item) for item in value]
        if isinstance(value,dict): return {key:self._redact_value(item) for key,item in value.items()}
        return value

    def _debug_snapshot(self, variables: dict[str, Any]) -> dict[str, Any]:
        """Return a redacted debugger view small enough for one event/log record."""
        snapshot: dict[str,Any]={}
        used=0
        omitted=0
        for name,value in list(variables.items())[:100]:
            safe=self._redact_value(value)
            try: encoded=json.dumps(safe,ensure_ascii=False,default=str,separators=(",",":"))
            except Exception: encoded=repr(safe)
            if len(encoded)>min(4096,8000-used):
                snapshot[str(name)[:128]]={"type":type(value).__name__,"truncated":True,"preview":encoded[:400]}
                used+=len(encoded[:400])+100
            else:
                snapshot[str(name)[:128]]=safe
                used+=len(encoded)+len(str(name))+16
            if used>=7800:
                omitted=max(0,len(variables)-len(snapshot))
                break
        if len(variables)>100:
            omitted+=len(variables)-100
        if omitted:
            snapshot["__debugger_note"]={"truncated":True,"omittedVariableCount":omitted,"reason":"Debugger snapshots are bounded to protect the event log."}
        return snapshot

    def _validate_debug_edit(self, workflow_name: str, variables: dict[str,Any], name: str, value: Any) -> None:
        if self._redact_value(variables[name]) != variables[name]:
            raise ValueError("Credential-bearing variables cannot be edited in the debugger")
        workflow=self.workflows.get(workflow_name,{})
        specs={item.get("name"):item for item in [*workflow.get("variables",[]),*workflow.get("arguments",[])] if isinstance(item,dict)}
        spec=specs.get(name)
        if spec and not _matches_type(value,spec.get("type","any")):
            raise ValueError(f"Variable '{name}' must remain type {spec.get('type')}")

    def _debug_editable_variables(self, workflow_name: str, variables: dict[str,Any]) -> list[str]:
        editable=[]
        for name,value in variables.items():
            try:
                encoded=json.dumps(value,ensure_ascii=False,allow_nan=False,separators=(",",":"))
                if len(encoded.encode("utf-8"))>65536: continue
                self._validate_debug_edit(workflow_name,variables,name,value)
                editable.append(name)
            except (TypeError,ValueError,RecursionError):
                continue
        return editable

    def _sequence(self, activities: list[dict[str, Any]], variables: dict[str, Any], logs: list[dict[str, Any]], states: dict[str, dict[str, Any]], workflow_name: str, depth: int = 0) -> None:
        for activity in activities:
            self._check_cancel()
            self._activity_executions+=1
            if self._activity_executions>self._max_activity_executions:
                raise WorkflowError(f"Run exceeded its {self._max_activity_executions} activity execution quota",activity.get("id"),"activity-quota")
            aid, kind = activity["id"], activity["type"]
            if self.options.debugger is not None:
                self.options.debugger.before_activity(aid,depth,variables,self.options.cancel,lambda event:self._emit(event,logs),
                    snapshot=self._debug_snapshot,editable=lambda current:self._debug_editable_variables(workflow_name,current),
                    validate_edit=lambda name,value:self._validate_debug_edit(workflow_name,variables,name,value),redact=self._redact_value)
                self._check_cancel()
            started = time.perf_counter()
            state = {"status": "running", "type": kind, "workflow": workflow_name}
            states[aid] = state
            self._emit({"event": "activity-started", "activityId": aid, "type": kind}, logs)
            try:
                self._activity(activity, variables, logs, states, workflow_name, depth)
                if self.options.debugger is not None:
                    self.options.debugger.after_activity(aid,depth)
                state.update(status="succeeded", durationMs=round((time.perf_counter() - started) * 1000, 3))
                self._emit({"event": "activity-completed", "activityId": aid, "type": kind, "durationMs": state["durationMs"]}, logs)
            except RunCancelled:
                state["status"] = "cancelled"
                raise
            except Exception as exc:
                state.update(status="failed", error=str(exc), durationMs=round((time.perf_counter() - started) * 1000, 3))
                if not getattr(exc, "activity_id", None):
                    try:
                        setattr(exc, "activity_id", aid)
                    except Exception:
                        pass
                try:
                    if not getattr(exc,"workflow",None):
                        setattr(exc,"workflow",workflow_name)
                    if kind=="invoke_workflow":
                        frames=list(getattr(exc,"workflow_stack",[]))
                        if len(frames)<40:
                            frames.insert(0,{"workflow":workflow_name,"activityId":aid,"targetWorkflow":str(activity.get("workflow",""))})
                            setattr(exc,"workflow_stack",frames)
                except Exception:
                    pass
                self._emit({"event": "activity-failed", "activityId": aid, "type": kind, "error": str(exc)}, logs)
                raise

    def _flowchart(self, graph: dict[str,Any], variables: dict[str,Any], logs: list[dict[str,Any]], states: dict[str,dict[str,Any]], workflow_name: str) -> None:
        nodes={node["id"]:node for node in graph["nodes"]}
        outgoing:dict[str,list[dict[str,Any]]]={node_id:[] for node_id in nodes}
        for transition in graph["transitions"]: outgoing[transition["from"]].append(transition)
        current=graph["start"]
        visits=0
        while current is not None:
            self._check_cancel()
            visits+=1
            if visits>self._max_activity_executions:
                raise WorkflowError(f"Flowchart exceeded its {self._max_activity_executions} node visit quota",current,"activity-quota")
            node=nodes[current];kind=node["kind"]
            if kind=="action":
                aid=current;node_type="flowchart_action";edges=outgoing[current]
                states[aid]={"status":"running","type":node_type,"workflow":workflow_name}
                self._emit({"event":"activity-started","activityId":aid,"type":node_type},logs)
                try:
                    self._sequence([node["activity"]],variables,logs,states,workflow_name)
                    next_node=edges[0]["to"] if edges else None
                    states[aid].update(status="succeeded",durationMs=0,nextNode=next_node)
                    self._emit({"event":"activity-completed","activityId":aid,"type":node_type,"nextNode":next_node},logs)
                    current=next_node
                except RunCancelled:
                    states[aid]["status"]="cancelled"
                    raise
                except Exception as exc:
                    states[aid].update(status="failed",error=str(exc))
                    self._emit({"event":"activity-failed","activityId":aid,"type":node_type,"error":str(exc)},logs)
                    raise
                continue
            aid=current;node_type=f"flowchart_{kind}"
            if self.options.debugger is not None:
                self.options.debugger.before_activity(aid,0,variables,self.options.cancel,lambda event:self._emit(event,logs),
                    snapshot=self._debug_snapshot,editable=lambda current:self._debug_editable_variables(workflow_name,current),
                    validate_edit=lambda name,value:self._validate_debug_edit(workflow_name,variables,name,value),redact=self._redact_value)
                self._check_cancel()
            states[aid]={"status":"running","type":node_type,"workflow":workflow_name}
            self._emit({"event":"activity-started","activityId":aid,"type":node_type},logs)
            try:
                if kind=="end": next_node=None
                elif kind=="decision":
                    outcome=evaluate(node["condition"],variables)
                    if not isinstance(outcome,bool): raise WorkflowError("Flowchart decision must evaluate to true or false",aid,"flowchart-condition-type")
                    edge=next(edge for edge in outgoing[aid] if edge["when"]==("true" if outcome else "false"))
                    next_node=edge["to"]
                else: next_node=outgoing[aid][0]["to"]
                states[aid].update(status="succeeded",durationMs=0,nextNode=next_node)
                self._emit({"event":"activity-completed","activityId":aid,"type":node_type,"nextNode":next_node},logs)
                if self.options.debugger is not None: self.options.debugger.after_activity(aid,0)
                current=next_node
            except Exception as exc:
                states[aid].update(status="failed",error=str(exc))
                if not getattr(exc,"activity_id",None):
                    try: setattr(exc,"activity_id",aid)
                    except Exception: pass
                self._emit({"event":"activity-failed","activityId":aid,"type":node_type,"error":str(exc)},logs)
                raise

    def _state_machine(self, machine: dict[str,Any], variables: dict[str,Any], logs: list[dict[str,Any]], states: dict[str,dict[str,Any]], workflow_name: str) -> None:
        """Run states in order, choosing the first true guarded transition.

        A transition without a condition is an explicit final fallback. State
        visit limits use the same project execution quota as flowcharts.
        """
        state_map={state["id"]:state for state in machine["states"]}
        outgoing:dict[str,list[dict[str,Any]]]={state_id:[] for state_id in state_map}
        for transition in machine["transitions"]:outgoing[transition["from"]].append(transition)
        current=machine["initial"]
        visits=0
        while current is not None:
            self._check_cancel()
            visits+=1
            if visits>self._max_activity_executions:
                raise WorkflowError(f"State machine exceeded its {self._max_activity_executions} state visit quota",current,"activity-quota")
            state=state_map[current]
            if self.options.debugger is not None:
                self.options.debugger.before_activity(current,0,variables,self.options.cancel,lambda event:self._emit(event,logs),
                    snapshot=self._debug_snapshot,editable=lambda active:self._debug_editable_variables(workflow_name,active),
                    validate_edit=lambda name,value:self._validate_debug_edit(workflow_name,variables,name,value),redact=self._redact_value)
                self._check_cancel()
            state_type="state_machine_state"
            states[current]={"status":"running","type":state_type,"workflow":workflow_name,"label":state.get("label",current)}
            self._emit({"event":"activity-started","activityId":current,"type":state_type,"state":current},logs)
            try:
                self._sequence(state.get("activities",[]),variables,logs,states,workflow_name)
                next_state=None
                if not state.get("final",False):
                    fallback=None
                    for transition in outgoing[current]:
                        condition=transition.get("condition")
                        if condition is None:
                            fallback=transition
                            continue
                        outcome=evaluate(condition,variables)
                        if not isinstance(outcome,bool):
                            raise WorkflowError(f"Transition condition from state {current} must evaluate to true or false",current,"state-machine-condition-type")
                        if outcome:
                            next_state=transition["to"]
                            break
                    if next_state is None and fallback is not None:next_state=fallback["to"]
                    if next_state is None:
                        raise WorkflowError(f"No transition condition matched from state {current}",current,"state-machine-no-transition-match")
                states[current].update(status="succeeded",durationMs=0,nextState=next_state)
                self._emit({"event":"activity-completed","activityId":current,"type":state_type,"nextState":next_state},logs)
                if self.options.debugger is not None:self.options.debugger.after_activity(current,0)
                current=next_state
            except Exception as exc:
                states[current].update(status="failed",error=str(exc))
                if not getattr(exc,"activity_id",None):
                    try:setattr(exc,"activity_id",current)
                    except Exception:pass
                self._emit({"event":"activity-failed","activityId":current,"type":state_type,"error":str(exc)},logs)
                raise

    def _activity(self, a: dict[str, Any], v: dict[str, Any], logs: list[dict[str, Any]], states: dict[str, dict[str, Any]], workflow_name: str, depth: int = 0) -> None:
        kind = a["type"]
        if kind == "assign":
            v[a["target"]] = interpolate(a.get("value"), v)
        elif kind == "credential_get":
            service=str(a.get("service", "")).strip()
            username=str(a.get("username", "")).strip()
            if not service or not username or len(service)>200 or len(username)>200:
                raise WorkflowError("Credential service and username are required", a["id"], "credential-config")
            try:
                keyring=get_safe_keyring()
            except RuntimeError as exc:
                code="credential-runtime-unavailable" if "optional keyring package" in str(exc) else "credential-backend-unsafe"
                raise WorkflowError(str(exc),a["id"],code) from exc
            secret=keyring.get_password(service,username)
            if not isinstance(secret,str) or not secret:
                raise WorkflowError("Credential was not found in the OS credential vault",a["id"],"credential-not-found")
            self._secret_values.add(secret)
            self.options.audit("credential.read", {"activityId":a["id"],"service":service,"username":username})
            v[a.get("result","credential")]=secret
        elif kind == "log":
            message = str(interpolate(a.get("message", ""), v))
            self._emit({"event": "log", "level": a.get("level", "info"), "activityId": a["id"], "message": message}, logs)
        elif kind == "if":
            self._sequence(a.get("then", []) if evaluate(a["condition"], v) else a.get("else", []), v, logs, states, workflow_name,depth+1)
        elif kind == "switch":
            actual=evaluate(a["expression"],v)
            selected=a.get("default",[])
            for case in a.get("cases",[]):
                candidate=case.get("value")
                numeric_pair=isinstance(actual,(int,float)) and not isinstance(actual,bool) and isinstance(candidate,(int,float)) and not isinstance(candidate,bool)
                if (numeric_pair or type(actual) is type(candidate)) and actual==candidate:
                    selected=case.get("activities",[])
                    break
            self._sequence(selected,v,logs,states,workflow_name,depth+1)
        elif kind == "for_each":
            values = evaluate(a["items"], v)
            if not isinstance(values, (list, tuple, str)):
                raise WorkflowError("For Each expression must produce a list, tuple, or string", a["id"], "loop-input")
            if len(values) > self._max_loop_items:
                raise WorkflowError("For Each exceeded the configured item limit", a["id"], "loop-limit")
            item_name = a["itemName"]
            had_old, old = item_name in v, v.get(item_name)
            for value in values:
                self._check_cancel()
                self._loop_items_executed+=1
                if self._loop_items_executed>1_000_000:
                    raise WorkflowError("Workflow exceeded the 1,000,000 total loop item quota",a["id"],"loop-quota")
                v[item_name] = value
                self._sequence(a.get("body", []), v, logs, states, workflow_name,depth+1)
            if had_old:
                v[item_name] = old
            else:
                v.pop(item_name, None)
        elif kind == "while":
            condition=a.get("condition","")
            while True:
                self._check_cancel()
                outcome=evaluate(condition,v)
                if not isinstance(outcome,bool):
                    raise WorkflowError("While condition must evaluate to true or false",a["id"],"while-condition-type")
                if not outcome:
                    break
                self._loop_items_executed+=1
                if self._loop_items_executed>self._max_loop_items or self._loop_items_executed>1_000_000:
                    raise WorkflowError("While exceeded the configured loop item limit",a["id"],"loop-limit")
                self._sequence(a.get("body",[]),v,logs,states,workflow_name,depth+1)
        elif kind == "retry":
            max_attempts=int(a["maxAttempts"])
            retry_on=set(a.get("retryOn",[]))
            delay=float(a.get("delaySeconds",0))
            factor=float(a.get("backoffFactor",1))
            for attempt in range(1,max_attempts+1):
                self._check_cancel()
                variable_checkpoint=copy_module.deepcopy(v)
                try:
                    self._sequence(a["body"],v,logs,states,workflow_name,depth+1)
                    if attempt>1:
                        self._emit({"event":"retry-succeeded","activityId":a["id"],"attempt":attempt,"maxAttempts":max_attempts},logs)
                    break
                except RunCancelled:
                    raise
                except Exception as exc:
                    error_code=getattr(exc,"code","activity-failed")
                    can_retry=attempt<max_attempts and error_code in retry_on
                    self._emit({"event":"retry-attempt-failed","activityId":a["id"],"attempt":attempt,"maxAttempts":max_attempts,"errorCode":error_code,"willRetry":can_retry},logs)
                    if not can_retry:
                        raise
                    v.clear();v.update(variable_checkpoint)
                    remaining=min(delay*(factor**(attempt-1)),60)
                    deadline=time.monotonic()+remaining
                    while time.monotonic()<deadline:
                        self._check_cancel()
                        time.sleep(min(.1,deadline-time.monotonic()))
        elif kind == "try":
            failure: Exception | None = None
            try:
                self._sequence(a.get("body", []), v, logs, states, workflow_name,depth+1)
            except Exception as exc:
                if isinstance(exc, RunCancelled):
                    raise
                failure = exc
                if a.get("errorName"):
                    v[a["errorName"]] = {"message": str(exc), "type": type(exc).__name__}
                self._sequence(a.get("catch", []), v, logs, states, workflow_name,depth+1)
            finally:
                self._sequence(a.get("finally", []), v, logs, states, workflow_name,depth+1)
            if failure is not None and not a.get("catch"):
                raise failure
        elif kind == "invoke_workflow":
            target = a["workflow"]
            raw_inputs=a.get("inputs",{})
            if not isinstance(raw_inputs,dict):
                raise WorkflowError("Invoked workflow inputs must be an object",a["id"],"argument-inputs")
            raw_outputs=a.get("outputs",{})
            if not isinstance(raw_outputs,dict):
                raise WorkflowError("Invoked workflow outputs must be an object",a["id"],"argument-outputs")
            child_args = {k: interpolate(value, v) for k, value in raw_inputs.items()}
            child_wf = self.workflows[target]
            child_vars = {item["name"]: item.get("value") for item in child_wf.get("variables", [])}
            child_specs={arg["name"]:arg for arg in child_wf.get("arguments", [])}
            unknown=set(child_args)-set(child_specs)
            if unknown:
                raise WorkflowError(f"Unknown invoked workflow input(s): {', '.join(sorted(map(str,unknown)))}",a["id"],"argument-unknown")
            invalid_outputs=set(raw_outputs)-{name for name,arg in child_specs.items() if arg.get("direction","in") in {"out","inout"}}
            if invalid_outputs:
                raise WorkflowError(f"Unknown invoked workflow output(s): {', '.join(sorted(map(str,invalid_outputs)))}",a["id"],"argument-output")
            for arg in child_specs.values():
                direction=arg.get("direction","in")
                if arg["name"] in child_args:
                    if direction=="out":
                        raise WorkflowError(f"Output-only argument '{arg['name']}' cannot be supplied as input",a["id"],"argument-direction")
                    child_vars[arg["name"]] = child_args[arg["name"]]
                elif "default" in arg:
                    child_vars[arg["name"]] = arg["default"]
                elif direction in {"in","inout"} and arg.get("required", True):
                    raise WorkflowError(f"Missing invoked workflow argument '{arg['name']}'", a["id"], "argument-missing")
                if arg["name"] in child_vars and not _matches_type(child_vars[arg["name"]],arg.get("type","any")):
                    raise WorkflowError(f"Invoked workflow argument '{arg['name']}' does not match declared type {arg.get('type')}",a["id"],"argument-type")
            before = dict(v)
            if child_wf.get("executionModel","sequence")=="flowchart":
                self._flowchart(child_wf["flowchart"],child_vars,logs,states,target)
            elif child_wf.get("executionModel")=="state_machine":
                self._state_machine(child_wf["stateMachine"],child_vars,logs,states,target)
            else:
                self._sequence(child_wf["activities"], child_vars, logs, states, target,depth+1)
            for arg in child_specs.values():
                if arg.get("direction") in {"out", "inout"} and arg["name"] in child_vars:
                    destination=raw_outputs.get(arg["name"],arg.get("target",arg["name"]))
                    if not isinstance(destination,str) or not destination.isidentifier() or len(destination)>256:
                        raise WorkflowError(f"Output '{arg['name']}' needs a valid caller variable name",a["id"],"argument-output-target")
                    v[destination] = child_vars[arg["name"]]
            if a.get("resultTarget"):
                v[a["resultTarget"]] = {k: value for k, value in child_vars.items() if k not in before}
        elif kind == "file_read_text":
            path = self._path(a["path"], v)
            max_bytes = self._max_file_bytes
            if path.stat().st_size > max_bytes: raise WorkflowError("File exceeds configured read limit", a["id"], "file-size-limit")
            v[a.get("result", "text")] = path.read_text(encoding=a.get("encoding", "utf-8"))
        elif kind == "file_write_text":
            output=str(interpolate(a.get("text", ""), v))
            if len(output.encode("utf-8"))>self._max_file_bytes: raise WorkflowError("Output file exceeds configured write limit",a["id"],"file-size-limit")
            self._atomic_write(self._path(a["path"], v, True),output.encode(a.get("encoding", "utf-8")))
        elif kind == "file_copy":
            src, dst = self._path(a["source"], v), self._path(a["destination"], v, True)
            if src.stat().st_size>self._max_file_bytes: raise WorkflowError("Copy source exceeds configured file limit",a["id"],"file-size-limit")
            shutil.copy2(src, dst)
        elif kind == "file_move":
            src, dst = self._path(a["source"], v), self._path(a["destination"], v, True)
            shutil.move(str(src), str(dst))
        elif kind == "file_delete":
            target = self._path(a["path"], v)
            if target.is_dir():
                raise WorkflowError("File Delete does not remove directories", a["id"], "file-type")
            target.unlink(missing_ok=bool(a.get("missingOk", False)))
        elif kind == "file_list":
            directory = self._path(a["directory"], v)
            pattern = str(a.get("pattern", "*"))
            matches=[str(p.relative_to(self.root)) for p in sorted(directory.glob(pattern)) if p.is_file()]
            if len(matches)>100_000: raise WorkflowError("File listing exceeds 100,000 entries",a["id"],"file-size-limit")
            v[a.get("result", "files")] = matches
        elif kind in {"regex_replace", "regex_match"}:
            source = str(interpolate(a.get("text", ""), v))
            pattern = str(interpolate(a.get("pattern", ""), v))
            if len(pattern)>1000 or len(source)>200_000: raise WorkflowError("Regex pattern or input exceeds its size limit",a["id"],"regex-size-limit")
            try:
                import regex as bounded_regex
                flags=bounded_regex.IGNORECASE if a.get("ignoreCase") else 0
                if kind == "regex_replace":
                    v[a.get("result", "text")] = bounded_regex.sub(pattern, str(interpolate(a.get("replacement", ""), v)), source, flags=flags, timeout=0.25)
                else:
                    match = bounded_regex.search(pattern, source, flags=flags, timeout=0.25)
                    v[a.get("result", "match")] = match.groupdict() if match and match.groupdict() else (match.group(0) if match else None)
            except TimeoutError as exc:
                raise WorkflowError("Regex operation exceeded its 250 ms time quota",a["id"],"regex-timeout") from exc
        elif kind == "transform_field":
            rows = interpolate(a.get("rows", "rows"), v)
            if isinstance(rows, str): rows = evaluate(rows, v)
            field = str(a["field"]); operation = a["operation"]
            result = []
            for row in rows:
                copy = dict(row)
                if field in copy and copy[field] is not None:
                    value = str(copy[field])
                    if operation == "trim": value = value.strip()
                    elif operation == "upper": value = value.upper()
                    elif operation == "lower": value = value.lower()
                    elif operation == "regex_replace":
                        if len(value)>200_000 or len(str(a.get("pattern","")))>1000: raise WorkflowError("Regex pattern or input exceeds its size limit",a["id"],"regex-size-limit")
                        try:
                            import regex as bounded_regex
                            value=bounded_regex.sub(str(a.get("pattern","")),str(a.get("replacement","")),value,timeout=0.25)
                        except TimeoutError as exc: raise WorkflowError("Regex operation exceeded its 250 ms time quota",a["id"],"regex-timeout") from exc
                    else: raise WorkflowError(f"Unsupported field transform: {operation}", a["id"], "transform-operation")
                    copy[field] = value
                result.append(copy)
            v[a.get("result", "rows")] = result
        elif kind == "filter_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str): rows=evaluate(rows,v)
            conditions=a.get("conditions")
            match_mode=a.get("match","all")
            allowed={"equals","not_equals","contains","not_contains","starts_with","ends_with","greater_than","greater_or_equal","less_than","less_or_equal","is_empty","is_not_empty"}
            if not isinstance(rows,list) or len(rows)>100_000:
                raise WorkflowError("Filter Rows input must be a list of at most 100,000 records",a["id"],"filter-input")
            if not isinstance(conditions,list) or not 1<=len(conditions)<=100:
                raise WorkflowError("Filter Rows needs 1–100 conditions",a["id"],"filter-conditions")
            if match_mode not in {"all","any"}:
                raise WorkflowError("Filter Rows match must be 'all' or 'any'",a["id"],"filter-match")
            for index,condition in enumerate(conditions):
                if not isinstance(condition,dict) or not isinstance(condition.get("field"),str) or not condition["field"].strip() or len(condition["field"])>256 or condition.get("operator") not in allowed:
                    raise WorkflowError(f"Filter Rows condition {index+1} needs a field and supported operator",a["id"],"filter-condition")
            def condition_matches(row:dict[str,Any],condition:dict[str,Any])->bool:
                operator=condition["operator"]
                actual=row.get(condition["field"])
                expected=interpolate(condition.get("value"),v)
                if operator=="is_empty": return actual is None or actual=="" or actual==[] or actual=={}
                if operator=="is_not_empty": return not (actual is None or actual=="" or actual==[] or actual=={})
                if actual is None: return operator=="not_equals" and expected is not None
                if operator in {"equals","not_equals"}:
                    if isinstance(actual,(int,float)) and not isinstance(actual,bool) and isinstance(expected,(int,float)) and not isinstance(expected,bool): equal=actual==expected
                    elif type(actual) is type(expected): equal=actual==expected
                    else: equal=False
                    if condition.get("ignoreCase") and isinstance(actual,str) and isinstance(expected,str): equal=actual.casefold()==expected.casefold()
                    return equal if operator=="equals" else not equal
                if operator in {"contains","not_contains","starts_with","ends_with"}:
                    if condition.get("ignoreCase") and isinstance(actual,str) and isinstance(expected,str): actual,expected=actual.casefold(),expected.casefold()
                    if isinstance(actual,str) and isinstance(expected,str): found=expected in actual
                    elif isinstance(actual,list): found=any(type(item) is type(expected) and item==expected for item in actual)
                    else: found=False
                    if operator=="starts_with": return isinstance(actual,str) and isinstance(expected,str) and actual.startswith(expected)
                    if operator=="ends_with": return isinstance(actual,str) and isinstance(expected,str) and actual.endswith(expected)
                    return found if operator=="contains" else not found
                numeric=lambda value:isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value)
                if numeric(actual) and numeric(expected): left,right=actual,expected
                elif isinstance(actual,str) and isinstance(expected,str): left,right=actual,expected
                else: return False
                return {"greater_than":left>right,"greater_or_equal":left>=right,"less_than":left<right,"less_or_equal":left<=right}[operator]
            result=[]
            for index,row in enumerate(rows):
                if index%256==0: self._check_cancel()
                if not isinstance(row,dict): raise WorkflowError(f"Filter Rows input record {index+1} is not an object",a["id"],"filter-row-type")
                matched=all(condition_matches(row,item) for item in conditions) if match_mode=="all" else any(condition_matches(row,item) for item in conditions)
                if matched: result.append(row)
            v[a.get("result","filteredRows")]=result
            self._emit({"event":"data-filtered","activityId":a["id"],"inputRows":len(rows),"outputRows":len(result),"conditionCount":len(conditions)},logs)
        elif kind == "join_rows":
            left=interpolate(a.get("left","leftRows"),v);right=interpolate(a.get("right","rightRows"),v)
            if isinstance(left,str): left=evaluate(left,v)
            if isinstance(right,str): right=evaluate(right,v)
            if not isinstance(left,list) or not isinstance(right,list) or len(left)>100_000 or len(right)>100_000:
                raise WorkflowError("Join Rows inputs must each be lists of at most 100,000 records",a["id"],"join-input")
            left_key=a.get("leftKey");right_key=a.get("rightKey");mode=a.get("joinType","inner");prefix=a.get("rightPrefix","right_");max_rows=a.get("maxRows",100_000)
            if not isinstance(left_key,str) or not left_key or not isinstance(right_key,str) or not right_key:
                raise WorkflowError("Join Rows needs non-empty left and right key fields",a["id"],"join-key")
            if mode not in {"inner","left"} or not isinstance(prefix,str) or len(prefix)>128 or not isinstance(max_rows,int) or isinstance(max_rows,bool) or not 1<=max_rows<=100_000:
                raise WorkflowError("Join Rows configuration is invalid",a["id"],"join-config")
            for label,records in (("left",left),("right",right)):
                if any(not isinstance(row,dict) or len(row)>2_000 for row in records):
                    raise WorkflowError(f"Join Rows {label} input must contain objects with at most 2,000 fields",a["id"],"join-row-type")
            left_fields:set[str]=set()
            right_fields:set[str]=set()
            def join_key(value:Any)->tuple[str,Any]|None:
                if value is None:return None
                if isinstance(value,bool):return ("boolean",value)
                if isinstance(value,str):return ("string",value)
                if isinstance(value,int):return ("number",Decimal(value))
                if isinstance(value,float) and math.isfinite(value):return ("number",Decimal(str(value)))
                return None
            index:dict[tuple[str,Any],list[dict[str,Any]]]={}
            for position,row in enumerate(right):
                if position%256==0:self._check_cancel()
                right_fields.update(row.keys())
                key=join_key(row.get(right_key))
                if key is not None:index.setdefault(key,[]).append(row)
            for position,row in enumerate(left):
                if position%256==0:self._check_cancel()
                left_fields.update(row.keys())
            if any(prefix+field in left_fields for field in right_fields):
                raise WorkflowError("A prefixed right-side field would overwrite a left-side field; choose another rightPrefix",a["id"],"join-field-collision")
            result=[]
            for position,row in enumerate(left):
                if position%256==0:self._check_cancel()
                matches=index.get(join_key(row.get(left_key)),[]) if join_key(row.get(left_key)) is not None else []
                if not matches and mode=="left":matches=[None]
                for other in matches:
                    if len(result)%256==0:self._check_cancel()
                    if len(result)>=max_rows:raise WorkflowError(f"Join Rows output exceeds the configured {max_rows:,} row limit",a["id"],"join-output-limit")
                    merged=dict(row)
                    if other is not None:
                        for field,value in other.items():merged[prefix+field]=value
                    else:
                        for field in (right[0].keys() if right else []):merged[prefix+field]=None
                    result.append(merged)
            name=a.get("result","joinedRows");v[name]=result
            self._emit({"event":"data-joined","activityId":a["id"],"leftRows":len(left),"rightRows":len(right),"outputRows":len(result),"joinType":mode,"matchedLeftRows":sum(1 for row in left if join_key(row.get(left_key)) in index)},logs)
        elif kind == "group_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            group_by=a.get("groupBy",[]);aggregations=a.get("aggregations",[]);max_groups=a.get("maxGroups",100_000)
            if not isinstance(rows,list) or len(rows)>100_000:
                raise WorkflowError("Group Rows input must be a list of at most 100,000 records",a["id"],"group-input")
            if not isinstance(group_by,list) or len(group_by)>100 or any(not isinstance(field,str) or not field.strip() or len(field)>256 for field in group_by) or len(set(group_by))!=len(group_by):
                raise WorkflowError("Group Rows needs up to 100 unique, non-empty group fields",a["id"],"group-fields")
            operations={"count_rows","count_values","sum","average","min","max"}
            if not isinstance(aggregations,list) or len(aggregations)>100:
                raise WorkflowError("Group Rows supports up to 100 aggregations",a["id"],"group-aggregations")
            aliases=set(group_by)
            for index,aggregation in enumerate(aggregations):
                if not isinstance(aggregation,dict) or aggregation.get("operation") not in operations or not isinstance(aggregation.get("as"),str) or not aggregation["as"].strip() or len(aggregation["as"])>256:
                    raise WorkflowError(f"Group Rows aggregation {index+1} needs a supported operation and output field",a["id"],"group-aggregation")
                if aggregation["operation"]!="count_rows" and (not isinstance(aggregation.get("field"),str) or not aggregation["field"].strip() or len(aggregation["field"])>256):
                    raise WorkflowError(f"Group Rows aggregation {index+1} needs an input field",a["id"],"group-aggregation-field")
                if aggregation["as"] in aliases:
                    raise WorkflowError(f"Group Rows output field '{aggregation['as']}' is duplicated",a["id"],"group-output-duplicate")
                aliases.add(aggregation["as"])
            if not isinstance(max_groups,int) or isinstance(max_groups,bool) or not 1<=max_groups<=100_000:
                raise WorkflowError("Group Rows maxGroups must be an integer from 1 to 100,000",a["id"],"group-limit")
            def group_atom(value:Any)->tuple[str,Any]:
                if value is None:return ("null",None)
                if isinstance(value,bool):return ("boolean",value)
                if isinstance(value,str):return ("string",value)
                if isinstance(value,int):return ("number",Decimal(value))
                if isinstance(value,float) and math.isfinite(value):return ("number",Decimal(str(value)))
                raise WorkflowError("Group Rows can group only finite JSON scalar values",a["id"],"group-key-type")
            groups:dict[tuple[Any,...],dict[str,Any]]={}
            for position,row in enumerate(rows):
                if position%256==0:self._check_cancel()
                if not isinstance(row,dict) or len(row)>2_000:
                    raise WorkflowError(f"Group Rows record {position+1} must be an object with at most 2,000 fields",a["id"],"group-row-type")
                values=[row.get(field) for field in group_by]
                key=tuple(group_atom(value) for value in values)
                group=groups.get(key)
                if group is None:
                    if len(groups)>=max_groups:raise WorkflowError(f"Group Rows exceeds the configured {max_groups:,} group limit",a["id"],"group-limit")
                    group={"keys":dict(zip(group_by,values)),"counts":[0]*len(aggregations),"totals":[Decimal(0) for _ in aggregations],"values":[None]*len(aggregations),"types":[None]*len(aggregations)}
                    groups[key]=group
                for agg_index,aggregation in enumerate(aggregations):
                    operation=aggregation["operation"]
                    if operation=="count_rows":group["counts"][agg_index]+=1;continue
                    value=row.get(aggregation.get("field"))
                    if value is None:continue
                    if operation=="count_values":group["counts"][agg_index]+=1;continue
                    if operation in {"sum","average"}:
                        if isinstance(value,bool) or not isinstance(value,(int,float)) or (isinstance(value,float) and not math.isfinite(value)):
                            raise WorkflowError(f"Group Rows {operation} field '{aggregation['field']}' must contain only finite numbers or nulls",a["id"],"group-number-type")
                        amount=Decimal(value) if isinstance(value,int) else Decimal(str(value))
                        current=group["totals"][agg_index]
                        precision=max(28,max(current.adjusted(),amount.adjusted())-min(current.as_tuple().exponent,amount.as_tuple().exponent)+len(str(max(1,len(rows))))+2)
                        with localcontext() as decimal_context:
                            decimal_context.prec=precision
                            group["totals"][agg_index]=current+amount
                        group["counts"][agg_index]+=1;continue
                    value_type="number" if isinstance(value,(int,float)) and not isinstance(value,bool) else "string" if isinstance(value,str) else None
                    if value_type is None or isinstance(value,float) and not math.isfinite(value):
                        raise WorkflowError(f"Group Rows {operation} field '{aggregation['field']}' must contain finite numbers, text or nulls",a["id"],"group-order-type")
                    previous_type=group["types"][agg_index]
                    if previous_type is not None and previous_type!=value_type:raise WorkflowError(f"Group Rows {operation} field '{aggregation['field']}' mixes text and numeric values",a["id"],"group-order-type")
                    group["types"][agg_index]=value_type;old=group["values"][agg_index]
                    group["values"][agg_index]=value if old is None or (operation=="min" and value<old) or (operation=="max" and value>old) else old
            result=[]
            for group in groups.values():
                output=dict(group["keys"])
                for agg_index,aggregation in enumerate(aggregations):
                    operation=aggregation["operation"];count=group["counts"][agg_index]
                    if operation in {"count_rows","count_values"}:value=count
                    elif operation=="sum":
                        total=group["totals"][agg_index]
                        if count==0:value=None
                        elif total==total.to_integral_value():value=int(total)
                        else:
                            value=float(total)
                            if not math.isfinite(value):raise WorkflowError(f"Group Rows sum '{aggregation['as']}' exceeds the finite output range",a["id"],"group-number-range")
                    elif operation=="average":
                        value=None if count==0 else float(group["totals"][agg_index]/count)
                        if value is not None and not math.isfinite(value):raise WorkflowError(f"Group Rows average '{aggregation['as']}' exceeds the finite output range",a["id"],"group-number-range")
                    else:value=group["values"][agg_index]
                    output[aggregation["as"]]=value
                result.append(output)
            if not result and not group_by and aggregations:
                result=[{aggregation["as"]:(0 if aggregation["operation"] in {"count_rows","count_values"} else None) for aggregation in aggregations}]
            v[a.get("result","groupedRows")]=result
            self._emit({"event":"data-grouped","activityId":a["id"],"inputRows":len(rows),"outputGroups":len(result),"groupByCount":len(group_by),"aggregationCount":len(aggregations)},logs)
        elif kind == "sort_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            keys=a.get("keys",[]);limit=a.get("maxRows",100_000)
            if not isinstance(limit,int) or isinstance(limit,bool) or not 1<=limit<=100_000:
                raise WorkflowError("Sort Rows maxRows must be an integer from 1 to 100,000",a["id"],"sort-limit")
            if not isinstance(rows,list) or len(rows)>100_000 or len(rows)>limit:
                raise WorkflowError("Sort Rows input must be a list within its configured 100,000 row limit",a["id"],"sort-input")
            if not isinstance(keys,list) or not 1<=len(keys)<=50:
                raise WorkflowError("Sort Rows needs 1–50 sort keys",a["id"],"sort-keys")
            for index,key in enumerate(keys):
                if not isinstance(key,dict) or not isinstance(key.get("field"),str) or not key["field"].strip() or not isinstance(key.get("direction","ascending"),str) or key.get("direction","ascending") not in ("ascending","descending") or not isinstance(key.get("nulls","last"),str) or key.get("nulls","last") not in ("first","last"):
                    raise WorkflowError(f"Sort Rows key {index+1} is invalid",a["id"],"sort-key")
            if len({key["field"] for key in keys})!=len(keys):raise WorkflowError("Sort Rows key fields must be unique",a["id"],"sort-key-duplicate")
            for index,row in enumerate(rows):
                if index%256==0:self._check_cancel()
                if not isinstance(row,dict) or len(row)>2_000:raise WorkflowError(f"Sort Rows record {index+1} must be an object with at most 2,000 fields",a["id"],"sort-row-type")
            from functools import cmp_to_key
            field_types:dict[str,str]={}
            def type_of(value:Any)->str:
                if isinstance(value,bool):return "boolean"
                if isinstance(value,str):return "string"
                if isinstance(value,int):return "number"
                if isinstance(value,float) and math.isfinite(value):return "number"
                raise WorkflowError("Sort Rows supports only finite numbers, text, booleans and null values",a["id"],"sort-value-type")
            for key in keys:
                field=key["field"];types=set()
                for index,row in enumerate(rows):
                    if index%256==0:self._check_cancel()
                    value=row.get(field)
                    if value is not None:types.add(type_of(value))
                if len(types)>1:raise WorkflowError(f"Sort Rows field '{field}' mixes incompatible value types",a["id"],"sort-mixed-types")
                field_types[field]=next(iter(types),"null")
            def compare(left:dict[str,Any],right:dict[str,Any])->int:
                for key in keys:
                    field=key["field"];a_value=left.get(field);b_value=right.get(field)
                    if a_value is None or b_value is None:
                        result=0 if a_value is None and b_value is None else (-1 if a_value is None else 1)
                        if key.get("nulls","last")=="last":result=-result
                    elif field_types[field]=="number":
                        a_num=Decimal(a_value) if isinstance(a_value,int) else Decimal(str(a_value));b_num=Decimal(b_value) if isinstance(b_value,int) else Decimal(str(b_value))
                        result=(a_num>b_num)-(a_num<b_num)
                    else:result=(a_value>b_value)-(a_value<b_value)
                    if result:return -result if key.get("direction","ascending")=="descending" and a_value is not None and b_value is not None else result
                return 0
            result=sorted(rows,key=cmp_to_key(compare))
            v[a.get("result","sortedRows")]=result
            self._emit({"event":"data-sorted","activityId":a["id"],"inputRows":len(rows),"outputRows":len(result),"keyCount":len(keys)},logs)
        elif kind == "map_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            mappings=a.get("mappings",[]);limit=a.get("maxRows",100_000)
            if not isinstance(limit,int) or isinstance(limit,bool) or not 1<=limit<=100_000:raise WorkflowError("Map Rows maxRows must be an integer from 1 to 100,000",a["id"],"map-limit")
            if not isinstance(rows,list) or len(rows)>limit:raise WorkflowError("Map Rows input must be a list within its configured row limit",a["id"],"map-input")
            if not isinstance(mappings,list) or not 1<=len(mappings)<=100:raise WorkflowError("Map Rows needs 1–100 field mappings",a["id"],"map-mappings")
            fields=set()
            for mapping_index,mapping in enumerate(mappings):
                if not isinstance(mapping,dict) or not isinstance(mapping.get("field"),str) or not mapping["field"].strip() or len(mapping["field"])>256 or not isinstance(mapping.get("expression"),str) or not mapping["expression"].strip() or len(mapping["expression"])>4096:raise WorkflowError(f"Map Rows mapping {mapping_index+1} is invalid",a["id"],"map-mapping")
                if mapping["field"] in fields:raise WorkflowError(f"Map Rows output field '{mapping['field']}' is duplicated",a["id"],"map-field-duplicate")
                fields.add(mapping["field"])
            result=[];scope=dict(v);scope["row"]=None;scope["rowIndex"]=0
            for row_index,row in enumerate(rows):
                if row_index%256==0:self._check_cancel()
                if not isinstance(row,dict) or len(row)>2_000:raise WorkflowError(f"Map Rows record {row_index+1} must be an object with at most 2,000 fields",a["id"],"map-row-type")
                scope["row"]=row;scope["rowIndex"]=row_index;output={}
                for mapping in mappings:
                    try:value=evaluate(mapping["expression"],scope)
                    except Exception as exc:raise WorkflowError(f"Map Rows expression for '{mapping['field']}' failed on input row {row_index+1}: {str(exc)[:400]}",a["id"],"map-expression") from exc
                    output[mapping["field"]]=value
                result.append(output)
            v[a.get("result","mappedRows")]=result
            self._emit({"event":"data-mapped","activityId":a["id"],"inputRows":len(rows),"outputRows":len(result),"mappedFieldCount":len(mappings)},logs)
        elif kind == "distinct_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            keys=a.get("keys",[]);keep=a.get("keep","first");limit=a.get("maxRows",100_000)
            if not isinstance(limit,int) or isinstance(limit,bool) or not 1<=limit<=100_000:raise WorkflowError("Distinct Rows maxRows must be an integer from 1 to 100,000",a["id"],"distinct-limit")
            if not isinstance(rows,list) or len(rows)>limit:raise WorkflowError("Distinct Rows input must be a list within its configured row limit",a["id"],"distinct-input")
            if not isinstance(keys,list) or not 1<=len(keys)<=100 or any(not isinstance(key,str) or not key.strip() or len(key)>256 for key in keys):raise WorkflowError("Distinct Rows needs 1–100 non-empty key fields",a["id"],"distinct-keys")
            if len(set(keys))!=len(keys):raise WorkflowError("Distinct Rows key fields must be unique",a["id"],"distinct-key-duplicate")
            if not isinstance(keep,str) or keep not in {"first","last"}:raise WorkflowError("Distinct Rows keep must be 'first' or 'last'",a["id"],"distinct-keep")
            seen=set();selected=[]
            def atom(value:Any)->tuple[str,Any]:
                if value is None:return ("null",None)
                if isinstance(value,bool):return ("boolean",value)
                if isinstance(value,str):return ("string",value)
                if isinstance(value,int) and abs(value)<=9_007_199_254_740_991:return ("number",Decimal(value))
                if isinstance(value,float) and math.isfinite(value):return ("number",Decimal(str(value)))
                raise WorkflowError("Distinct Rows key values must be null, text, booleans, or finite safe numbers",a["id"],"distinct-key-type")
            for index,row in enumerate(rows):
                if index%256==0:self._check_cancel()
                if not isinstance(row,dict) or len(row)>2_000:raise WorkflowError(f"Distinct Rows record {index+1} must be an object with at most 2,000 fields",a["id"],"distinct-row-type")
                signature=tuple(atom(row.get(field)) for field in keys)
                if keep=="first":
                    if signature not in seen:seen.add(signature);selected.append(row)
                else:selected.append((signature,row))
            if keep=="last":
                last_indices={}
                for index,(signature,_row) in enumerate(selected):last_indices[signature]=index
                selected=[row for index,(signature,row) in enumerate(selected) if last_indices[signature]==index]
            v[a.get("result","distinctRows")]=selected
            self._emit({"event":"data-distinct","activityId":a["id"],"inputRows":len(rows),"outputRows":len(selected),"duplicateRowsRemoved":len(rows)-len(selected),"keyCount":len(keys),"kept":keep},logs)
        elif kind == "validate_rows":
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            limit=a.get("maxRows",100_000);rules=a.get("rules",[])
            if not isinstance(limit,int) or isinstance(limit,bool) or not 1<=limit<=100_000:raise WorkflowError("Validate Rows maxRows must be an integer from 1 to 100,000",a["id"],"validate-limit")
            if not isinstance(rows,list) or len(rows)>limit:raise WorkflowError("Validate Rows input must be a list within its configured row limit",a["id"],"validate-input")
            if not isinstance(rules,list) or not 1<=len(rules)<=100:raise WorkflowError("Validate Rows needs 1–100 field rules",a["id"],"validate-rules")
            valid=[];invalid=[]
            for index,row in enumerate(rows):
                if index%256==0:self._check_cancel()
                if not isinstance(row,dict) or len(row)>2_000:raise WorkflowError(f"Validate Rows record {index+1} must be an object with at most 2,000 fields",a["id"],"validate-row-type")
                errors=[]
                for rule in rules:
                    field=rule["field"];expected=rule["type"];present=field in row;value=row.get(field)
                    if rule.get("required",False) and (not present or value is None):errors.append({"field":field,"code":"required","message":"Required value is missing"});continue
                    if not present or value is None:continue
                    safe_number=isinstance(value,int) and not isinstance(value,bool) and abs(value)<=9_007_199_254_740_991 or isinstance(value,float) and math.isfinite(value) and (not value.is_integer() or abs(value)<=9_007_199_254_740_991)
                    type_ok=(expected=="text" and isinstance(value,str)) or (expected=="number" and safe_number) or (expected=="boolean" and isinstance(value,bool)) or (expected=="object" and isinstance(value,dict)) or (expected=="array" and isinstance(value,list))
                    if not type_ok:errors.append({"field":field,"code":"type","message":f"Expected {expected}"});continue
                    if rule.get("nonEmpty",False) and ((isinstance(value,str) and not value.strip()) or (isinstance(value,(dict,list)) and not value)):
                        errors.append({"field":field,"code":"empty","message":"Value must not be empty"})
                if errors:invalid.append({"row":row,"errors":errors})
                else:valid.append(row)
            valid_name=a.get("validResult","validRows");invalid_name=a.get("invalidResult","invalidRows")
            v[valid_name]=valid;v[invalid_name]=invalid
            self._emit({"event":"data-validated","activityId":a["id"],"inputRows":len(rows),"validRows":len(valid),"invalidRows":len(invalid),"ruleCount":len(rules)},logs)
        elif kind == "csv_read":
            path = self._path(a["path"], v)
            max_bytes = self._max_file_bytes
            if path.stat().st_size > max_bytes: raise WorkflowError("CSV exceeds configured read limit", a["id"], "file-size-limit")
            with path.open("r", newline="", encoding=a.get("encoding", "utf-8-sig")) as f:
                v[a.get("result", "rows")] = list(csv.DictReader(f, delimiter=a.get("delimiter", ",")))
        elif kind == "csv_write":
            rows = interpolate(a.get("rows", ""), v)
            if isinstance(rows, str):
                rows = evaluate(rows, v)
            if not isinstance(rows, list):
                raise WorkflowError("CSV Write rows must be a list of objects", a["id"], "csv-input")
            fields = a.get("fields") or list(dict.fromkeys(k for row in rows for k in row.keys()))
            if len(rows)>100_000 or len(fields)>2_000 or sum(sum(len(str(value)) for value in row.values()) for row in rows)>self._max_file_bytes:
                raise WorkflowError("CSV output exceeds its row, field, or size quota",a["id"],"file-size-limit")
            rendered=io.StringIO(newline="")
            writer = csv.DictWriter(rendered, fieldnames=fields, delimiter=a.get("delimiter", ","), extrasaction="ignore")
            writer.writeheader();writer.writerows(rows)
            encoded=rendered.getvalue().encode(a.get("encoding", "utf-8"))
            if len(encoded)>self._max_file_bytes: raise WorkflowError("CSV output exceeds configured file size limit",a["id"],"file-size-limit")
            self._atomic_write(self._path(a["path"], v, True),encoded)
        elif kind == "xlsx_read":
            path=self._path(a["path"],v)
            if path.suffix.lower()!='.xlsx': raise WorkflowError("Spreadsheet Read supports .xlsx files only",a["id"],"spreadsheet-format")
            if path.stat().st_size>self._max_file_bytes: raise WorkflowError("Spreadsheet exceeds configured read limit",a["id"],"file-size-limit")
            try:
                with zipfile.ZipFile(path) as archive:
                    infos=archive.infolist()
                    if len(infos)>2_000 or sum(info.file_size for info in infos)>min(self._max_file_bytes*10,200_000_000):
                        raise WorkflowError("Spreadsheet expands beyond its archive quota",a["id"],"file-size-limit")
                    for info in infos:
                        if not info.filename.lower().endswith((".xml",".rels")):continue
                        tail=b""
                        with archive.open(info) as stream:
                            while chunk:=stream.read(65_536):
                                self._check_cancel()
                                probe=(tail+chunk).upper()
                                if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
                                    raise WorkflowError("Spreadsheet XML must not contain DTD or entity declarations",a["id"],"spreadsheet-xml-entity")
                                tail=probe[-16:]
                from openpyxl import load_workbook
                workbook=load_workbook(path,read_only=True,data_only=bool(a.get("dataOnly",True)))
            except WorkflowError: raise
            except Exception as exc: raise WorkflowError(f"Cannot open XLSX spreadsheet: {exc}",a["id"],"spreadsheet-read-error") from exc
            try:
                sheet_name=str(a.get("sheetName") or workbook.sheetnames[0])
                if sheet_name not in workbook.sheetnames: raise WorkflowError(f"Worksheet not found: {sheet_name}",a["id"],"spreadsheet-sheet")
                sheet=workbook[sheet_name];header_row=int(a.get("headerRow",1))
                if header_row>sheet.max_row:raise WorkflowError("Spreadsheet headerRow is beyond the worksheet data",a["id"],"spreadsheet-header")
                headers=[cell.value for cell in next(sheet.iter_rows(min_row=header_row,max_row=header_row))]
                while headers and headers[-1] is None:headers.pop()
                if not headers or len(headers)>2_000 or any(not isinstance(value,(str,int,float)) or not str(value).strip() for value in headers):
                    raise WorkflowError("Spreadsheet header row needs 1–2,000 non-empty text or numeric headers",a["id"],"spreadsheet-header")
                headers=[str(value).strip() for value in headers]
                if len(headers)!=len(set(headers)): raise WorkflowError("Spreadsheet header names must be unique",a["id"],"spreadsheet-header-duplicate")
                rows=[];text_bytes=0
                for values in sheet.iter_rows(min_row=header_row+1,max_col=len(headers),values_only=True):
                    self._check_cancel()
                    if all(value is None for value in values):continue
                    if len(rows)>=100_000:raise WorkflowError("Spreadsheet exceeds 100,000 data rows",a["id"],"spreadsheet-row-limit")
                    if (len(rows)+1)*len(headers)>2_000_000:raise WorkflowError("Spreadsheet exceeds two million cells",a["id"],"spreadsheet-cell-limit")
                    row={}
                    for name,value in zip(headers,values):
                        if isinstance(value,(datetime.date,datetime.datetime,datetime.time)):value=value.isoformat()
                        if isinstance(value,str):text_bytes+=len(value.encode("utf-8"))
                        if text_bytes>self._max_file_bytes:raise WorkflowError("Spreadsheet cell text exceeds configured data limit",a["id"],"file-size-limit")
                        row[name]=value
                    rows.append(row)
                v[a.get("result","rows")]=rows
            finally:workbook.close()
        elif kind == "pdf_read_text":
            path=self._path(a["path"],v)
            if path.suffix.lower()!=".pdf":raise WorkflowError("PDF Read Text supports .pdf files only",a["id"],"pdf-format")
            if path.stat().st_size>self._max_file_bytes:raise WorkflowError("PDF exceeds configured read limit",a["id"],"file-size-limit")
            self._check_cancel()
            try:
                from pypdf import PdfReader
                reader=PdfReader(path,strict=True)
                if reader.is_encrypted:raise WorkflowError("Encrypted PDFs are not supported; provide an authorised unencrypted copy",a["id"],"pdf-encrypted")
                page_count=len(reader.pages)
                max_pages=int(a.get("maxPages",500))
                if page_count>max_pages:raise WorkflowError(f"PDF has {page_count} pages, exceeding the configured {max_pages}-page limit",a["id"],"pdf-page-limit")
                pages=[];total_bytes=0
                for page in reader.pages:
                    self._check_cancel()
                    text=page.extract_text(extraction_mode="layout" if a.get("layoutMode",False) else "plain") or ""
                    total_bytes+=len(text.encode("utf-8"))
                    if total_bytes>self._max_file_bytes:raise WorkflowError("Extracted PDF text exceeds configured data limit",a["id"],"file-size-limit")
                    pages.append(text)
                v[a.get("result","text") ]="\n\f\n".join(pages)
                v[a.get("pageCountResult","pageCount")]=page_count
            except WorkflowError:raise
            except Exception as exc:raise WorkflowError(f"Cannot extract text from PDF: {type(exc).__name__}: {exc}",a["id"],"pdf-read-error") from exc
        elif kind == "image_ocr":
            path=self._path(a["path"],v)
            if path.suffix.lower() not in {".png",".jpg",".jpeg",".bmp"}:raise WorkflowError("OCR Image supports PNG, JPEG, and BMP files only",a["id"],"ocr-format")
            if not path.is_file():raise WorkflowError("OCR Image input file does not exist",a["id"],"ocr-not-found")
            self.options.audit("file.read",{"activityId":a["id"],"path":self._audit_path(path)})
            if path.stat().st_size>self._max_file_bytes:raise WorkflowError("OCR Image input exceeds configured file size limit",a["id"],"file-size-limit")
            executable=shutil.which("tesseract")
            if not executable:raise WorkflowError("Tesseract OCR is unavailable; install Tesseract and ensure it is on PATH",a["id"],"ocr-runtime")
            try:
                from PIL import Image
                with Image.open(path) as image:
                    if image.format not in {"PNG","JPEG","BMP"}:raise WorkflowError("OCR Image supports PNG, JPEG, and BMP image data only",a["id"],"ocr-format")
                    width,height=image.size
                    if width<1 or height<1 or width*height>30_000_000:raise WorkflowError("OCR Image dimensions must not exceed 30 megapixels",a["id"],"ocr-pixel-limit")
                    image.verify()
            except WorkflowError:raise
            except ImportError as exc:raise WorkflowError("Pillow is unavailable; install the optional Open RPA Studio OCR dependency",a["id"],"ocr-runtime") from exc
            except Exception as exc:raise WorkflowError(f"Cannot validate OCR Image input: {type(exc).__name__}",a["id"],"ocr-image-invalid") from exc
            language=a.get("language","eng");psm=a.get("pageSegmentationMode",3);timeout=float(a.get("timeoutSeconds",60))
            self.options.audit("image.ocr.attempt",{"activityId":a["id"],"path":self._audit_path(path),"language":language,"pageSegmentationMode":psm})
            secret_markers=("API_KEY","APIKEY","TOKEN","SECRET","PASSWORD","CREDENTIAL","AUTHORIZATION")
            environment={key:value for key,value in os.environ.items() if not any(marker in key.upper() for marker in secret_markers)}
            environment["OMP_THREAD_LIMIT"]="1"
            flags=getattr(subprocess,"CREATE_NO_WINDOW",0)
            deadline=time.monotonic()+timeout
            try:
                with tempfile.TemporaryFile() as output:
                    process=subprocess.Popen([executable,str(path),"stdout","-l",language,"--psm",str(psm)],cwd=str(self.root),stdin=subprocess.DEVNULL,stdout=output,stderr=subprocess.DEVNULL,env=environment,creationflags=flags)
                    try:
                        while process.poll() is None:
                            self._check_cancel()
                            if time.monotonic()>=deadline:
                                process.terminate()
                                try:process.wait(timeout=2)
                                except subprocess.TimeoutExpired:process.kill();process.wait()
                                raise WorkflowError("OCR Image exceeded its configured timeout",a["id"],"ocr-timeout")
                            output.seek(0,os.SEEK_END)
                            if output.tell()>self._max_file_bytes:
                                process.terminate()
                                try:process.wait(timeout=2)
                                except subprocess.TimeoutExpired:process.kill();process.wait()
                                raise WorkflowError("OCR Image output exceeds configured data limit",a["id"],"file-size-limit")
                            time.sleep(min(0.05,max(0.005,deadline-time.monotonic())))
                    except RunCancelled:
                        process.terminate()
                        try:process.wait(timeout=2)
                        except subprocess.TimeoutExpired:process.kill();process.wait()
                        raise
                    exit_code=process.wait()
                    if exit_code!=0:raise WorkflowError("Tesseract failed to recognize the image; check the image and installed language data",a["id"],"ocr-failed")
                    output.seek(0);raw=output.read(self._max_file_bytes+1)
                    if len(raw)>self._max_file_bytes:raise WorkflowError("OCR Image output exceeds configured data limit",a["id"],"file-size-limit")
                    try:recognized=raw.decode("utf-8")
                    except UnicodeDecodeError as exc:raise WorkflowError("Tesseract returned invalid UTF-8 text",a["id"],"ocr-output") from exc
                v[a.get("result","recognizedText")]=recognized
                self.options.audit("image.ocr.completed",{"activityId":a["id"],"path":self._audit_path(path),"outputBytes":len(raw)})
                self._emit({"event":"image-ocr-completed","activityId":a["id"],"outputBytes":len(raw)},logs)
            except WorkflowError:raise
            except OSError as exc:raise WorkflowError(f"Cannot start Tesseract OCR: {type(exc).__name__}",a["id"],"ocr-runtime") from exc
        elif kind == "sqlite_query":
            path=self._path(a["path"],v)
            if path.suffix.lower() not in {".db",".sqlite",".sqlite3"}:raise WorkflowError("SQLite Query supports .db, .sqlite, or .sqlite3 files only",a["id"],"sqlite-format")
            parameters=interpolate(a.get("parameters",[]),v)
            if not isinstance(parameters,(list,dict)):raise WorkflowError("SQLite Query parameters must resolve to a list or object",a["id"],"sqlite-parameters")
            self.options.audit("database.query.attempt",{"activityId":a["id"],"path":self._audit_path(path),"parameterCount":len(parameters)})
            if not path.is_file():raise WorkflowError("SQLite database file does not exist",a["id"],"sqlite-not-found")
            if path.stat().st_size>self._max_file_bytes:raise WorkflowError("SQLite database exceeds configured read limit",a["id"],"file-size-limit")
            query=interpolate(a["query"],v)
            if not isinstance(query,str) or not query.strip() or len(query)>100_000:raise WorkflowError("SQLite Query needs a SQL statement no longer than 100,000 characters",a["id"],"sqlite-query")
            bound_values=parameters.values() if isinstance(parameters,dict) else parameters
            if len(parameters)>1000 or any(value is not None and not isinstance(value,(str,int,float,bool,bytes)) for value in bound_values):
                raise WorkflowError("SQLite Query accepts at most 1,000 scalar parameters",a["id"],"sqlite-parameters")
            timeout=float(a.get("timeoutSeconds",30));deadline=time.monotonic()+timeout;max_rows=int(a.get("maxRows",10_000))
            connection=None
            try:
                connection=sqlite3.connect(path.as_uri()+"?mode=ro",uri=True,timeout=min(timeout,0.5),isolation_level=None)
                connection.enable_load_extension(False)
                connection.execute("PRAGMA query_only=ON")
                connection.execute("PRAGMA trusted_schema=OFF")
                connection.execute(f"PRAGMA busy_timeout={min(int(timeout*1000),500)}")
                allowed={sqlite3.SQLITE_SELECT,sqlite3.SQLITE_READ,sqlite3.SQLITE_FUNCTION}
                if hasattr(sqlite3,"SQLITE_RECURSIVE"):allowed.add(sqlite3.SQLITE_RECURSIVE)
                def authorize(action:int,arg1:str|None,arg2:str|None,_database:str|None,_trigger:str|None)->int:
                    if action==sqlite3.SQLITE_READ and _database!="main":return sqlite3.SQLITE_DENY
                    return sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY
                connection.set_authorizer(authorize)
                def progress()->int:
                    return int(self.options.cancel() or time.monotonic()>deadline)
                connection.set_progress_handler(progress,1000)
                cursor=connection.execute(query,parameters)
                columns=[str(item[0]) for item in (cursor.description or [])]
                names=[];seen_names={}
                for column in columns:
                    seen_names[column]=seen_names.get(column,0)+1
                    names.append(column if seen_names[column]==1 else f"{column}_{seen_names[column]}")
                rows=[];result_bytes=0
                while True:
                    self._check_cancel()
                    if time.monotonic()>deadline:raise WorkflowError("SQLite Query exceeded its time limit",a["id"],"sqlite-timeout")
                    batch=cursor.fetchmany(min(128,max_rows-len(rows)+1))
                    if not batch:break
                    for raw in batch:
                        if len(rows)>=max_rows:raise WorkflowError(f"SQLite Query exceeded its {max_rows}-row result limit",a["id"],"sqlite-row-limit")
                        row={name:( {"$binaryBase64":base64.b64encode(value).decode("ascii")} if isinstance(value,bytes) else value) for name,value in zip(names,raw)}
                        try:encoded=json.dumps(row,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode("utf-8")
                        except (TypeError,ValueError) as exc:raise WorkflowError("SQLite Query returned a value that cannot be represented as JSON",a["id"],"sqlite-value") from exc
                        result_bytes+=len(encoded)
                        if result_bytes>self._max_file_bytes:raise WorkflowError("SQLite Query result exceeds configured data limit",a["id"],"file-size-limit")
                        rows.append(row)
                v[a.get("result","rows")]=rows
                self.options.audit("database.query.completed",{"activityId":a["id"],"path":self._audit_path(path),"rowCount":len(rows),"parameterCount":len(parameters)})
            except WorkflowError:raise
            except sqlite3.Error as exc:
                if self.options.cancel():raise RunCancelled() from exc
                if time.monotonic()>deadline or "interrupted" in str(exc).lower():raise WorkflowError("SQLite Query exceeded its time limit",a["id"],"sqlite-timeout") from exc
                message="SQLite Query was blocked by the read-only policy" if "not authorized" in str(exc).lower() else f"SQLite Query failed: {exc}"
                code="sqlite-readonly-policy" if "not authorized" in str(exc).lower() else "sqlite-query-error"
                raise WorkflowError(message,a["id"],code) from exc
            finally:
                if connection is not None:connection.close()
        elif kind == "xlsx_write":
            if Path(str(interpolate(a["path"],v))).suffix.lower()!='.xlsx':raise WorkflowError("Spreadsheet Write supports .xlsx output only",a["id"],"spreadsheet-format")
            rows=interpolate(a.get("rows","rows"),v)
            if isinstance(rows,str):rows=evaluate(rows,v)
            if not isinstance(rows,list) or any(not isinstance(row,dict) for row in rows):raise WorkflowError("Spreadsheet Write rows must be a list of objects",a["id"],"spreadsheet-input")
            if len(rows)>100_000:raise WorkflowError("Spreadsheet output exceeds 100,000 rows",a["id"],"spreadsheet-row-limit")
            fields=a.get("fields") or list(dict.fromkeys(key for row in rows for key in row))
            if not fields or len(fields)>2_000 or any(not isinstance(name,str) or not name.strip() for name in fields) or len(fields)!=len(set(fields)):
                raise WorkflowError("Spreadsheet Write needs 1–2,000 unique text column names",a["id"],"spreadsheet-fields")
            if any(not isinstance(key,str) for row in rows for key in row) or any(key not in fields for row in rows for key in row):raise WorkflowError("Spreadsheet row contains a column not listed in fields",a["id"],"spreadsheet-fields")
            if sum(len(str(value).encode("utf-8")) for row in rows for value in row.values())>self._max_file_bytes:raise WorkflowError("Spreadsheet cell data exceeds configured write limit",a["id"],"file-size-limit")
            if len(rows)*len(fields)>2_000_000:raise WorkflowError("Spreadsheet exceeds two million cells",a["id"],"spreadsheet-cell-limit")
            from openpyxl import Workbook
            workbook=Workbook();sheet=workbook.active;sheet.title=str(a.get("sheetName") or "Sheet1");sheet.append(fields)
            for row in rows:
                self._check_cancel()
                values=[]
                for name in fields:
                    value=row.get(name)
                    if isinstance(value,(dict,list,tuple,set)):raise WorkflowError("Spreadsheet cells must be scalar values",a["id"],"spreadsheet-cell-type")
                    values.append(value)
                sheet.append(values)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value,str) and cell.value.startswith("="):cell.data_type="s"
            output=io.BytesIO();workbook.save(output);workbook.close();encoded=output.getvalue()
            if len(encoded)>self._max_file_bytes:raise WorkflowError("Spreadsheet exceeds configured write limit",a["id"],"file-size-limit")
            self._atomic_write(self._path(a["path"],v,True),encoded)
        elif kind == "json_parse":
            value = interpolate(a.get("text", ""), v)
            if not isinstance(value, str):
                value = json.dumps(value)
            parsed = json.loads(value)
            selector = a.get("select")
            if selector:
                for part in selector.split("."):
                    parsed = parsed[int(part)] if isinstance(parsed, list) else parsed[part]
            v[a.get("result", "json")]=parsed
        elif kind == "json_write":
            value = interpolate(a.get("value", ""), v)
            if isinstance(value, str):
                value = evaluate(value, v)
            rendered = json.dumps(value, ensure_ascii=False, indent=2)
            if len(rendered.encode("utf-8"))>self._max_file_bytes: raise WorkflowError("JSON output exceeds configured file limit",a["id"],"file-size-limit")
            if a.get("path"):
                self._atomic_write(self._path(a["path"], v, True),(rendered+"\n").encode("utf-8"))
            if a.get("result"):
                v[a["result"]] = rendered
        elif kind == "wait":
            seconds = float(interpolate(a.get("seconds", 0), v))
            if seconds < 0 or seconds > 300:
                raise WorkflowError("Wait duration must be between 0 and 300 seconds", a["id"], "wait-range")
            end = time.monotonic() + seconds
            while time.monotonic() < end:
                self._check_cancel()
                time.sleep(min(0.05, end - time.monotonic()))
        elif kind == "api_request":
            self._api_activity(a,v)
        elif kind == "http_request":
            if not self.options.allow_network:
                raise WorkflowError("HTTP access requires explicit per-run network permission", a["id"], "network-permission-denied")
            url = str(interpolate(a["url"], v))
            body = a.get("body")
            if body is not None:
                body = str(interpolate(body, v)).encode("utf-8")
            parsed_target=urlsplit(url)
            self.options.audit("http.request", {"activityId":a["id"],"method":str(a.get("method","GET")).upper(),"host":parsed_target.hostname,"scheme":parsed_target.scheme})
            try:
                v[a.get("result", "response")] = restricted_http_request(
                    url, str(a.get("method", "GET")),
                    {str(k): str(interpolate(val, v)) for k, val in a.get("headers", {}).items()},
                    body, float(a.get("timeoutSeconds", 20)), int(a.get("maxResponseBytes", 5_000_000)))
            except EgressDenied as exc:
                raise WorkflowError(str(exc), a["id"], "http-egress-denied") from exc
            except (OSError, TimeoutError, ValueError) as exc:
                code="http-timeout" if isinstance(exc,TimeoutError) else "http-error"
                raise WorkflowError(f"HTTP request failed: {exc}", a["id"], code) from exc
        elif kind == "http_paginate":
            if not self.options.allow_network:
                raise WorkflowError("API access requires explicit per-run network permission", a["id"], "network-permission-denied")
            try:
                base_url = str(interpolate(a["url"], v))
                parsed = urlsplit(base_url)
                if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username is not None or parsed.password is not None or parsed.fragment:
                    raise ValueError("Paginated API URL must be an absolute HTTP(S) URL without embedded credentials or a fragment")
                page_param = str(a.get("pageParam", "page"))
                size_param = str(a.get("pageSizeParam", "limit"))
                if page_param and page_param == size_param:
                    raise ValueError("Page number and page size parameters must be different")
                page_size = int(a.get("pageSize", 100))
                max_pages = int(a.get("maxPages", 10))
                max_items = int(a.get("maxItems", 10_000))
                max_total_bytes = int(a.get("maxTotalBytes", 5_000_000))
                timeout = float(a.get("timeoutSeconds", 20))
                items_path = str(a.get("itemsPath", ""))
                headers = {str(k): str(interpolate(val, v)) for k, val in a.get("headers", {}).items()}
                items: list[Any] = []
                total_bytes = 0
                pages = 0
                truncated = False
                result_name = a.get("result", "items")
                metadata_name = a.get("metadataResult", "pagination")
                for page_number in range(1, max_pages + 1):
                    self._check_cancel()
                    query = [(key, value) for key, value in parse_qsl(parsed.query, keep_blank_values=True) if key not in {page_param, size_param}]
                    if page_param:
                        query.append((page_param, str(page_number)))
                    if size_param:
                        query.append((size_param, str(page_size)))
                    page_url = urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urlencode(query), ""))
                    self.options.audit("http.request", {"activityId": a["id"], "method": "GET", "host": parsed.hostname, "scheme": parsed.scheme, "page": page_number})
                    response = restricted_http_request(page_url, "GET", headers, None, timeout, min(max_total_bytes, 5_000_000))
                    status = response.get("status")
                    if not isinstance(status, int) or not 200 <= status < 300:
                        raise WorkflowError(f"Paginated API returned HTTP status {status}", a["id"], "api-http-status")
                    body_text = response.get("body", "")
                    if not isinstance(body_text, str):
                        raise WorkflowError("Paginated API response body was not text", a["id"], "api-response")
                    total_bytes += len(body_text.encode("utf-8"))
                    if total_bytes > max_total_bytes:
                        raise WorkflowError("Paginated API exceeded the configured total response limit", a["id"], "api-total-size-limit")
                    try:
                        document = json.loads(body_text)
                        selected: Any = document
                        if items_path:
                            for part in items_path.split("."):
                                selected = selected[int(part)] if isinstance(selected, list) and part.isdigit() else selected[part]
                    except (json.JSONDecodeError, KeyError, IndexError, TypeError, ValueError) as exc:
                        raise WorkflowError(f"Could not read JSON items at '{items_path}': {exc}", a["id"], "api-items-path") from exc
                    if not isinstance(selected, list):
                        raise WorkflowError("The configured JSON items path must resolve to an array", a["id"], "api-items-type")
                    pages += 1
                    remaining = max_items - len(items)
                    items.extend(selected[:remaining])
                    if len(selected) > remaining:
                        truncated = True
                        break
                    if len(selected) < page_size:
                        break
                    if page_number == max_pages:
                        truncated = True
                v[result_name] = items
                if metadata_name:
                    v[metadata_name] = {"pagesFetched": pages, "itemCount": len(items), "responseBytes": total_bytes, "truncated": truncated}
            except WorkflowError:
                raise
            except (OSError, TimeoutError, ValueError, TypeError) as exc:
                code = "http-timeout" if isinstance(exc, TimeoutError) else "api-request-error"
                raise WorkflowError(f"Paginated API request failed: {exc}", a["id"], code) from exc
        elif kind.startswith("browser_"):
            self._browser_activity(a, v, logs)
        elif kind.startswith("desktop_"):
            self._desktop_activity(a, v)
        else:
            raise WorkflowError(f"Unsupported activity: {kind}", a.get("id"), "activity-unsupported")

        if kind in {"file_read_text","csv_read","xlsx_read","pdf_read_text","sqlite_query","file_list"}:
            raw_path=a.get("path",a.get("directory",""))
            target=self._path(raw_path,v) if raw_path else self.root
            self.options.audit("file.read", {"activityId":a["id"],"path":self._audit_path(target)})
        elif kind in {"file_write_text","csv_write","xlsx_write","json_write"}:
            raw_path=a.get("path")
            if raw_path:
                target=self._path(raw_path,v)
                self.options.audit("file.write", {"activityId":a["id"],"path":self._audit_path(target)})
        elif kind in {"file_copy","file_move"}:
            for field,action in (("source","file.read"),("destination","file.write")):
                target=self._path(a[field],v)
                self.options.audit(action,{"activityId":a["id"],"path":self._audit_path(target)})
        elif kind=="file_delete":
            target=self._path(a["path"],v)
            self.options.audit("file.delete",{"activityId":a["id"],"path":self._audit_path(target)})

    def _api_activity(self, a: dict[str, Any], v: dict[str, Any]) -> None:
        if a.get("pagination") is not None:
            self._api_paginate(a,v,a["pagination"])
            return
        if not self.options.allow_network:
            raise WorkflowError("API access requires explicit per-run network permission", a["id"], "network-permission-denied")
        if self.options.api_connections_path is None:
            raise WorkflowError("API connections are unavailable in this execution environment", a["id"], "api-connections-unavailable")
        from . import api_connections
        try:
            profile=api_connections.get_profile(self.options.api_connections_path,str(a.get("connectionId","")))
        except (OSError,ValueError) as exc:
            raise WorkflowError(str(exc),a["id"],"api-connection-not-found") from exc
        raw_path=str(interpolate(a.get("path",""),v))
        parts=urlsplit(raw_path)
        if parts.scheme or parts.netloc or parts.fragment or raw_path.startswith("//") or any(part==".." for part in unquote(parts.path).replace("\\","/").split("/")):
            raise WorkflowError("API path must be relative to the selected connection and cannot change its host",a["id"],"api-path-invalid")
        if len(raw_path)>8192:
            raise WorkflowError("API path exceeds 8,192 characters",a["id"],"api-path-limit")
        base=profile["baseUrl"].rstrip("/")+"/"
        target=urljoin(base,raw_path.lstrip("/"))
        base_parts=urlsplit(base);target_parts=urlsplit(target)
        if (target_parts.scheme,target_parts.hostname,target_parts.port)!=(base_parts.scheme,base_parts.hostname,base_parts.port):
            raise WorkflowError("API request path changed the connection origin",a["id"],"api-path-invalid")
        query=a.get("query",{})
        if not isinstance(query,dict) or len(query)>100:
            raise WorkflowError("API query parameters must be an object with at most 100 entries",a["id"],"api-query-invalid")
        omit_empty_query=set(a.get("omitEmptyQuery",[]))
        query_items=parse_qsl(target_parts.query,keep_blank_values=True)
        for key,value in query.items():
            if not isinstance(key,str) or not key or len(key)>256:
                raise WorkflowError("API query parameter names must contain 1–256 characters",a["id"],"api-query-invalid")
            resolved=interpolate(value,v) if isinstance(value,str) else value
            if isinstance(resolved,(dict,list,tuple)):
                raise WorkflowError(f"API query parameter {key!r} must be scalar",a["id"],"api-query-invalid")
            if key in omit_empty_query and resolved in (None,""):
                continue
            query_items.append((key,"" if resolved is None else str(resolved)))
        url=urlunsplit((target_parts.scheme,target_parts.netloc,target_parts.path,urlencode(query_items,doseq=False),""))
        method=str(a.get("method","GET")).upper()
        if method not in {"GET","POST","PUT","PATCH","DELETE"}:
            raise WorkflowError("API Request supports GET, POST, PUT, PATCH, and DELETE",a["id"],"api-method-invalid")
        if method=="DELETE" and not self.options.allow_destructive:
            raise WorkflowError("API DELETE requires explicit destructive permission for this run",a["id"],"destructive-permission-denied")
        raw_headers=a.get("headers",{})
        if not isinstance(raw_headers,dict) or len(raw_headers)>50:
            raise WorkflowError("API request headers must be an object with at most 50 entries",a["id"],"api-headers-invalid")
        headers=dict(profile["headers"])
        omit_empty_headers=set(a.get("omitEmptyHeaders",[]))
        for name,value in raw_headers.items():
            if not isinstance(name,str) or not name or len(name)>128 or "\r" in name or "\n" in name:
                raise WorkflowError("API request contains an invalid header name",a["id"],"api-headers-invalid")
            if not isinstance(value,str) or len(value)>8192 or "\r" in value or "\n" in value:
                raise WorkflowError(f"API request header {name!r} must be single-line text no longer than 8,192 characters",a["id"],"api-headers-invalid")
            resolved=interpolate(value,v)
            if name in omit_empty_headers and resolved in (None,""):
                continue
            headers[name]="" if resolved is None else str(resolved)
        blocked_auth_headers={"authorization",profile["apiKeyHeader"].lower()}
        if any(name.lower() in blocked_auth_headers for name in raw_headers):
            raise WorkflowError("API Request authentication headers are managed by the saved connection",a["id"],"api-auth-header-managed")
        secret=None
        try:
            if profile["authType"] in {"bearer","api_key_header","oauth_client_credentials"}:
                secret=api_connections.read_secret(profile["id"])
                self._secret_values.add(secret)
                if profile["authType"]=="bearer":headers["Authorization"]="Bearer "+secret
                elif profile["authType"]=="api_key_header":headers[profile["apiKeyHeader"]]=secret
                else:
                    cached=self._api_tokens.get(profile["id"])
                    now=time.monotonic()
                    if cached and cached[1]>now+1:token=cached[0]
                    else:
                        try:token,expires=api_connections.request_oauth_token(profile,secret,restricted_http_request,min(float(a.get("timeoutSeconds",20)),20.0))
                        except EgressDenied as exc:raise WorkflowError(str(exc),a["id"],"http-egress-denied") from exc
                        except (OSError,TimeoutError,ValueError,RuntimeError) as exc:raise WorkflowError(f"OAuth token exchange failed: {exc}",a["id"],"api-oauth-token-failed") from exc
                        safety=min(30.0,expires*.1);self._api_tokens[profile["id"]]=(token,now+max(1.0,expires-safety))
                        token_host=urlsplit(profile["tokenUrl"]).hostname
                        self.options.audit("api.oauth.token",{"connectionId":profile["id"],"host":token_host})
                    self._secret_values.add(token);headers["Authorization"]="Bearer "+token
            elif profile["authType"]=="oauth_authorization_code":
                client_secret=api_connections.read_secret(profile["id"]) if profile["clientAuthMethod"]!="none" else ""
                if client_secret:self._secret_values.add(client_secret)
                try:token=api_connections.authorization_code_access_token(profile,restricted_http_request,client_secret,min(float(a.get("timeoutSeconds",20)),20.0))
                except EgressDenied as exc:raise WorkflowError(str(exc),a["id"],"http-egress-denied") from exc
                except (OSError,TimeoutError,ValueError,RuntimeError) as exc:raise WorkflowError(f"OAuth authorization failed: {exc}",a["id"],"api-oauth-authorization-failed") from exc
                self._secret_values.add(token);headers["Authorization"]="Bearer "+token
                self.options.audit("api.oauth.token",{"connectionId":profile["id"],"host":urlsplit(profile["tokenUrl"]).hostname,"grant":"authorization_code"})
            body=a.get("body")
            if body is not None:body=str(interpolate(body,v)).encode("utf-8")
            retry=a.get("retryPolicy") or {}
            if not isinstance(retry,dict):raise WorkflowError("API retryPolicy must be an object",a["id"],"api-retry-policy")
            attempts_value=retry.get("maxAttempts",1);delay_value=retry.get("delaySeconds",.5);backoff_value=retry.get("backoffFactor",2)
            if not isinstance(attempts_value,int) or isinstance(attempts_value,bool) or not 1<=attempts_value<=5:
                raise WorkflowError("API retry maxAttempts must be from 1 to 5",a["id"],"api-retry-policy")
            if not isinstance(delay_value,(int,float)) or isinstance(delay_value,bool) or not math.isfinite(delay_value) or not 0<=delay_value<=10:
                raise WorkflowError("API retry delaySeconds must be from 0 to 10",a["id"],"api-retry-policy")
            if not isinstance(backoff_value,(int,float)) or isinstance(backoff_value,bool) or not math.isfinite(backoff_value) or not 1<=backoff_value<=3:
                raise WorkflowError("API retry backoffFactor must be from 1 to 3",a["id"],"api-retry-policy")
            max_attempts=attempts_value;base_delay=float(delay_value);backoff=float(backoff_value)
            if max_attempts>1 and (method!="GET" or body is not None):
                raise WorkflowError("Automatic API retries require GET without a request body",a["id"],"api-retry-safety")
            transient_statuses={408,425,429,500,502,503,504}
            response=None
            for attempt in range(1,max_attempts+1):
                wait=max(0.0,profile["minIntervalSeconds"]-(time.monotonic()-self._api_last_request_at.get(profile["id"],0.0)))
                until=time.monotonic()+wait
                while time.monotonic()<until:
                    self._check_cancel();time.sleep(min(.05,until-time.monotonic()))
                self._check_cancel()
                audit={"activityId":a["id"],"connectionId":profile["id"],"method":method,"host":target_parts.hostname}
                if max_attempts>1:audit.update({"attempt":attempt,"maxAttempts":max_attempts})
                self.options.audit("api.request",audit)
                try:
                    response=restricted_http_request(url,method,headers,body,float(a.get("timeoutSeconds",20)),int(a.get("maxResponseBytes",5_000_000)))
                    self._api_last_request_at[profile["id"]]=time.monotonic()
                except EgressDenied:
                    raise
                except (OSError,TimeoutError) as exc:
                    self._api_last_request_at[profile["id"]]=time.monotonic()
                    if attempt>=max_attempts:
                        if max_attempts>1:self.options.audit("api.retry.exhausted",{"activityId":a["id"],"attempts":attempt,"reason":"transport-error","errorType":type(exc).__name__})
                        raise
                    delay=min(base_delay*(backoff**(attempt-1)),30.0)
                    self.options.audit("api.retry",{"activityId":a["id"],"attempt":attempt,"maxAttempts":max_attempts,"reason":"transport-error","delaySeconds":delay})
                    until=time.monotonic()+delay
                    while time.monotonic()<until:
                        self._check_cancel();time.sleep(min(.05,until-time.monotonic()))
                    continue
                status=response.get("status") if isinstance(response,dict) else None
                if status not in transient_statuses:break
                if attempt>=max_attempts:
                    self.options.audit("api.retry.exhausted",{"activityId":a["id"],"attempts":attempt,"reason":"transient-http-status","status":status})
                    break
                retry_after=_retry_after_seconds(response.get("headers") if isinstance(response,dict) else None)
                delay=retry_after if retry_after is not None else min(base_delay*(backoff**(attempt-1)),30.0)
                self.options.audit("api.retry",{"activityId":a["id"],"attempt":attempt,"maxAttempts":max_attempts,"reason":"transient-http-status","status":status,"delaySeconds":delay,"delaySource":"retry-after" if retry_after is not None else "exponential-backoff"})
                until=time.monotonic()+delay
                while time.monotonic()<until:
                    self._check_cancel();time.sleep(min(.05,until-time.monotonic()))
            if response is None:raise WorkflowError("API request did not produce a response",a["id"],"api-request-empty-response")
            v[a.get("result","response")]=response
            mapping=a.get("responseMapping")
            if mapping is not None:
                status=response.get("status") if isinstance(response,dict) else None
                if not isinstance(status,int) or isinstance(status,bool) or not 200<=status<300:
                    raise WorkflowError("JSON response mapping requires a successful 2xx HTTP response",a["id"],"api-response-status")
                try:
                    body_text=response.get("body")
                    if not isinstance(body_text,str):raise ValueError
                    document=json.loads(body_text)
                except (TypeError,ValueError,json.JSONDecodeError) as exc:
                    raise WorkflowError("API response body is not valid JSON for the configured mapping",a["id"],"api-response-json") from exc
                try:
                    records=_json_pointer_get(document,mapping.get("itemsPointer",""))
                except (KeyError,IndexError,ValueError,TypeError) as exc:
                    raise WorkflowError("API response items pointer did not resolve",a["id"],"api-response-pointer") from exc
                if not isinstance(records,list):
                    raise WorkflowError("API response items pointer must select a JSON array",a["id"],"api-response-array")
                v[mapping["result"]]=self._map_api_records(records,mapping,a)
        except WorkflowError:
            raise
        except EgressDenied as exc:
            raise WorkflowError(str(exc),a["id"],"http-egress-denied") from exc
        except (OSError,TimeoutError,ValueError,RuntimeError) as exc:
            code="http-timeout" if isinstance(exc,TimeoutError) else "http-error"
            raise WorkflowError(f"API request failed: {exc}",a["id"],code) from exc

    def _api_paginate(self, a: dict[str, Any], v: dict[str, Any], config: dict[str, Any]) -> None:
        """Fetch bounded cursor or next-link pages through the normal connection path."""
        if not self.options.allow_network:
            raise WorkflowError("API access requires explicit per-run network permission",a["id"],"network-permission-denied")
        if self.options.api_connections_path is None:
            raise WorkflowError("API connections are unavailable in this execution environment",a["id"],"api-connections-unavailable")
        from . import api_connections
        try: profile=api_connections.get_profile(self.options.api_connections_path,str(a.get("connectionId","")))
        except (OSError,ValueError) as exc: raise WorkflowError(str(exc),a["id"],"api-connection-not-found") from exc
        strategy=config.get("strategy")
        max_pages=int(config.get("maxPages",10));max_items=int(config.get("maxItems",10000));max_bytes=int(config.get("maxTotalBytes",5000000))
        base=urlsplit(profile["baseUrl"].rstrip("/")+"/")
        base_path=base.path.rstrip("/")
        base_prefix=(base_path+"/") if base_path else "/"
        page_path=str(a.get("path","/"));page_query=dict(a.get("query",{}));cursor_param=str(config.get("cursorParam",""))
        omit_empty_query=set(a.get("omitEmptyQuery",[]))
        current_url=urljoin(profile["baseUrl"].rstrip("/")+"/",page_path.lstrip("/"))
        current_parts=urlsplit(current_url);initial_query=parse_qsl(current_parts.query,keep_blank_values=True)
        for key,value in page_query.items():
            resolved=interpolate(value,v) if isinstance(value,str) else value
            if key in omit_empty_query and resolved in (None,""):continue
            if not isinstance(resolved,(dict,list,tuple)):initial_query.append((key,"" if resolved is None else str(resolved)))
        current_url=urlunsplit((current_parts.scheme,current_parts.netloc,current_parts.path,urlencode(initial_query,doseq=False),""))
        initial_cursor=page_query.get(cursor_param) if strategy=="cursor" else None
        if isinstance(initial_cursor,str):initial_cursor=interpolate(initial_cursor,v)
        seen_cursors={str(initial_cursor)} if isinstance(initial_cursor,(str,int)) and not isinstance(initial_cursor,bool) and str(initial_cursor) else set()
        all_items=[];mapped_items=[];mapped_output_bytes=0;mapping=a.get("responseMapping");total_bytes=0;last_response=None;seen_links={current_url};more=False
        page_result="__openrpa_internal_page_response__"
        for page_number in range(max_pages):
            self._check_cancel()
            page=dict(a);page.pop("pagination",None);page.pop("responseMapping",None)
            page["method"]="GET";page.pop("body",None);page["result"]=page_result
            page["path"]=page_path;page["query"]=dict(page_query)
            page_vars=dict(v)
            self._api_activity(page,page_vars)
            response=page_vars.get(page_result)
            status=response.get("status") if isinstance(response,dict) else None
            if not isinstance(status,int) or isinstance(status,bool) or not 200<=status<300:
                raise WorkflowError("API pagination requires successful 2xx responses",a["id"],"api-pagination-status")
            body_text=response.get("body")
            if not isinstance(body_text,str):raise WorkflowError("API page response body is missing",a["id"],"api-pagination-json")
            page_bytes=len(body_text.encode("utf-8"));total_bytes+=page_bytes
            if total_bytes>max_bytes:raise WorkflowError("API pagination exceeded maxTotalBytes",a["id"],"api-pagination-byte-limit")
            try:document=json.loads(body_text)
            except (ValueError,TypeError) as exc:raise WorkflowError("API page response is not valid JSON",a["id"],"api-pagination-json") from exc
            try:items=_json_pointer_get(document,config.get("itemsPointer",""))
            except (KeyError,IndexError,ValueError,TypeError) as exc:raise WorkflowError("API pagination itemsPointer did not resolve",a["id"],"api-pagination-items-pointer") from exc
            if not isinstance(items,list):raise WorkflowError("API pagination itemsPointer must select a JSON array",a["id"],"api-pagination-items-array")
            if len(all_items)+len(items)>max_items:raise WorkflowError("API pagination exceeded maxItems",a["id"],"api-pagination-item-limit")
            if mapping is not None:
                map_limit=int(mapping.get("maxItems",10_000))
                if len(mapped_items)+len(items)>map_limit:
                    raise WorkflowError("API pagination exceeded the configured mapped-row limit",a["id"],"api-response-limit")
                page_mapped=self._map_api_records(items,mapping,a)
                mapped_output_bytes+=sum(len(json.dumps(row,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode("utf-8")) for row in page_mapped)
                if mapped_output_bytes>int(mapping.get("maxOutputBytes",5_000_000)):
                    raise WorkflowError("API pagination exceeded the configured mapped-output byte limit",a["id"],"api-response-output-limit")
                mapped_items.extend(page_mapped)
            all_items.extend(items);last_response=response;more=False
            if strategy=="cursor":
                try:next_cursor=_json_pointer_get(document,config.get("cursorPointer",""))
                except (KeyError,IndexError,ValueError,TypeError) as exc:raise WorkflowError("API pagination cursorPointer did not resolve",a["id"],"api-pagination-cursor-pointer") from exc
                if next_cursor is None or next_cursor=="":break
                if isinstance(next_cursor,bool) or not isinstance(next_cursor,(str,int)):
                    raise WorkflowError("API pagination cursor must be a string, integer, or null",a["id"],"api-pagination-cursor-type")
                token=str(next_cursor)
                if not token or len(token)>2048:raise WorkflowError("API pagination cursor is empty or exceeds 2,048 characters",a["id"],"api-pagination-cursor-limit")
                if token in seen_cursors:raise WorkflowError("API pagination detected a repeated cursor",a["id"],"api-pagination-cursor-loop")
                seen_cursors.add(token);page_query=dict(a.get("query",{}));page_query[cursor_param]=token;more=True
            elif strategy=="next_link":
                try:next_link=_json_pointer_get(document,config.get("nextLinkPointer",""))
                except (KeyError,IndexError,ValueError,TypeError) as exc:raise WorkflowError("API pagination nextLinkPointer did not resolve",a["id"],"api-pagination-next-link-pointer") from exc
                if next_link is None or next_link=="":break
                if not isinstance(next_link,str) or len(next_link)>8192:raise WorkflowError("API pagination next link must be text no longer than 8,192 characters",a["id"],"api-pagination-next-link")
                target=urljoin(current_url,next_link);parts=urlsplit(target)
                if parts.username is not None or parts.password is not None or parts.fragment or (parts.scheme,parts.hostname,parts.port)!=(base.scheme,base.hostname,base.port):
                    raise WorkflowError("API next link must remain on the saved connection origin",a["id"],"api-pagination-next-link-origin")
                decoded_path=unquote(parts.path).replace("\\","/")
                if any(segment==".." for segment in decoded_path.split("/")):
                    raise WorkflowError("API next link cannot traverse outside the saved connection path",a["id"],"api-pagination-next-link-path")
                if base_prefix!="/":
                    if parts.path==base_path:relative_path=""
                    elif parts.path.startswith(base_prefix):relative_path=parts.path[len(base_prefix):]
                    else:raise WorkflowError("API next link must remain within the saved connection base path",a["id"],"api-pagination-next-link-path")
                else:relative_path=parts.path.lstrip("/")
                normalized=urlunsplit((parts.scheme,parts.netloc,parts.path,parts.query,""))
                if normalized in seen_links:raise WorkflowError("API pagination detected a repeated next link",a["id"],"api-pagination-next-link-loop")
                seen_links.add(normalized);page_path=relative_path+("?"+parts.query if parts.query else "");page_query={};current_url=normalized;more=True
            else:
                raise WorkflowError("Unsupported API pagination strategy",a["id"],"api-pagination-strategy")
            if not more:break
            if page_number+1>=max_pages:raise WorkflowError("API pagination reached maxPages while another page remained",a["id"],"api-pagination-page-limit")
        if last_response is not None:
            v[a.get("result","response")]=last_response
            v[config.get("result","items")]=all_items
            if mapping is not None:v[mapping["result"]]=mapped_items
            metadata_name=config.get("metadataResult","pagination")
            if metadata_name:
                metadata={"strategy":strategy,"pages":page_number+1,"itemCount":len(all_items),"responseBytes":total_bytes}
                if mapping is not None:
                    metadata["mappedItemCount"]=len(mapped_items)
                    metadata["mappedOutputBytes"]=mapped_output_bytes
                v[metadata_name]=metadata

    def _map_api_records(self, records: list[Any], mapping: dict[str, Any], activity: dict[str, Any]) -> list[dict[str, Any]]:
        """Project JSON values from one bounded API record array."""
        if len(records)>int(mapping.get("maxItems",10_000)):
            raise WorkflowError("API response array exceeds the configured mapped-row limit",activity["id"],"api-response-limit")
        mapped=[];output_bytes=0;max_output_bytes=int(mapping.get("maxOutputBytes",5_000_000))
        for record in records:
            self._check_cancel()
            if not isinstance(record,dict):
                raise WorkflowError("API response array items must be JSON objects",activity["id"],"api-response-record")
            row={}
            for field_mapping in mapping.get("fields",[]):
                try:
                    if isinstance(field_mapping.get("expression"),str):
                        value=evaluate(field_mapping["expression"],{"item":record})
                    else:
                        value=_json_pointer_get(record,field_mapping["path"])
                except Exception as exc:
                    code="api-response-field-expression" if isinstance(field_mapping.get("expression"),str) else "api-response-field-pointer"
                    message="An API response field expression failed for a record" if code.endswith("expression") else "An API response field pointer did not resolve"
                    raise WorkflowError(message,activity["id"],code) from exc
                row[field_mapping["name"]]=value
            try:row_bytes=len(json.dumps(row,ensure_ascii=False,separators=(",",":"),allow_nan=False).encode("utf-8"))
            except (TypeError,ValueError,OverflowError) as exc:
                raise WorkflowError("Mapped API response record is not valid JSON",activity["id"],"api-response-field-type") from exc
            output_bytes+=row_bytes
            if output_bytes>max_output_bytes:
                raise WorkflowError("API response mapping exceeded the configured mapped-output byte limit",activity["id"],"api-response-output-limit")
            mapped.append(row)
        return mapped

    def _browser_activity(self, a: dict[str, Any], v: dict[str, Any], logs: list[dict[str, Any]] | None = None) -> None:
        kind = a["type"]
        if kind == "browser_launch":
            if self._browser_page is not None:
                raise WorkflowError("Browser is already open", a["id"], "browser-state")
            browser_type=a.get("browserType","chromium")
            channel=a.get("channel","chromium")
            profile_name=a.get("profileName","")
            if profile_name and not self.options.allow_desktop:
                raise WorkflowError("Reusing a persistent signed-in browser profile requires per-run desktop permission; review the profile and enable Allow desktop control.",a["id"],"desktop-permission-denied")
            try:
                from playwright.sync_api import sync_playwright
            except ImportError as exc:
                raise WorkflowError("Browser automation requires optional Playwright and its browser runtime. See docs/AUTOMATION_ADAPTERS.md", a["id"], "runtime-unavailable") from exc
            self._browser_runtime = sync_playwright().start()
            launcher=getattr(self._browser_runtime,browser_type,None)
            if launcher is None:
                raise WorkflowError(f"Playwright browser engine '{browser_type}' is unavailable",a["id"],"browser-engine-unavailable")
            options={"headless":bool(a.get("headless",False))}
            executable_path=a.get("executablePath","")
            if browser_type=="chromium":
                if channel=="custom":
                    try:
                        custom_path=Path(str(executable_path)).expanduser().resolve(strict=True)
                    except (OSError,ValueError) as exc:
                        raise WorkflowError("Custom Chromium executable does not exist",a["id"],"browser-executable-not-found") from exc
                    if not custom_path.is_file():
                        raise WorkflowError("Custom Chromium executable path is not a file",a["id"],"browser-executable-not-found")
                    options["executable_path"]=str(custom_path)
                elif channel!="chromium":options["channel"]=channel
            try:
                if profile_name:
                    import hashlib
                    from .app_paths import app_data_dir
                    project_key=hashlib.sha256(str(self.root).encode("utf-8")).hexdigest()[:24]
                    data_root=app_data_dir().resolve()
                    profiles_base=data_root/"browser-profiles"
                    profiles_base.mkdir(parents=True,exist_ok=True)
                    profile_root=profiles_base.resolve()/project_key
                    try:profile_root.relative_to(data_root)
                    except ValueError as exc:raise WorkflowError("App-managed browser profile folder escapes local app data",a["id"],"browser-profile-path") from exc
                    profile_path=profile_root/profile_name.casefold()
                    profile_root.mkdir(parents=True,exist_ok=True)
                    profile_path.mkdir(parents=True,exist_ok=True)
                    resolved_profile=profile_path.resolve()
                    try:resolved_profile.relative_to(profile_root.resolve())
                    except ValueError as exc:raise WorkflowError("Persistent browser profile path escaped the app-managed profile folder",a["id"],"browser-profile-path") from exc
                    self._browser_persistent_profile=True
                    self._browser_context=launcher.launch_persistent_context(str(resolved_profile),**options)
                    self._browser=self._browser_context.browser
                else:
                    self._browser = launcher.launch(**options)
            except Exception as exc:
                self._browser_persistent_profile=False
                self._browser=self._browser_context=None
                if isinstance(exc,WorkflowError):raise
                if browser_type=="chromium" and channel=="chromium" and "Executable doesn't exist" in str(exc):
                    raise WorkflowError("Playwright Chromium is not installed in this environment. Install it with 'playwright install chromium', or set Browser Launch channel to 'chrome' or 'msedge' to use an installed browser.",a["id"],"browser-runtime-missing") from exc
                if browser_type in {"firefox","webkit"} and "Executable doesn't exist" in str(exc):
                    raise WorkflowError(f"Playwright {browser_type.title()} is not installed in this environment. Install it with 'playwright install {browser_type}'.",a["id"],"browser-runtime-missing") from exc
                raise WorkflowError(f"Could not launch Playwright {browser_type} browser (channel '{channel}'): {exc}",a["id"],"browser-launch-failed") from exc
            if not profile_name:self._browser_context = self._browser.new_context()
            pages=list(getattr(self._browser_context,"pages",[]))
            self._browser_page=pages[0] if pages else self._browser_context.new_page()
            self._browser_tabs = {}
            self._browser_unclaimed_popups = []
            self._browser_popup_counter = 0
            self._register_browser_page(self._browser_page, "main")
            self._browser_context.on("page", self._on_browser_popup)
            if profile_name:
                profile_id=hashlib.sha256(f"{project_key}:{profile_name.casefold()}".encode("utf-8")).hexdigest()[:16]
                self._emit({"event":"browser-profile-opened","activityId":a["id"],"profileId":profile_id},logs)
                self.options.audit("browser.profile_opened",{"activityId":a["id"],"profileId":profile_id,"browserType":browser_type})
            return
        if kind == "browser_close":
            self._close_browser()
            return
        if kind in {"browser_new_tab","browser_capture_popup","browser_switch_tab","browser_close_tab"} and self._browser_context is None:
            raise WorkflowError("Launch Browser must run before tab operations", a["id"], "browser-state")
        if kind == "browser_new_tab":
            name = str(a.get("name", ""))
            if name in self._browser_tabs:
                raise WorkflowError(f"Browser tab name '{name}' is already in use", a["id"], "browser-tab-name")
            page = self._browser_context.new_page()
            self._register_browser_page(page, name)
            self._browser_page = page
            url = str(interpolate(a.get("url", ""), v)).strip()
            if url:
                self._navigate_browser_page(page, url, a, v, timeout_seconds=float(a.get("timeoutSeconds", 15)))
            return
        if kind == "browser_capture_popup":
            name = str(a.get("name", ""))
            if name in self._browser_tabs:
                raise WorkflowError(f"Browser tab name '{name}' is already in use", a["id"], "browser-tab-name")
            deadline = time.monotonic() + float(a.get("timeoutSeconds", 15))
            while time.monotonic() < deadline:
                self._check_cancel()
                for candidate in self._browser_context.pages:
                    if not candidate.is_closed() and not any(item is candidate for item in self._browser_tabs.values()):
                        self._on_browser_popup(candidate)
                if self._browser_unclaimed_popups:
                    break
                time.sleep(min(.05, max(0, deadline - time.monotonic())))
            if not self._browser_unclaimed_popups:
                raise WorkflowError("No new browser popup opened before the timeout", a["id"], "browser-popup-timeout")
            page = self._browser_unclaimed_popups.pop(0)
            self._register_browser_page(page, name)
            self._browser_page = page
            return
        if kind == "browser_switch_tab":
            name = str(a.get("name", ""))
            page = self._browser_tabs.get(name)
            if page is None or page.is_closed():
                raise WorkflowError(f"Browser tab '{name}' is not open", a["id"], "browser-tab-not-found")
            self._browser_page = page
            return
        if kind == "browser_close_tab":
            name = str(a.get("name", ""))
            page = self._browser_tabs.get(name)
            open_tabs = [(tab_name, tab_page) for tab_name, tab_page in self._browser_tabs.items() if not tab_page.is_closed()]
            if page is None or page.is_closed():
                raise WorkflowError(f"Browser tab '{name}' is not open", a["id"], "browser-tab-not-found")
            if len(open_tabs) <= 1:
                raise WorkflowError("Cannot close the last browser tab with Close Browser Tab; use Close Browser to end the session", a["id"], "browser-last-tab")
            page.close()
            self._browser_tabs.pop(name, None)
            self._browser_responses.pop(page,None)
            self._browser_unclaimed_popups = [item for item in self._browser_unclaimed_popups if item is not page]
            if self._browser_page is page:
                self._browser_page = next(tab_page for tab_name, tab_page in open_tabs if tab_name != name)
            return
        if kind == "browser_wait_response":
            if self._browser_page is None:
                raise WorkflowError("Launch Browser must run before waiting for a browser response", a["id"], "browser-state")
            if not self.options.allow_network:
                raise WorkflowError("Browser response observation requires explicit per-run network permission", a["id"], "network-permission-denied")
            return self._wait_browser_response(a, v, logs)
        if self._browser_page is None:
            raise WorkflowError("Launch Browser must run before browser interaction", a["id"], "browser-state")
        timeout = min(max(float(a.get("timeoutSeconds", 15)),1),120) * 1000
        if kind == "browser_wait":
            state=a.get("state","visible")
            self._locator(a,logs,allow_fallback=state not in {"hidden","detached"}).wait_for(state=state,timeout=timeout)
            return
        if kind == "browser_assert":
            assertion=a.get("assertion","visible")
            locator=self._locator(a,logs,allow_fallback=assertion!="hidden")
            if assertion in {"visible","hidden"}:
                actual=locator.is_visible(timeout=timeout) if assertion=="visible" else locator.is_hidden(timeout=timeout)
                passed=actual
            elif assertion in {"enabled","disabled"}:
                actual=locator.is_enabled(timeout=timeout);passed=actual if assertion=="enabled" else not actual
            elif assertion=="checked":
                actual=locator.is_checked(timeout=timeout);passed=actual
            elif assertion=="text":
                actual=locator.inner_text(timeout=timeout)
                expected=str(interpolate(a.get("expectedText",""),v))
                passed=actual==expected if a.get("textMatch","exact")=="exact" else expected in actual
            else: raise WorkflowError(f"Unsupported browser assertion: {assertion}",a["id"],"browser-assertion")
            if not passed:
                detail=f"expected {assertion}={a.get('expectedText','')!r}, observed {actual!r}" if assertion=="text" else f"expected {assertion}, observed {actual}"
                raise WorkflowError(f"Browser assertion failed for selector {a.get('selector')!r}: {detail}",a["id"],"browser-assertion-failed")
            return
        if kind == "browser_navigate":
            url = str(interpolate(a["url"], v))
            self._navigate_browser_page(self._browser_page, url, a, v, timeout_seconds=timeout/1000)
        elif kind == "browser_click": self._locator(a,logs).click(timeout=timeout)
        elif kind == "browser_hover": self._locator(a,logs).hover(timeout=timeout)
        elif kind == "browser_press": self._locator(a,logs).press(str(a.get("key","Enter")),timeout=timeout)
        elif kind == "browser_upload":
            raw_paths = a.get("paths", [])
            paths = [self._path(raw_path, v) for raw_path in raw_paths]
            sizes = []
            for path in paths:
                if not path.is_file():
                    raise WorkflowError("Browser Upload path must identify an existing file", a["id"], "browser-upload-file")
                size = path.stat().st_size
                if size > self._max_file_bytes:
                    raise WorkflowError("Browser Upload file exceeds the configured file size limit", a["id"], "file-size-limit")
                sizes.append(size)
            locator = self._locator(a, logs)
            if locator.count() != 1:
                raise WorkflowError("Browser Upload selector must match exactly one file input", a["id"], "browser-upload-selector")
            is_file_input = locator.evaluate("element => element instanceof HTMLInputElement && element.type === 'file'")
            if not is_file_input:
                raise WorkflowError("Browser Upload selector must identify an HTML file input", a["id"], "browser-upload-selector")
            locator.set_input_files([str(path) for path in paths], timeout=timeout)
            self.options.audit("browser.upload", {"activityId": a["id"], "fileCount": len(paths), "totalBytes": sum(sizes)})
            if logs is not None:
                self._emit({"event": "browser-upload-complete", "activityId": a["id"], "fileCount": len(paths), "totalBytes": sum(sizes)}, logs)
        elif kind == "browser_download":
            destination = self._path(a["path"], v, True)
            overwrite = bool(a.get("overwrite", False))
            if destination.exists() and not overwrite:
                raise WorkflowError("Browser Download destination already exists; enable overwrite to replace it", a["id"], "browser-download-exists")
            try:
                with self._browser_page.expect_download(timeout=timeout) as pending_download:
                    self._locator(a, logs).click(timeout=timeout)
                download = pending_download.value
                source = Path(download.path())
            except Exception as exc:
                error_name = type(exc).__name__
                if "Timeout" in error_name or "Timeout" in str(exc):
                    raise WorkflowError("No browser download began before the configured timeout", a["id"], "browser-download-timeout") from exc
                raise WorkflowError(f"Browser download failed before it could be saved ({error_name})", a["id"], "browser-download-failed") from exc
            size = source.stat().st_size
            if size > self._max_file_bytes:
                try: download.cancel()
                except Exception: pass
                raise WorkflowError("Browser Download exceeds the configured file size limit", a["id"], "file-size-limit")
            descriptor, temporary = tempfile.mkstemp(prefix=".openrpa-download-", dir=destination.parent)
            copied = 0
            try:
                with os.fdopen(descriptor, "wb") as output, source.open("rb") as input_stream:
                    while True:
                        self._check_cancel()
                        chunk = input_stream.read(min(1024 * 1024, self._max_file_bytes - copied + 1))
                        if not chunk: break
                        copied += len(chunk)
                        if copied > self._max_file_bytes:
                            raise WorkflowError("Browser Download grew beyond the configured file size limit", a["id"], "file-size-limit")
                        output.write(chunk)
                    output.flush(); os.fsync(output.fileno())
                if destination.exists() and not overwrite:
                    raise WorkflowError("Browser Download destination appeared during the download", a["id"], "browser-download-exists")
                os.replace(temporary, destination)
            finally:
                if os.path.exists(temporary): os.unlink(temporary)
            result_name = a.get("result", "download")
            v[result_name] = {"path": self._audit_path(destination), "fileName": str(download.suggested_filename)[:255], "bytes": copied}
            self.options.audit("browser.download", {"activityId": a["id"], "path": self._audit_path(destination), "bytes": copied})
            if logs is not None:
                self._emit({"event": "browser-download-complete", "activityId": a["id"], "bytes": copied}, logs)
        elif kind == "browser_fill": self._locator(a,logs).fill(str(interpolate(a.get("value", ""), v)), timeout=timeout)
        elif kind == "browser_select": self._locator(a,logs).select_option(value=str(interpolate(a.get("value", ""),v)),timeout=timeout)
        elif kind == "browser_select_many":
            raw_values=interpolate(a.get("values",[]),v)
            if isinstance(raw_values,list):values=[interpolate(value,v) for value in raw_values]
            else:values=raw_values
            if not isinstance(values,list) or not 1<=len(values)<=100 or any(not isinstance(value,str) or not value or len(value)>1000 for value in values):
                raise WorkflowError("Select Multiple needs 1 to 100 non-empty option values, each no longer than 1,000 characters",a["id"],"browser-option-values")
            if len(set(values))!=len(values):raise WorkflowError("Select Multiple option values must be unique",a["id"],"browser-option-values-duplicate")
            locator=self._locator(a,logs)
            try:
                from .browser_selection import select_multiple_options
                selected=select_multiple_options(locator,values,int(timeout*1000),self._check_cancel)
            except RunCancelled:raise
            except WorkflowError:raise
            except Exception as exc:raise WorkflowError(f"Could not select multiple browser options ({type(exc).__name__}): {str(exc)[:500]}",a["id"],"browser-select-many") from exc
            v[a.get("result","selectedValues")]=selected
            if logs is not None:self._emit({"event":"browser-select-many-complete","activityId":a["id"],"selectedCount":len(selected)},logs)
        elif kind == "browser_check": self._locator(a,logs).check(timeout=timeout)
        elif kind == "browser_uncheck": self._locator(a,logs).uncheck(timeout=timeout)
        elif kind == "browser_text": v[a.get("result", "text")] = self._locator(a,logs).inner_text(timeout=timeout)
        elif kind == "browser_accessibility_snapshot":
            limit=int(a.get("maxOutputChars",50_000))
            try:
                snapshot=self._browser_page.locator("body").aria_snapshot(timeout=timeout)
            except Exception as exc:
                raise WorkflowError(f"Could not read the browser accessibility tree ({type(exc).__name__}): {str(exc)[:500]}",a["id"],"browser-accessibility-snapshot") from exc
            if len(snapshot)>limit:
                raise WorkflowError(f"Accessibility snapshot has {len(snapshot)} characters; limit is {limit}. Increase maxOutputChars or scope the page before reading it.",a["id"],"browser-accessibility-limit")
            v[a.get("result","accessibilitySnapshot")]=snapshot
            if logs is not None:self._emit({"event":"browser-accessibility-snapshot","activityId":a["id"],"characters":len(snapshot)},logs)
        elif kind == "browser_extract_records":
            append=bool(a.get("append",False));result_name=a.get("result","records")
            try:
                if a.get("selectorKind", "css") != "css":
                    raise WorkflowError("Extract Repeated Web Items requires a CSS selector for the repeated item container", a["id"], "browser-extract-selector")
                fields = a.get("fields", [])
                max_records=int(a.get("maxRecords",500));max_chars=int(a.get("maxOutputChars",1_000_000))
                limits={"maxRecords":max_records,"maxOutputChars":max_chars}
                previous=v.get(result_name,[]) if append else []
                deduplicate_by=a.get("deduplicateBy",[])
                if not isinstance(deduplicate_by,list) or len(deduplicate_by)>10 or any(not isinstance(name,str) or not name for name in deduplicate_by):
                    raise WorkflowError("Extract Repeated Web Items deduplicateBy must be a list of at most 10 field names",a["id"],"browser-extract-deduplicate")
                if deduplicate_by and not append:
                    raise WorkflowError("Extract Repeated Web Items deduplication requires append mode",a["id"],"browser-extract-deduplicate")
                mapped_names={field.get("name") for field in fields if isinstance(field,dict) and isinstance(field.get("name"),str)}
                if deduplicate_by and any(name not in mapped_names for name in deduplicate_by):
                    raise WorkflowError("Every deduplicateBy value must match a mapped field name",a["id"],"browser-extract-deduplicate")
                previous_chars=0
                if append:
                    if not isinstance(previous,list) or any(not isinstance(row,dict) or any(not isinstance(key,str) or not isinstance(value,str) for key,value in row.items()) for row in previous):
                        raise WorkflowError("Append mode needs its result variable to be a list of string-valued records",a["id"],"browser-extract-append-input")
                    previous_chars=sum(len(key.encode("utf-16-le","surrogatepass"))//2+len(value.encode("utf-16-le","surrogatepass"))//2 for row in previous for key,value in row.items())
                    total_record_limit=int(a.get("maxTotalRecords",10000))
                    total_char_limit=int(a.get("maxTotalOutputChars",5_000_000))
                    if len(previous)>total_record_limit or previous_chars>total_char_limit:
                        raise WorkflowError("Append extraction reached its aggregate record or text limit; the accumulated result was discarded",a["id"],"browser-extract-total-limit")
                    limits["maxRecords"]=max_records
                    limits["maxOutputChars"]=min(max_chars,total_char_limit)
                locator = self._browser_scope(a).locator(str(a.get("selector", "")))
                locator.first.wait_for(state="attached", timeout=timeout)
                extracted = locator.evaluate_all("""(items, config) => {
                  if (items.length > config.maxRecords) return {error: `matched ${items.length} items; limit is ${config.maxRecords}`};
                  let chars = 0;
                  const records = [];
                  for (let index = 0; index < items.length; index++) {
                    const item = items[index];
                    const record = {};
                    for (const field of config.fields) {
                      const matches = item.querySelectorAll(field.selector);
                      if (matches.length > 1) return {error: `item ${index + 1}, field '${field.name}' matched ${matches.length} elements; use a more specific selector`};
                      if (!matches.length) {
                        if (field.required) return {error: `item ${index + 1} is missing required field '${field.name}'`};
                        record[field.name] = '';
                        continue;
                      }
                      const element = matches[0];
                      let value = field.source === 'attribute' ? (element.getAttribute(field.attribute) || '') : (element.innerText || element.textContent || '');
                      value = String(value).replace(/\\s+/g, ' ').trim();
                      if (field.tokenIndex !== undefined) value = (value.match(/\\S+/g) || [])[field.tokenIndex] || '';
                      if (field.required && !value) return {error: `item ${index + 1} has an empty required field '${field.name}'`};
                      if (field.source === 'attribute' && field.attribute.toLowerCase() === 'href' && value) value = new URL(value, document.baseURI).href;
                      record[field.name] = value;
                      chars += field.name.length + value.length;
                      if (chars > config.maxOutputChars) return {error: `extracted text exceeds output limit of ${config.maxOutputChars} characters`};
                    }
                    records.push(record);
                  }
                  return {records};
                }""", {"fields": fields, **limits})
                if extracted.get("error"):
                    raise WorkflowError(f"Extract Repeated Web Items failed: {extracted['error']}", a["id"], "browser-extract-records")
                records = extracted.get("records", [])
                duplicates_skipped=0
                if append and deduplicate_by:
                    combined=[];seen_keys=set()
                    for row in previous+records:
                        key=tuple(row.get(name,"") for name in deduplicate_by)
                        if any(value for value in key):
                            if key in seen_keys:
                                duplicates_skipped+=1
                                continue
                            seen_keys.add(key)
                        combined.append(row)
                else:
                    combined=previous+records if append else records
                if append:
                    combined_chars=sum(len(key.encode("utf-16-le","surrogatepass"))//2+len(value.encode("utf-16-le","surrogatepass"))//2 for row in combined for key,value in row.items())
                    if len(combined)>int(a.get("maxTotalRecords",10000)) or combined_chars>int(a.get("maxTotalOutputChars",5_000_000)):
                        raise WorkflowError("Append extraction reached its aggregate record or text limit; the accumulated result was discarded",a["id"],"browser-extract-total-limit")
                v[result_name]=combined
                if logs is not None:
                    self._emit({"event":"browser-records-extracted","activityId":a["id"],"records":len(records),"duplicatesSkipped":duplicates_skipped,"totalRecords":len(combined),"appended":append,"fields":[field.get("name") for field in fields]},logs)
            except Exception:
                if append:v.pop(result_name,None)
                raise
        elif kind == "browser_table_read":
            if a.get("selectorKind","css")!="css":raise WorkflowError("Read Web Table requires a CSS selector for a table element",a["id"],"browser-table-selector")
            locator=self._browser_scope(a).locator(str(a.get("selector","")))
            locator.wait_for(state="visible",timeout=timeout)
            matches=locator.count()
            if matches!=1:raise WorkflowError(f"Read Web Table selector must match exactly one table; found {matches}",a["id"],"browser-table-ambiguous")
            limits={"maxRows":int(a.get("maxRows",1000)),"maxColumns":int(a.get("maxColumns",50)),"maxOutputChars":int(a.get("maxOutputChars",1_000_000)),"firstRowHeaders":bool(a.get("firstRowHeaders",True))}
            extracted=locator.evaluate("""(table, limits) => {
              if (!(table instanceof HTMLTableElement)) return {error: 'selector did not identify an HTML table'};
              const rowCount = table.rows.length;
              const headerOffset = limits.firstRowHeaders && rowCount ? 1 : 0;
              const dataRows = rowCount - headerOffset;
              if (dataRows > limits.maxRows) return {error: `table has ${dataRows} data rows; limit is ${limits.maxRows}`};
              const sourceRows = Array.from(table.rows);
              let width = 0;
              for (const row of sourceRows) width = Math.max(width, row.cells.length);
              if (width > limits.maxColumns) return {error: `table has ${width} columns; limit is ${limits.maxColumns}`};
              if (!width) return {rows: []};
              const clean = (cell) => (cell.innerText || '').replace(/\\s+/g, ' ').trim();
              let headers = headerOffset
                ? Array.from({length: width}, (_, index) => clean(sourceRows[0].cells[index] || {innerText: ''}) || `column_${index + 1}`)
                : Array.from({length: width}, (_, index) => `column_${index + 1}`);
              const used = new Map();
              headers = headers.map((header, index) => {
                const base = header.slice(0, 256) || `column_${index + 1}`;
                const count = (used.get(base) || 0) + 1; used.set(base, count);
                return count === 1 ? base : `${base}_${count}`;
              });
              let chars = 0;
              const rows = [];
              for (const row of sourceRows.slice(headerOffset)) {
                const record = {};
                for (let column = 0; column < width; column++) {
                  const value = clean(row.cells[column] || {innerText: ''});
                  chars += headers[column].length + value.length;
                  if (chars > limits.maxOutputChars) return {error: `table text exceeds output limit of ${limits.maxOutputChars} characters`};
                  record[headers[column]] = value;
                }
                rows.push(record);
              }
              return {rows};
            }""",limits)
            if extracted.get("error"):raise WorkflowError(f"Read Web Table failed: {extracted['error']}",a["id"],"browser-table-limit")
            v[a.get("result","rows")]=extracted["rows"]
            if logs is not None:logs.append({"event":"browser-table-read","activityId":a["id"],"rows":len(extracted["rows"]),"columns":len(extracted["rows"][0]) if extracted["rows"] else 0})
        elif kind == "browser_scroll":
            direction = a.get("direction", "down")
            distance = int(a.get("distancePx", 600))
            delay = int(a.get("delayMilliseconds", 700)) / 1000
            to_bottom = bool(a.get("toBottom", False))
            max_steps = int(a.get("maxSteps", 100))
            container_selector = str(a.get("containerSelector", ""))
            scroll_container = None
            if container_selector:
                try:
                    scroll_container = self._browser_page.locator(container_selector)
                    match_count = scroll_container.count()
                    if match_count != 1:
                        raise WorkflowError(f"Scroll container selector must match exactly one element; found {match_count}", a["id"], "browser-scroll-container-match")
                    if not scroll_container.evaluate("element => element instanceof Element"):
                        raise WorkflowError("Scroll container selector must identify an element", a["id"], "browser-scroll-container-type")
                except WorkflowError:
                    raise
                except Exception as exc:
                    raise WorkflowError(f"Could not resolve scroll container selector ({type(exc).__name__})", a["id"], "browser-scroll-container-selector") from exc
            sign = 1 if direction == "down" else -1
            completed = 0
            previous = None
            stable_bottom = 0
            reached_bottom = False
            metrics = {"scrollY": 0, "scrollTop": 0, "viewportHeight": 0, "documentHeight": 0, "progressPercent": 0}
            steps = max_steps if to_bottom else 1
            for step in range(steps):
                self._check_cancel()
                input_mode = "wheel"
                if scroll_container is not None:
                    scroll_container.scroll_into_view_if_needed(timeout=timeout)
                    box = scroll_container.bounding_box()
                    if not box or box["width"] <= 0 or box["height"] <= 0:
                        raise WorkflowError("Scroll container is not visible or has no size", a["id"], "browser-scroll-container-hidden")
                    container_top_before = scroll_container.evaluate("element => element.scrollTop || 0")
                    page_top_before = self._browser_page.evaluate("window.scrollY || 0")
                    self._browser_page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                self._browser_page.mouse.wheel(0, sign * distance)
                deadline = time.monotonic() + delay
                while time.monotonic() < deadline:
                    self._check_cancel()
                    time.sleep(min(.05, deadline - time.monotonic()))
                if scroll_container is not None:
                    container_top_after_wheel = scroll_container.evaluate("element => element.scrollTop || 0")
                    can_move_container = scroll_container.evaluate("(element, direction) => direction > 0 ? element.scrollTop + element.clientHeight < element.scrollHeight - 1 : element.scrollTop > 1", sign)
                    if abs(container_top_after_wheel - container_top_before) < 0.5 and can_move_container:
                        # Some browser engines route synthetic wheel input to the document even
                        # when the pointer is over a nested overflow region. Restore that outer
                        # position and use the element's native scroll API, which still emits the
                        # scroll event used by virtualized and lazy-loaded lists.
                        page_top_after_wheel = self._browser_page.evaluate("window.scrollY || 0")
                        if abs(page_top_after_wheel - page_top_before) > 0.5:
                            self._browser_page.evaluate("position => window.scrollTo(0, position)", page_top_before)
                        scroll_container.evaluate("(element, delta) => element.scrollBy({top: delta, behavior: 'instant'})", sign * distance)
                        input_mode = "native-scroll-fallback"
                        deadline = time.monotonic() + delay
                        while time.monotonic() < deadline:
                            self._check_cancel()
                            time.sleep(min(.05, deadline - time.monotonic()))
                    metrics = scroll_container.evaluate("""element => {
                      const scrollTop = element.scrollTop || 0;
                      const viewportHeight = element.clientHeight || 0;
                      const documentHeight = element.scrollHeight || 0;
                      return {scrollY: window.scrollY || 0, scrollTop, viewportHeight, documentHeight, progressPercent: documentHeight <= viewportHeight ? 100 : Math.min(100, Math.round((scrollTop + viewportHeight) / documentHeight * 100))};
                    }""")
                else:
                    metrics = self._browser_page.evaluate("""() => {
                  const root = document.documentElement;
                  const scrollY = window.scrollY || root.scrollTop || 0;
                  const viewportHeight = window.innerHeight || root.clientHeight || 0;
                  const documentHeight = Math.max(root.scrollHeight, document.body?.scrollHeight || 0);
                  return {scrollY, scrollTop: scrollY, viewportHeight, documentHeight, progressPercent: documentHeight <= viewportHeight ? 100 : Math.min(100, Math.round((scrollY + viewportHeight) / documentHeight * 100))};
                }""")
                completed += 1
                at_bottom = metrics["scrollTop"] + metrics["viewportHeight"] >= metrics["documentHeight"] - 2
                if at_bottom and metrics == previous:
                    stable_bottom += 1
                elif at_bottom:
                    stable_bottom = 1
                else:
                    stable_bottom = 0
                reached_bottom = at_bottom
                if logs is not None:
                    self._emit({"event": "browser-scroll-progress", "activityId": a["id"], "step": completed, "maxSteps": steps, "target": "container" if scroll_container is not None else "page", "inputMode": input_mode, **metrics}, logs)
                previous = metrics
                if to_bottom and stable_bottom >= 2:
                    break
            v[a.get("result", "scrollResult")] = {"scrollsCompleted": completed, "reachedBottom": reached_bottom, "target": "container" if scroll_container is not None else "page", **metrics}
        elif kind == "browser_screenshot":
            path = self._path(a["path"], v, True)
            self._browser_page.screenshot(path=str(path), full_page=bool(a.get("fullPage", False)), timeout=timeout)
        else: raise WorkflowError(f"Unsupported browser activity: {kind}", a["id"], "activity-unsupported")

    def _close_browser(self) -> None:
        if self._browser_persistent_profile and self._browser_context is not None:
            try:self._browser_context.close()
            except Exception:pass
        elif self._browser is not None:
            try: self._browser.close()
            except Exception: pass
        if self._browser_runtime is not None:
            try: self._browser_runtime.stop()
            except Exception: pass
        self._browser_runtime = self._browser = self._browser_context = self._browser_page = None
        self._browser_persistent_profile=False
        self._browser_tabs = {}
        self._browser_unclaimed_popups = []
        self._browser_popup_counter = 0
        self._browser_responses = {}

    def _register_browser_page(self, page: Any, name: str) -> None:
        # Keep this method usable by focused adapter tests and partially
        # constructed engines as well as normal WorkflowEngine instances.
        if not hasattr(self, "_browser_responses"):
            self._browser_responses = {}
        already_registered = any(item is page for item in self._browser_tabs.values())
        for existing, item in list(self._browser_tabs.items()):
            if item is page and existing != name:
                self._browser_tabs.pop(existing, None)
        self._browser_unclaimed_popups = [item for item in self._browser_unclaimed_popups if item is not page]
        if not already_registered:
            page.set_default_timeout(30_000)
            page.set_default_navigation_timeout(60_000)
            page.route("**/*", self._guard_browser_request)
            route_web_socket = getattr(page, "route_web_socket", None)
            if callable(route_web_socket):
                route_web_socket("**/*", lambda socket: socket.close(code=1008, reason="Open RPA Studio blocks WebSocket connections"))
            self._browser_responses[page] = []
            page.on("response", lambda response, observed_page=page: self._record_browser_response(observed_page, response))
        self._browser_tabs[name] = page

    def _record_browser_response(self, page: Any, response: Any) -> None:
        """Keep bounded response metadata for synchronization; bodies are read only by explicit JSON mapping."""
        try:
            from urllib.parse import urlsplit, urlunsplit
            parts=urlsplit(str(response.url))
            host=parts.hostname or ""
            if ":" in host and not host.startswith("["):host=f"[{host}]"
            port=f":{parts.port}" if parts.port is not None else ""
            safe_url=urlunsplit((parts.scheme,host+port,parts.path,"",""))[:2048]
            content_type=str(response.headers.get("content-type","")).split(";",1)[0][:128]
            length_header=response.headers.get("content-length","")
            encoding=str(response.headers.get("content-encoding","identity")).lower()
            body_unavailable="not-json"
            response_reference=None
            if content_type.lower()=="application/json" or content_type.lower().endswith("+json"):
                if not length_header or not str(length_header).isdigit():body_unavailable="size-unknown"
                elif int(length_header)>2_000_000:body_unavailable="size-limit"
                elif encoding not in {"","identity"}:body_unavailable="compressed"
                else:response_reference=response;body_unavailable=""
            metadata={
                "url":safe_url,
                "method":str(response.request.method).upper()[:12],
                "status":int(response.status),
                "contentType":content_type,
                "observedAt":time.monotonic(),
                "_response":response_reference,
                "_bodyUnavailable":body_unavailable,
            }
            entries=self._browser_responses.setdefault(page,[])
            entries.append(metadata)
            if len(entries)>100:del entries[:-100]
            body_entries=[entry for entry in entries if entry.get("_response") is not None]
            for entry in body_entries[:-10]:
                entry["_response"]=None;entry["_bodyUnavailable"]="expired"
        except Exception:
            return

    def _wait_browser_response(self, activity: dict[str, Any], variables: dict[str, Any],
                               logs: list[dict[str, Any]] | None = None) -> None:
        page=self._browser_page
        url_contains=str(activity.get("urlContains",""))
        expected_method=str(activity.get("method","any")).upper()
        expected_status=activity.get("status")
        timeout=float(activity.get("timeoutSeconds",15))
        max_age=float(activity.get("maxAgeSeconds",30))
        deadline=time.monotonic()+timeout
        while True:
            self._check_cancel()
            now=time.monotonic()
            entries=self._browser_responses.setdefault(page,[])
            entries[:]=[item for item in entries if now-item["observedAt"]<=max_age]
            match=next((item for item in entries
                        if url_contains in item["url"]
                        and (expected_method=="ANY" or item["method"]==expected_method)
                        and (expected_status is None or item["status"]==expected_status)),None)
            if match is not None:
                entries.remove(match)
                result={"method":match["method"],"status":match["status"],"contentType":match["contentType"]}
                variables[activity.get("result","response")]=result
                mapping=activity.get("jsonMapping")
                mapped_count=0
                if mapping is not None:
                    mapped=self._map_browser_json_response(match,mapping,activity)
                    variables[mapping["result"]]=mapped
                    mapped_count=len(mapped)
                schema_discovered=False
                if activity.get("discoverJsonSchema"):
                    discovered=self._inspect_browser_json_response(match,activity)
                    variables[activity["schemaResult"]]=discovered
                    schema_discovered=True
                self.options.audit("browser.response.observed",{"activityId":activity["id"],**result})
                if logs is not None:
                    self._emit({"event":"browser-response-observed","activityId":activity["id"],**result,**({"mappedItems":mapped_count,"mappedFields":len(mapping.get("fields",[]))} if mapping is not None else {}),**({"schemaDiscovered":True,"schemaObjectArrays":len(discovered["objectArrayPointers"]),"schemaTruncated":discovered["truncated"]} if schema_discovered else {})},logs)
                return
            remaining=deadline-now
            if remaining<=0:
                raise WorkflowError(
                    f"No browser response matched the configured URL substring, method, and status within {timeout:g}s",
                    activity["id"],"browser-response-timeout")
            time.sleep(min(.05,remaining))

    def _read_bounded_browser_json_response(self, match: dict[str,Any], activity: dict[str,Any]) -> Any:
        aid=activity["id"]
        if not 200<=match["status"]<300:
            raise WorkflowError("Browser JSON inspection requires a successful 2xx response",aid,"browser-response-json-status")
        content_type=match["contentType"].lower()
        if content_type!="application/json" and not content_type.endswith("+json"):
            raise WorkflowError("Matched browser response is not a JSON content type",aid,"browser-response-json-content-type")
        response=match.get("_response")
        if response is None:
            reason=match.get("_bodyUnavailable")
            messages={"size-unknown":"Browser JSON mapping requires a response with a declared Content-Length","size-limit":"Browser JSON response exceeds the 2,000,000-byte runtime cap","compressed":"Browser JSON mapping does not read compressed response bodies","expired":"The bounded browser response body reference expired; trigger the response again"}
            message=messages.get(reason,"Matched browser response body is unavailable or is not JSON")
            raise WorkflowError(message,aid,"browser-response-body-unavailable")
        max_bytes=int(activity.get("maxResponseBytes",500_000))
        try:
            length_header=response.header_value("content-length")
            if not length_header or not length_header.isdigit():
                raise WorkflowError("Browser JSON inspection requires a response with a bounded Content-Length header",aid,"browser-response-size-unknown")
            declared_length=int(length_header)
            if declared_length>max_bytes:
                raise WorkflowError(f"Browser JSON response declares {declared_length} bytes; configured limit is {max_bytes}",aid,"browser-response-size-limit")
            encoding=(response.header_value("content-encoding") or "identity").lower()
            if encoding not in {"","identity"}:
                raise WorkflowError("Browser JSON inspection accepts only uncompressed responses so its byte limit is enforceable",aid,"browser-response-encoding")
            body=response.body()
        except WorkflowError:
            raise
        except Exception as exc:
            raise WorkflowError(f"Could not read the matched browser JSON response ({type(exc).__name__})",aid,"browser-response-body-read") from exc
        if len(body)>max_bytes:
            raise WorkflowError(f"Browser JSON response exceeds the configured {max_bytes}-byte limit",aid,"browser-response-size-limit")
        try:
            return json.loads(body.decode("utf-8-sig"))
        except (UnicodeDecodeError,json.JSONDecodeError,RecursionError) as exc:
            raise WorkflowError("Matched browser response is not valid UTF-8 JSON",aid,"browser-response-json") from exc

    def _inspect_browser_json_response(self, match: dict[str,Any], activity: dict[str,Any]) -> dict[str,Any]:
        """Inspect JSON structure only; returned data contains no response values."""
        document=self._read_bounded_browser_json_response(match,activity)
        return _inspect_json_shape(document)

    def _map_browser_json_response(self, match: dict[str,Any], mapping: dict[str,Any], activity: dict[str,Any]) -> list[dict[str,Any]]:
        """Extract only explicitly mapped scalar fields from a bounded browser JSON response."""
        aid=activity["id"]
        document=self._read_bounded_browser_json_response(match,activity)
        try:
            records=_json_pointer_get(document,mapping.get("itemsPointer",""))
        except (KeyError,IndexError,ValueError,TypeError,RecursionError) as exc:
            raise WorkflowError("Matched browser response array pointer did not resolve",aid,"browser-response-json") from exc
        if not isinstance(records,list):
            raise WorkflowError("Browser JSON response array pointer must select a JSON array",aid,"browser-response-json-array")
        if len(records)>int(mapping.get("maxItems",500)):
            raise WorkflowError("Browser JSON response array exceeds the configured item limit",aid,"browser-response-json-limit")
        mapped=[]
        for record in records:
            self._check_cancel()
            if not isinstance(record,dict):
                raise WorkflowError("Browser JSON response array items must be objects",aid,"browser-response-json-record")
            row={}
            for field_mapping in mapping.get("fields",[]):
                try:value=_json_pointer_get(record,field_mapping["path"])
                except (KeyError,IndexError,ValueError,TypeError) as exc:
                    raise WorkflowError("A browser JSON response field pointer did not resolve",aid,"browser-response-json-field-pointer") from exc
                if value is not None and not isinstance(value,(str,bool,int,float)):
                    raise WorkflowError("Mapped browser JSON fields must be scalar values or null",aid,"browser-response-json-field-type")
                if isinstance(value,float) and not math.isfinite(value):
                    raise WorkflowError("Mapped browser JSON number is not finite",aid,"browser-response-json-field-type")
                row[field_mapping["name"]]=value
            mapped.append(row)
        return mapped

    def _on_browser_popup(self, page: Any) -> None:
        if any(item is page for item in self._browser_tabs.values()):
            return
        self._browser_popup_counter += 1
        name = f"popup_{self._browser_popup_counter}"
        while name in self._browser_tabs:
            self._browser_popup_counter += 1
            name = f"popup_{self._browser_popup_counter}"
        self._register_browser_page(page, name)
        self._browser_unclaimed_popups.append(page)

    def _guard_browser_request(self, route: Any) -> None:
        try:
            url = route.request.url
            if url.startswith(("http://", "https://")):
                if not self.options.allow_network:
                    raise EgressDenied("Browser requests require explicit per-run network permission")
                assert_public_url(url)
            route.continue_()
        except Exception:
            route.abort()

    def _navigate_browser_page(self, page: Any, url: str, activity: dict[str, Any], variables: dict[str, Any], timeout_seconds: float) -> None:
        if not url.startswith(("https://", "http://")):
            raise WorkflowError("Browser URL must use http or https", activity["id"], "browser-url")
        if not self.options.allow_network:
            raise WorkflowError("Browser access requires explicit per-run network permission", activity["id"], "network-permission-denied")
        try:
            assert_public_url(url)
        except EgressDenied as exc:
            raise WorkflowError(str(exc), activity["id"], "http-egress-denied") from exc
        timeout = min(max(float(timeout_seconds), 1), 120) * 1000
        page.goto(url, wait_until=activity.get("waitUntil", "domcontentloaded"), timeout=timeout)

    def _browser_scope(self,a:dict[str,Any])->Any:
        selector=a.get("frameSelector")
        if selector is None or selector=="":return self._browser_page
        if not isinstance(selector,str) or not selector.strip():
            raise WorkflowError("Browser frameSelector must be a non-empty CSS selector",a["id"],"browser-frame-selector")
        try:
            iframe=self._browser_page.locator(selector)
            count=iframe.count()
            if count!=1:raise WorkflowError(f"Browser frameSelector must match exactly one iframe; found {count}",a["id"],"browser-frame-ambiguous")
            if not iframe.evaluate("element => element instanceof HTMLIFrameElement"):
                raise WorkflowError("Browser frameSelector must identify an iframe element",a["id"],"browser-frame-selector")
            return self._browser_page.frame_locator(selector)
        except WorkflowError:raise
        except Exception as exc:
            raise WorkflowError(f"Could not resolve the configured iframe selector ({type(exc).__name__})",a["id"],"browser-frame-selector") from exc

    def _locator(self, a: dict[str, Any], logs: list[dict[str, Any]] | None = None, allow_fallback: bool = True) -> Any:
        scope=self._browser_scope(a)
        def make_locator(kind: str, value: str, exact: bool = True) -> Any:
            if kind == "role":
                role, _, name = value.partition(":")
                return scope.get_by_role(role, name=name or None, exact=exact)
            if kind == "label": return scope.get_by_label(value, exact=exact)
            if kind == "text": return scope.get_by_text(value, exact=exact)
            if kind == "placeholder": return scope.get_by_placeholder(value, exact=exact)
            if kind == "test_id": return scope.get_by_test_id(value)
            if kind == "css": return scope.locator(value)
            raise WorkflowError(f"Unknown selector kind: {kind}", a["id"], "selector-kind")

        kind, value = a.get("selectorKind", "role"), str(a.get("selector", ""))
        if not value: raise WorkflowError("Browser selector is required", a["id"], "selector-empty")
        exact = bool(a.get("exact", True))
        primary=make_locator(kind,value,exact)
        fallbacks=a.get("selectorFallbacks",[])
        if not allow_fallback:
            return primary
        if not fallbacks and a.get("selectorHealing","off")!="accessible_role":
            return primary
        primary_count=primary.count()
        if primary_count==1:
            return primary
        matches=[]
        ambiguous=[]
        for index,fallback in enumerate(fallbacks):
            locator=make_locator(str(fallback["selectorKind"]),str(fallback["selector"]),bool(fallback.get("exact",True)))
            count=locator.count()
            if count==1: matches.append((index,locator))
            elif count>1: ambiguous.append(index)
        if ambiguous or len(matches)>1 or primary_count>1:
            reason="ambiguous" if ambiguous or len(matches)>1 or primary_count>1 else "not found"
            code="selector-fallback-ambiguous" if reason=="ambiguous" else "selector-fallback-not-found"
            raise WorkflowError(f"Browser selector {reason}: primary matched {primary_count}; {len(matches)} fallback(s) matched uniquely and {len(ambiguous)} fallback(s) were ambiguous. Make the explicit selector chain resolve to exactly one control.",a["id"],code)
        if matches:
            fallback_index,locator=matches[0]
            details={"activityId":a["id"],"fallbackIndex":fallback_index}
            self.options.audit("browser.selector_fallback",details)
            if logs is not None:
                self._emit({"event":"browser-selector-fallback","activityId":a["id"],"fallbackIndex":fallback_index},logs)
            return locator
        if a.get("selectorHealing","off")=="accessible_role" and primary_count==0:
            role,separator,expected=value.partition(":") if kind=="role" else ("","","")
            if separator and role and expected and exact:
                return self._heal_accessible_role(scope,role,expected,a,logs)
        if fallbacks:
            raise WorkflowError(f"Browser selector not found: primary matched {primary_count}; no configured fallback matched uniquely. Make the explicit selector chain resolve to exactly one control.",a["id"],"selector-fallback-not-found")
        return primary

    def _heal_accessible_role(self, scope: Any, role: str, expected: str, activity: dict[str, Any], logs: list[dict[str, Any]] | None) -> Any:
        """Recover only a missing exact role/name selector with one strong visible match."""
        expected_normalized=_normalize_accessible_name(expected)
        if not expected_normalized:
            raise WorkflowError("Accessible selector recovery needs a non-empty accessible name",activity["id"],"browser-selector-heal-not-found")
        try:
            candidates=scope.get_by_role(role)
            count=candidates.count()
        except Exception as exc:
            raise WorkflowError(f"Could not inspect same-role browser controls for safe selector recovery ({type(exc).__name__})",activity["id"],"browser-selector-heal-failed") from exc
        if count>250:
            raise WorkflowError("Accessible selector recovery stopped: this role has more than 250 controls",activity["id"],"browser-selector-heal-limit")
        matches=[]
        for index in range(count):
            candidate=candidates.nth(index)
            try:
                if not candidate.is_visible() or not candidate.is_enabled():
                    continue
                snapshot=candidate.aria_snapshot()
                accessible_name=_accessible_role_name(snapshot,role)
                normalized=_normalize_accessible_name(accessible_name)
                if not normalized:
                    continue
                if normalized==expected_normalized:
                    score=1.0
                elif min(len(normalized),len(expected_normalized))>=10:
                    score=SequenceMatcher(None,expected_normalized,normalized,autojunk=False).ratio()
                else:
                    continue
                if score>=0.94:
                    matches.append((score,index,candidate))
            except Exception:
                continue
        matches.sort(key=lambda item:item[0],reverse=True)
        if not matches:
            raise WorkflowError(f"Accessible selector recovery found no safe {role} match for the configured name",activity["id"],"browser-selector-heal-not-found")
        if len(matches)>1:
            if matches[0][0]==1.0 or matches[0][0]-matches[1][0]<0.10:
                raise WorkflowError(f"Accessible selector recovery found {len(matches)} plausible {role} controls; no unique match is safe",activity["id"],"browser-selector-heal-ambiguous")
        score,index,locator=matches[0]
        self.options.audit("browser.selector_healed",{"activityId":activity["id"],"role":role,"candidateIndex":index,"similarity":round(score,3)})
        if logs is not None:
            self._emit({"event":"browser-selector-healed","activityId":activity["id"],"role":role,"similarity":round(score,3)},logs)
        return locator

    def _desktop_activity(self, a: dict[str, Any], v: dict[str, Any]) -> None:
        if not self.options.allow_desktop:
            raise WorkflowError("Desktop control requires explicit per-run permission", a["id"], "desktop-permission-denied")
        kind=a["type"]
        if kind == "desktop_ocr_click":
            self._check_cancel()
            try:
                from .windows_uia import ocr_click
                title=str(interpolate(a.get("windowTitle", ""),v))
                target=str(interpolate(a.get("text", ""),v))
                self.options.audit("desktop.ocr_click.attempt",{"activityId":a["id"],"processId":a.get("processId"),"confidenceThreshold":a.get("confidenceThreshold",80)})
                result=ocr_click(title,target,int(a.get("confidenceThreshold",80)),float(a.get("timeoutSeconds",15)),a.get("processId"),self.options.cancel)
                self._check_cancel()
                v[a.get("result","ocrMatch")]=result
                self.options.audit("desktop.ocr_click.succeeded",{"activityId":a["id"],"processId":a.get("processId"),"confidence":result["confidence"]})
                return
            except WorkflowError:
                raise
            except Exception as exc:
                self._check_cancel()
                message=f"{type(exc).__name__}: {exc}"
                raise WorkflowError(f"OCR desktop click failed: {message}",a["id"],"desktop-ocr-click-failed") from exc
        if kind.startswith("desktop_window_"):
            try:
                from . import desktop_accessibility
                if not desktop_accessibility.status().get("available"):
                    raise WorkflowError("Desktop accessibility is unavailable",a["id"],"runtime-unavailable")
                title=str(interpolate(a.get("windowTitle", ""),v) or "").strip() or None
                process_id=a.get("processId")
                if kind=="desktop_window_wait":
                    desired=a.get("state","exists");deadline=time.monotonic()+float(a.get("timeoutSeconds",30))
                    while True:
                        self._check_cancel()
                        exists=desktop_accessibility.window_exists(title,process_id)
                        if exists is (desired=="exists"):return
                        remaining=deadline-time.monotonic()
                        if remaining<=0:raise WorkflowError(f"Timed out waiting for the target desktop window to {desired.replace('_',' ')}",a["id"],"desktop-window-timeout")
                        time.sleep(min(0.2,remaining))
                self._check_cancel()
                action={"desktop_window_activate":desktop_accessibility.activate_window,
                        "desktop_window_minimize":desktop_accessibility.minimize_window,
                        "desktop_window_restore":desktop_accessibility.restore_window,
                        "desktop_window_close":desktop_accessibility.close_window}[kind]
                action(title,process_id)
                self.options.audit("desktop.window.action",{"activityId":a["id"],"action":kind.removeprefix("desktop_window_"),"processId":process_id})
                if kind=="desktop_window_close":
                    deadline=time.monotonic()+float(a.get("timeoutSeconds",15))
                    while True:
                        self._check_cancel()
                        if not desktop_accessibility.window_exists(title,process_id):return
                        remaining=deadline-time.monotonic()
                        if remaining<=0:raise WorkflowError("The close request was sent, but the target window remained open; it may be waiting for an in-app confirmation",a["id"],"desktop-window-close-timeout")
                        time.sleep(min(0.2,remaining))
                return
            except WorkflowError:
                raise
            except Exception as exc:
                raise WorkflowError(f"Desktop window action failed: {type(exc).__name__}: {exc}",a["id"],"desktop-window-error") from exc
        if kind.startswith("desktop_ui_"):
            try:
                from . import desktop_accessibility
                runtime=desktop_accessibility.status()
                if not runtime.get("available"):
                    raise WorkflowError(runtime.get("reason") or "Windows UI Automation is unavailable",a["id"],"runtime-unavailable")
                title=str(interpolate(a.get("windowTitle",""),v))
                if kind=="desktop_ui_select_many":
                    selected=desktop_accessibility.set_selection(
                        title,a.get("name"),a.get("automationId"),a.get("controlType","List"),
                        a.get("values",[]),a.get("mode","replace"),a.get("processId"))
                    v[a.get("result","selectedItems")]=selected
                    return
                if kind == "desktop_ui_grid_read":
                    name=a.get("name");automation_id=a.get("automationId");control_type=a.get("controlType","DataGrid");process_id=a.get("processId")
                    timeout=min(max(float(a.get("timeoutSeconds",10)),1),120);deadline=time.monotonic()+timeout
                    while True:
                        self._check_cancel()
                        observed=desktop_accessibility.control_state(title,name,automation_id,control_type,process_id)
                        if observed is not None and observed.get("visible") and observed.get("enabled"):break
                        remaining=deadline-time.monotonic()
                        if remaining<=0:raise WorkflowError(f"Timed out after {timeout:g}s waiting for the selected UIA grid to become visible and enabled",a["id"],"desktop-uia-timeout")
                        time.sleep(min(0.2,remaining))
                    rows, metadata = desktop_accessibility.read_grid(
                        title, name, automation_id, control_type,
                        a.get("maxRows",100), a.get("maxColumns",30), a.get("maxOutputChars",50000), process_id)
                    v[a.get("result","rows")]=rows
                    metadata_result=a.get("metadataResult","gridMetadata")
                    if metadata_result:v[metadata_result]=metadata
                    return
                if kind == "desktop_ui_set_range_value":
                    name=a.get("name");automation_id=a.get("automationId");control_type=a.get("controlType","Slider");process_id=a.get("processId")
                    timeout=min(max(float(a.get("timeoutSeconds",10)),1),120);deadline=time.monotonic()+timeout
                    while True:
                        self._check_cancel()
                        observed=desktop_accessibility.control_state(title,name,automation_id,control_type,process_id)
                        if observed is not None and observed.get("visible") and observed.get("enabled"):break
                        remaining=deadline-time.monotonic()
                        if remaining<=0:raise WorkflowError(f"Timed out after {timeout:g}s waiting for the selected Slider or Spinner to become visible and enabled",a["id"],"desktop-uia-timeout")
                        time.sleep(min(0.2,remaining))
                    requested=interpolate(a.get("value"),v)
                    if isinstance(requested,bool) or not isinstance(requested,(int,float)) or not math.isfinite(requested):raise WorkflowError("Set UI Range Value expression must produce a finite number",a["id"],"desktop-uia-range-value")
                    desktop_accessibility.set_range_value(title,name,automation_id,control_type,float(requested),process_id)
                    return
                if kind == "desktop_ui_drag":
                    self._check_cancel()
                    desktop_accessibility.drag(title, a.get("sourceName"), a.get("sourceAutomationId"), a.get("sourceControlType"),
                                               a.get("targetName"), a.get("targetAutomationId"), a.get("targetControlType"),
                                               a.get("durationMilliseconds",500), a.get("processId"))
                    return
                name=a.get("name")
                automation_id=a.get("automationId")
                control_type=a.get("controlType")
                process_id=a.get("processId")
                timeout=min(max(float(a.get("timeoutSeconds",10)),1),120)
                if kind=="desktop_ui_scroll_into_view":
                    self._check_cancel()
                    desktop_accessibility.scroll_into_view(title,name,automation_id,control_type,process_id)
                    return
                def state() -> dict[str,bool] | None:
                    self._check_cancel()
                    return desktop_accessibility.control_state(title,name,automation_id,control_type,process_id)
                desired={"desktop_ui_wait":a.get("state","exists"),
                         "desktop_ui_click":"enabled","desktop_ui_press":"enabled","desktop_ui_fill":"enabled",
                         "desktop_ui_invoke":"enabled",
                         "desktop_ui_text":"exists","desktop_ui_select":"enabled",
                         "desktop_ui_check":"enabled","desktop_ui_uncheck":"enabled","desktop_ui_set_range_value":"enabled",
                         "desktop_ui_expand":"enabled","desktop_ui_collapse":"enabled"}[kind]
                deadline=time.monotonic()+timeout
                observed=None
                while True:
                    observed=state()
                    ready=(observed is not None if desired=="exists" else
                           observed is None or not observed["visible"] if desired=="not_visible" else
                           observed is not None and bool(observed.get("checked")) if desired=="checked" else
                           observed is not None and observed.get("checked") is False if desired=="unchecked" else
                           observed is not None and bool(observed.get("selected")) if desired=="selected" else
                           observed is not None and observed.get("selected") is False if desired=="unselected" else
                           observed is not None and observed.get("expanded") is True if desired=="expanded" else
                           observed is not None and observed.get("expanded") is False if desired=="collapsed" else
                           observed is not None and observed.get(desired,False))
                    if ready:break
                    remaining=deadline-time.monotonic()
                    if remaining<=0:
                        selector=f"window={title!r}, name={name!r}, automationId={automation_id!r}, controlType={control_type!r}"
                        raise WorkflowError(f"Timed out after {timeout:g}s waiting for UIA {desired}: {selector}",a["id"],"desktop-uia-timeout")
                    time.sleep(min(0.2,remaining))
                if kind=="desktop_ui_wait":return
                if kind=="desktop_ui_click":desktop_accessibility.click(title,name,automation_id,control_type,process_id)
                elif kind=="desktop_ui_invoke":desktop_accessibility.invoke(title,name,automation_id,control_type,process_id)
                elif kind=="desktop_ui_press":desktop_accessibility.press(title,name,automation_id,control_type,a.get("key","Enter"),process_id)
                elif kind=="desktop_ui_fill":desktop_accessibility.fill(title,name,automation_id,control_type,str(interpolate(a.get("value",""),v)),process_id)
                elif kind=="desktop_ui_set_range_value":
                    requested=interpolate(a.get("value"),v)
                    if isinstance(requested,bool) or not isinstance(requested,(int,float)) or not math.isfinite(requested):raise WorkflowError("Set UI Range Value expression must produce a finite number",a["id"],"desktop-uia-range-value")
                    desktop_accessibility.set_range_value(title,name,automation_id,control_type,float(requested),process_id)
                elif kind=="desktop_ui_select":desktop_accessibility.select(title,name,automation_id,control_type,str(interpolate(a.get("value",""),v)),process_id)
                elif kind=="desktop_ui_check":desktop_accessibility.set_checked(title,name,automation_id,control_type,True,process_id)
                elif kind=="desktop_ui_uncheck":desktop_accessibility.set_checked(title,name,automation_id,control_type,False,process_id)
                elif kind=="desktop_ui_expand":desktop_accessibility.set_expanded(title,name,automation_id,control_type,True,process_id)
                elif kind=="desktop_ui_collapse":desktop_accessibility.set_expanded(title,name,automation_id,control_type,False,process_id)
                elif kind=="desktop_ui_text":v[a.get("result","text")]=desktop_accessibility.read_text(title,name,automation_id,control_type,process_id)
                return
            except WorkflowError:
                raise
            except Exception as exc:
                raise WorkflowError(f"Desktop accessibility failed: {type(exc).__name__}: {exc}",a["id"],"desktop-uia-error") from exc
        if kind=="desktop_launch_app":
            self._check_cancel()
            raw_executable=str(interpolate(a.get("executable",""),v))
            try:
                executable=Path(raw_executable).expanduser().resolve(strict=True)
                if not executable.is_file():
                    raise WorkflowError("The configured application path is not a file",a["id"],"desktop-app-not-found")
                arguments=a.get("arguments",[])
                command=[str(executable),*(str(interpolate(argument,v)) for argument in arguments)]
                # Never invoke a shell or inherit the workflow engine's standard streams.
                creationflags=getattr(subprocess,"CREATE_NEW_PROCESS_GROUP",0) if os.name=="nt" else 0
                process=subprocess.Popen(command,cwd=str(self.root),stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,shell=False,close_fds=True,creationflags=creationflags,
                    start_new_session=(os.name!="nt"))
                pid=int(process.pid)
                import threading
                threading.Thread(target=process.wait,name=f"openrpa-child-reaper-{pid}",daemon=True).start()
                result=str(a.get("result","processId"))
                v[result]=pid
                self.options.audit("desktop.application_launched",{"activityId":a["id"],"executable":executable.name,"argumentCount":len(arguments),"processId":pid})
                return
            except WorkflowError:
                raise
            except (OSError,ValueError,TypeError) as exc:
                code="desktop-app-not-found" if isinstance(exc,FileNotFoundError) else "desktop-app-launch-error"
                raise WorkflowError(f"Could not launch the configured application ({type(exc).__name__}): {exc}",a["id"],code) from exc
        try:
            import pyautogui
        except ImportError as exc:
            raise WorkflowError("Desktop automation requires the optional PyAutoGUI adapter. See docs/AUTOMATION_ADAPTERS.md", a["id"], "runtime-unavailable") from exc
        if kind=="desktop_screenshot":
            path=self._path(a["path"],v,True);pyautogui.screenshot().save(path)
        elif kind=="desktop_click": pyautogui.click(int(a["x"]),int(a["y"]),clicks=int(a.get("clicks",1)))
        elif kind=="desktop_type": pyautogui.write(str(interpolate(a.get("text",""),v)),interval=float(a.get("interval",0.02)))
        elif kind=="desktop_hotkey": pyautogui.hotkey(*[x.strip() for x in str(a["keys"]).split("+")])
        elif kind=="desktop_press": pyautogui.press(str(a["key"]))
        else: raise WorkflowError(f"Unsupported desktop activity: {kind}",a["id"],"activity-unsupported")


def run_project(project: dict[str, Any], root: Path, workflow: str | None = None, inputs: dict[str, Any] | None = None, options: RunOptions | None = None) -> RunResult:
    engine = WorkflowEngine(project, root, options)
    name = workflow or project.get("main") or next(iter(engine.workflows))
    return engine.run(name, inputs)


def _matches_type(value: Any, declared: str) -> bool:
    return {"any":lambda:True,"string":lambda:isinstance(value,str),"integer":lambda:isinstance(value,int) and not isinstance(value,bool),
        "number":lambda:isinstance(value,(int,float)) and not isinstance(value,bool),"boolean":lambda:isinstance(value,bool),
        "object":lambda:isinstance(value,dict),"array":lambda:isinstance(value,list)}.get(declared,lambda:False)()
