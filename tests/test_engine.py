from __future__ import annotations
import json
from pathlib import Path
import pytest
from openrpa.engine import RunOptions, run_project
from openrpa.model import WorkflowError, load_project, validate_workflow
from openrpa.package import pack_project, unpack_project
from openrpa.expressions import ExpressionError, evaluate
import hashlib
from openpyxl import Workbook, load_workbook
from openpyxl.xml.functions import fromstring
from defusedxml.common import EntitiesForbidden

EXAMPLE = Path(__file__).parents[1] / "examples" / "csv-cleanup"

def test_real_csv_workflow_reads_transforms_and_writes(tmp_path: Path):
    project = {"format":"openrpa.project","formatVersion":1,"name":"CSV","main":"Main","workflows":{"Main":"main.json"}}
    workflow = {"format":"openrpa.workflow","formatVersion":1,"name":"CSV","variables":[],"arguments":[],"activities":[
        {"id":"read","type":"csv_read","path":"input.csv","result":"rows"},
        {"id":"norm","type":"transform_field","rows":"rows","field":"name","operation":"trim","result":"rows"},
        {"id":"write","type":"csv_write","path":"out.csv","rows":"rows"}]}
    (tmp_path/"main.json").write_text(json.dumps(workflow))
    (tmp_path/"input.csv").write_text("name,email\nAda,ada@example.org\nGrace,grace@example.org\n")
    result=run_project(project,tmp_path)
    assert result.status == "succeeded"
    assert len(result.variables["rows"]) == 2
    assert result.variables["rows"][0]["name"] == "Ada"
    assert "Grace" in (tmp_path/"out.csv").read_text()
    assert result.activities["write"]["status"] == "succeeded"

def test_flowchart_action_executes_nested_control_flow_activity(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Nested flowchart","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Nested flowchart","executionModel":"flowchart","variables":[],"arguments":[],"activities":[],"flowchart":{
        "start":"start","nodes":[
            {"id":"start","kind":"start"},
            {"id":"action-if","kind":"action","activity":{"id":"if-check","type":"if","condition":"True","then":[{"id":"set-result","type":"assign","target":"result","value":"nested branch executed"}],"else":[{"id":"set-result-other","type":"assign","target":"result","value":"wrong branch"}]}},
            {"id":"end","kind":"end"},
        ],"transitions":[{"from":"start","to":"action-if"},{"from":"action-if","to":"end"}],
    }}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["result"]=="nested branch executed"
    assert result.activities["set-result"]["status"]=="succeeded"

def test_xlsx_activities_round_trip_with_literal_formula_text_and_audit(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Excel","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Excel","variables":[{"name":"records","type":"array","value":[{"name":"=2+2","count":4},{"name":"Ada","count":7}]}],"arguments":[],"activities":[
        {"id":"write","type":"xlsx_write","path":"out.xlsx","sheetName":"Packages","rows":"records"},
        {"id":"read","type":"xlsx_read","path":"out.xlsx","sheetName":"Packages","result":"loaded"},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    audits=[]
    result=run_project(project,tmp_path,options=RunOptions(audit=lambda action,details:audits.append((action,details))))
    assert result.status=="succeeded",result.error
    assert result.variables["loaded"]==[{"name":"=2+2","count":4},{"name":"Ada","count":7}]
    assert result.activities["write"]["status"]=="succeeded" and result.activities["read"]["status"]=="succeeded"
    wb=load_workbook(tmp_path/"out.xlsx",data_only=False)
    assert wb["Packages"]["A2"].data_type=="s"
    wb.close()
    assert [(action,row["path"]) for action,row in audits if action in {"file.read","file.write"}]==[("file.write","out.xlsx"),("file.read","out.xlsx")]

def test_xlsx_read_uses_explicit_sheet_and_header_row(tmp_path: Path):
    wb=Workbook();sheet=wb.active;sheet.title="Orders";sheet.append(["Export"]);sheet.append(["order_id","amount"]);sheet.append(["O-1",12.5]);wb.save(tmp_path/"input.xlsx");wb.close()
    project={"format":"openrpa.project","formatVersion":1,"name":"Excel","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Excel","variables":[],"arguments":[],"activities":[{"id":"read","type":"xlsx_read","path":"input.xlsx","sheetName":"Orders","headerRow":2,"result":"orders"}]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded",result.error
    assert result.variables["orders"]==[{"order_id":"O-1","amount":12.5}]

def test_xlsx_rejects_oversized_expanded_archives_and_invalid_schema(tmp_path: Path):
    import zipfile
    archive=tmp_path/"bomb.xlsx"
    with zipfile.ZipFile(archive,"w",compression=zipfile.ZIP_DEFLATED) as zf:zf.writestr("xl/worksheets/sheet1.xml","A"*100_000)
    project={"format":"openrpa.project","formatVersion":1,"name":"Excel","main":"Main","workflows":{"Main":"main.json"},"limits":{"maxFileBytes":1000}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Excel","variables":[],"arguments":[],"activities":[{"id":"read","type":"xlsx_read","path":"bomb.xlsx"}]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed" and result.error and result.error["code"]=="file-size-limit"
    invalid={**workflow,"activities":[{"id":"read","type":"xlsx_read","path":"input.xlsx","headerRow":0,"sheetName":"bad/name"}]}
    problems=validate_workflow(invalid)
    assert {problem["code"] for problem in problems}>={"spreadsheet-header-row","spreadsheet-sheet-name"}
    invalid_writer={**workflow,"activities":[{"id":"write","type":"xlsx_write","path":"out.xlsx","rows":"records","fields":[]}]}
    assert any(problem["code"]=="spreadsheet-fields" for problem in validate_workflow(invalid_writer))

def test_conditional_breakpoint_schema_requires_enabled_breakpoint_and_bounded_text():
    base={"format":"openrpa.workflow","formatVersion":1,"name":"Debug","variables":[],"arguments":[],"activities":[
        {"id":"inspect","type":"log","breakpoint":True,"breakpointCondition":"amount > 10"}]}
    assert validate_workflow(base)==[]
    disabled={**base,"activities":[{**base["activities"][0],"breakpoint":False}]}
    assert any(problem["code"]=="breakpoint-condition-disabled" for problem in validate_workflow(disabled))
    oversized={**base,"activities":[{**base["activities"][0],"breakpointCondition":"x"*2049}]}
    assert any(problem["code"]=="breakpoint-condition" for problem in validate_workflow(oversized))

def test_conditional_breakpoints_are_validated_on_flowchart_nodes_and_states():
    flowchart={"format":"openrpa.workflow","formatVersion":1,"name":"Graph debug","variables":[],"arguments":[],"activities":[],
        "executionModel":"flowchart","flowchart":{"start":"start","nodes":[
            {"id":"start","kind":"start"},{"id":"step","kind":"action","activity":{"id":"log","type":"log","message":"ok"},"breakpoint":True,"breakpointCondition":"amount > 10"},
            {"id":"end","kind":"end"}],"transitions":[{"from":"start","to":"step"},{"from":"step","to":"end"}]}}
    assert validate_workflow(flowchart)==[]
    disabled={**flowchart,"flowchart":{**flowchart["flowchart"],"nodes":[{**flowchart["flowchart"]["nodes"][1],"breakpoint":False}]}}
    assert any(problem["code"]=="breakpoint-condition-disabled" for problem in validate_workflow(disabled))
    oversized={**flowchart,"flowchart":{**flowchart["flowchart"],"nodes":[{**flowchart["flowchart"]["nodes"][1],"breakpointCondition":"x"*2049}]}}
    assert any(problem["code"]=="breakpoint-condition" for problem in validate_workflow(oversized))

    machine={"format":"openrpa.workflow","formatVersion":1,"name":"State debug","variables":[],"arguments":[],"activities":[],
        "executionModel":"state_machine","stateMachine":{"initial":"review","states":[
            {"id":"review","label":"Review","breakpoint":True,"breakpointCondition":"amount > 10","activities":[]},
            {"id":"done","label":"Done","final":True,"activities":[]}],"transitions":[{"from":"review","to":"done"}]}}
    assert validate_workflow(machine)==[]
    invalid={**machine,"stateMachine":{**machine["stateMachine"],"states":[{**machine["stateMachine"]["states"][0],"breakpoint":False}]}}
    assert any(problem["code"]=="breakpoint-condition-disabled" for problem in validate_workflow(invalid))

def test_openpyxl_uses_defused_xml_parser_for_workbook_content():
    with pytest.raises(EntitiesForbidden):
        fromstring(b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///private-data">]><x>&secret;</x>')

def test_xlsx_rejects_dtd_and_entity_declarations_before_openpyxl(tmp_path: Path):
    import zipfile
    with zipfile.ZipFile(tmp_path/"entity.xlsx","w",compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml",b'<!DOCTYPE x [<!ENTITY e "expanded">]><x>&e;</x>')
    project={"format":"openrpa.project","formatVersion":1,"name":"Excel","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Excel","variables":[],"arguments":[],"activities":[{"id":"read","type":"xlsx_read","path":"entity.xlsx"}]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed" and result.error and result.error["code"]=="spreadsheet-xml-entity"


def test_while_activity_updates_values_and_stops_on_condition(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"While","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"While","variables":[{"name":"count","type":"integer","value":0}],"arguments":[],"activities":[
        {"id":"repeat","type":"while","condition":"count < 3","body":[{"id":"increment","type":"assign","target":"count","value":"{{count + 1}}"}]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["count"]==3
    assert result.activities["increment"]["status"]=="succeeded"

def test_while_activity_enforces_iteration_quota(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Quota","main":"Main","workflows":{"Main":"main.json"},"limits":{"maxLoopItems":2}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Quota","variables":[],"arguments":[],"activities":[
        {"id":"forever","type":"while","condition":"True","body":[]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed"
    assert result.error["code"]=="loop-limit"

def test_switch_activity_executes_matching_case_and_default(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Switch","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Switch","variables":[{"name":"route","type":"string","value":"approved"}],"arguments":[],"activities":[
        {"id":"route-choice","type":"switch","expression":"route","cases":[
            {"value":"approved","activities":[{"id":"approved-log","type":"log","message":"approved"}]},
            {"value":"rejected","activities":[{"id":"rejected-log","type":"log","message":"rejected"}]},
        ],"default":[{"id":"default-log","type":"log","message":"default"}]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert [entry["message"] for entry in result.logs if entry.get("event")=="log"]==["approved"]
    workflow["variables"][0]["value"]="pending"
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    default_result=run_project(project,tmp_path)
    assert default_result.status=="succeeded"
    assert [entry["message"] for entry in default_result.logs if entry.get("event")=="log"]==["default"]

def test_switch_activity_uses_strict_json_scalar_matching(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Switch types","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Switch types","variables":[{"name":"route","type":"integer","value":1}],"arguments":[],"activities":[
        {"id":"route-choice","type":"switch","expression":"route","cases":[
            {"value":True,"activities":[{"id":"boolean-log","type":"log","message":"boolean"}]},
            {"value":1,"activities":[{"id":"integer-log","type":"log","message":"integer"}]},
        ],"default":[]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert [entry["message"] for entry in result.logs if entry.get("event")=="log"]==["integer"]

def test_switch_validation_rejects_duplicate_values(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Bad switch","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Bad switch","variables":[],"arguments":[],"activities":[
        {"id":"switch","type":"switch","expression":"route","cases":[{"value":"a","activities":[]},{"value":"a","activities":[]}],"default":[]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    problems=validate_workflow(workflow,set(project["workflows"]))
    assert any(problem["code"]=="switch-case-duplicate" for problem in problems)
    workflow["activities"][0]["cases"]=[{"value":1,"activities":[]},{"value":1.0,"activities":[]}]
    numeric_problems=validate_workflow(workflow,set(project["workflows"]))
    assert any(problem["code"]=="switch-case-duplicate" for problem in numeric_problems)

def test_retry_scope_retries_selected_error_and_restores_variables(tmp_path: Path,monkeypatch):
    project={"format":"openrpa.project","formatVersion":1,"name":"Retry","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Retry","variables":[{"name":"count","type":"integer","value":0}],"arguments":[],"activities":[
        {"id":"retry","type":"retry","maxAttempts":3,"retryOn":["http-timeout"],"delaySeconds":0,"backoffFactor":2,"repeatSafe":True,"body":[
            {"id":"increment","type":"assign","target":"count","value":"{{count + 1}}"},
            {"id":"flaky","type":"http_request","url":"https://example.org","method":"GET","timeoutSeconds":1},
        ]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    calls={"count":0}
    def fail_once(url,method,headers,body,timeout,max_bytes):
        calls["count"]+=1
        if calls["count"]==1:raise TimeoutError("temporary network timeout")
        return {"status":200,"headers":{},"body":"ok","redirectBlocked":False}
    monkeypatch.setattr("openrpa.engine.restricted_http_request",fail_once)
    result=run_project(project,tmp_path,options=RunOptions(allow_network=True))
    assert result.status=="succeeded"
    assert calls["count"]==2
    assert result.variables["count"]==1
    assert result.variables["response"]["body"]=="ok"
    attempts=[event for event in result.logs if event.get("event")=="retry-attempt-failed"]
    assert attempts==[{"event":"retry-attempt-failed","activityId":"retry","attempt":1,"maxAttempts":3,"errorCode":"http-timeout","willRetry":True}]
    assert any(event.get("event")=="retry-succeeded" for event in result.logs)

def test_retry_scope_respects_error_filter_and_requires_safety(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Retry filter","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Retry filter","variables":[],"arguments":[],"activities":[
        {"id":"retry","type":"retry","maxAttempts":3,"retryOn":["different-error"],"delaySeconds":0,"backoffFactor":1,"repeatSafe":True,"body":[{"id":"bad","type":"file_read_text","path":"../outside","result":"text"}]},
    ]}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed"
    assert result.error["code"]=="file-permission-denied"
    assert len([event for event in result.logs if event.get("event")=="retry-attempt-failed"])==1
    workflow["activities"][0]["repeatSafe"]=False
    problems=validate_workflow(workflow,set(project["workflows"]))
    assert any(problem["code"]=="retry-safety" for problem in problems)
    workflow["activities"][0]["repeatSafe"]=True
    workflow["activities"][0]["retryOn"]=[]
    problems=validate_workflow(workflow,set(project["workflows"]))
    assert any(problem["code"]=="retry-codes-empty" for problem in problems)

def test_flowchart_executes_selected_decision_path(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Flowchart","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Flowchart","executionModel":"flowchart","variables":[{"name":"approved","type":"boolean","value":True}],"arguments":[],"activities":[],"flowchart":{
        "start":"start","nodes":[
            {"id":"start","kind":"start","label":"Start"},
            {"id":"decision","kind":"decision","label":"Approved?","condition":"approved"},
            {"id":"yes-node","kind":"action","activity":{"id":"yes-action","type":"assign","target":"result","value":"yes"}},
            {"id":"no-node","kind":"action","activity":{"id":"no-action","type":"assign","target":"result","value":"no"}},
            {"id":"end","kind":"end","label":"End"},
        ],"transitions":[
            {"from":"start","to":"decision"},{"from":"decision","to":"yes-node","when":"true"},{"from":"decision","to":"no-node","when":"false"},{"from":"yes-node","to":"end"},{"from":"no-node","to":"end"},
        ],
    }}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    problems=validate_workflow(workflow,set(project["workflows"]))
    assert problems==[]
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["result"]=="yes"
    assert result.activities["yes-action"]["status"]=="succeeded"
    assert "no-action" not in result.activities
    workflow["variables"][0]["value"]=False
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded" and result.variables["result"]=="no"

def test_flowchart_validator_rejects_unreachable_nodes_and_bad_decision_edges():
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Bad graph","executionModel":"flowchart","variables":[],"arguments":[],"activities":[],"flowchart":{
        "start":"start","nodes":[{"id":"start","kind":"start"},{"id":"decision","kind":"decision","condition":"True"},{"id":"end","kind":"end"},{"id":"orphan","kind":"end"}],
        "transitions":[{"from":"start","to":"decision"},{"from":"decision","to":"end","when":"true"}],
    }}
    problems=validate_workflow(workflow)
    codes={problem["code"] for problem in problems}
    assert "flowchart-decision-edges" in codes
    assert "flowchart-unreachable" in codes

def test_flowchart_node_coordinates_are_bounded_and_backward_compatible():
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Positions","executionModel":"flowchart","variables":[],"arguments":[],"activities":[],"flowchart":{
        "start":"start","nodes":[{"id":"start","kind":"start"},{"id":"end","kind":"end","x":850,"y":420}],"transitions":[{"from":"start","to":"end"}],
    }}
    assert validate_workflow(workflow)==[]
    workflow["flowchart"]["nodes"][1]["x"]=10001
    assert any(problem["code"]=="flowchart-position" for problem in validate_workflow(workflow))

def test_invoke_workflow_runs_flowchart_and_maps_typed_outputs(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Reusable","main":"Main","workflows":{"Main":"main.json","Child":"child.json"}}
    main={"format":"openrpa.workflow","formatVersion":1,"name":"Main","variables":[],"arguments":[],"activities":[
        {"id":"call","type":"invoke_workflow","workflow":"Child","inputs":{"amount":"7"}},
        {"id":"done","type":"log","message":"Result={{total}}"},
    ]}
    child={"format":"openrpa.workflow","formatVersion":1,"name":"Child","executionModel":"flowchart","variables":[],"arguments":[
        {"name":"amount","type":"string","direction":"in"},{"name":"total","type":"integer","direction":"out","target":"total"}],"activities":[],
        "flowchart":{"start":"start","nodes":[{"id":"start","kind":"start"},{"id":"set","kind":"action","activity":{"id":"set-total","type":"assign","target":"total","value":"{{ 7 }}"}},{"id":"end","kind":"end"}],"transitions":[{"from":"start","to":"set"},{"from":"set","to":"end"}]}}
    (tmp_path/"main.json").write_text(json.dumps(main),encoding="utf-8")
    (tmp_path/"child.json").write_text(json.dumps(child),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["total"]==7
    assert result.activities["set-total"]["workflow"]=="Child"
    assert any(event.get("message")=="Result=7" for event in result.logs)

def test_invoke_workflow_rejects_bad_argument_names_directions_and_types(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Typed calls","main":"Main","workflows":{"Main":"main.json","Child":"child.json"}}
    child={"format":"openrpa.workflow","formatVersion":1,"name":"Child","variables":[],"arguments":[
        {"name":"amount","type":"integer","direction":"in"},{"name":"result","type":"string","direction":"out"}],"activities":[]}
    (tmp_path/"child.json").write_text(json.dumps(child),encoding="utf-8")
    for inputs,code in [({"missing":"1"},"argument-unknown"),({"amount":"wrong"},"argument-type"),({"amount":1,"result":"no"},"argument-direction")]:
        main={"format":"openrpa.workflow","formatVersion":1,"name":"Main","variables":[],"arguments":[],"activities":[{"id":"call","type":"invoke_workflow","workflow":"Child","inputs":inputs}]}
        (tmp_path/"main.json").write_text(json.dumps(main),encoding="utf-8")
        result=run_project(project,tmp_path)
        assert result.status=="failed"
        assert result.error["code"]==code

def test_workflow_argument_defaults_must_match_declared_type():
    base={"format":"openrpa.workflow","formatVersion":1,"name":"Defaults","variables":[],"activities":[]}
    valid={**base,"arguments":[{"name":"limit","type":"integer","direction":"in","required":False,"default":12}]}
    assert validate_workflow(valid)==[]
    for arg in [
        {"name":"limit","type":"integer","direction":"in","default":1.5},
        {"name":"enabled","type":"boolean","direction":"in","default":"yes"},
        {"name":"payload","type":"object","direction":"in","default":[]},
    ]:
        assert any(problem["code"]=="argument-default-type" for problem in validate_workflow({**base,"arguments":[arg]}))

def test_invoke_workflow_maps_outputs_per_call_site(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Output mapping","main":"Main","workflows":{"Main":"main.json","Child":"child.json"}}
    main={"format":"openrpa.workflow","formatVersion":1,"name":"Main","variables":[{"name":"callerValue","type":"integer","value":0}],"arguments":[],"activities":[
        {"id":"call","type":"invoke_workflow","workflow":"Child","outputs":{"answer":"callerValue"}}]}
    child={"format":"openrpa.workflow","formatVersion":1,"name":"Child","variables":[],"arguments":[{"name":"answer","type":"integer","direction":"out"}],"activities":[
        {"id":"set-answer","type":"assign","target":"answer","value":42}]}
    (tmp_path/"main.json").write_text(json.dumps(main),encoding="utf-8")
    (tmp_path/"child.json").write_text(json.dumps(child),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["callerValue"]==42

def test_cross_workflow_binding_validation_checks_required_unknown_and_output_names():
    from openrpa.model import validate_workflow_bindings
    workflows={
        "Main":{"name":"Main","activities":[{"id":"call","type":"invoke_workflow","workflow":"Child","inputs":{"extra":"x"},"outputs":{"amount":"other","answer":"bad-name"}}]},
        "Child":{"name":"Child","arguments":[{"name":"amount","type":"integer","direction":"in"},{"name":"answer","type":"integer","direction":"out"}],"activities":[]},
    }
    problems=validate_workflow_bindings(workflows)
    codes={item["code"] for item in problems}
    assert {"argument-unknown","argument-missing","argument-output","argument-output-target"} <= codes
    workflows["Main"]["activities"][0]["inputs"]={"amount":"not-an-integer"}
    assert any(item["code"]=="argument-type" for item in validate_workflow_bindings(workflows))


def test_nested_workflow_failure_reports_origin_and_bounded_call_path(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Error path","main":"Main","workflows":{
        "Main":"main.json","Child":"child.json","Grandchild":"grandchild.json"}}
    main={"format":"openrpa.workflow","formatVersion":1,"name":"Main","variables":[],"arguments":[],"activities":[
        {"id":"call-child","type":"invoke_workflow","workflow":"Child"}]}
    child={"format":"openrpa.workflow","formatVersion":1,"name":"Child","variables":[],"arguments":[],"activities":[
        {"id":"call-grandchild","type":"invoke_workflow","workflow":"Grandchild"}]}
    grandchild={"format":"openrpa.workflow","formatVersion":1,"name":"Grandchild","variables":[],"arguments":[],"activities":[
        {"id":"read-missing","type":"file_read_text","path":"missing.txt","result":"text"}]}
    for filename,value in [("main.json",main),("child.json",child),("grandchild.json",grandchild)]:
        (tmp_path/filename).write_text(json.dumps(value),encoding="utf-8")

    result=run_project(project,tmp_path)

    assert result.status=="failed"
    assert result.error["activityId"]=="read-missing"
    assert result.error["workflow"]=="Grandchild"
    assert result.error["workflowStack"]==[
        {"workflow":"Main","activityId":"call-child","targetWorkflow":"Child"},
        {"workflow":"Child","activityId":"call-grandchild","targetWorkflow":"Grandchild"},
    ]
    assert "Call path: Main/call-child → Child → Child/call-grandchild → Grandchild" in result.error["message"]

def test_invocation_cycle_through_switch_case_is_rejected(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Cycle","main":"A","workflows":{"A":"a.json","B":"b.json"}}
    a={"format":"openrpa.workflow","formatVersion":1,"name":"A","variables":[],"arguments":[],"activities":[{"id":"route","type":"switch","expression":"True","cases":[{"value":True,"activities":[{"id":"to-b","type":"invoke_workflow","workflow":"B"}]}],"default":[]}]}
    b={"format":"openrpa.workflow","formatVersion":1,"name":"B","variables":[],"arguments":[],"activities":[{"id":"to-a","type":"invoke_workflow","workflow":"A"}]}
    (tmp_path/"a.json").write_text(json.dumps(a),encoding="utf-8");(tmp_path/"b.json").write_text(json.dumps(b),encoding="utf-8")
    with pytest.raises(WorkflowError,match="invocation cycle"):
        run_project(project,tmp_path)

def test_file_escape_is_rejected(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[{"id":"read","type":"file_read_text","path":"../outside.txt","result":"x"}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path)
    assert result.status == "failed"
    assert result.error["code"] == "file-permission-denied"
    assert result.error["activityId"] == "read"

def test_http_requires_permission(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[{"id":"http","type":"http_request","url":"https://example.org"}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path)
    assert result.status == "failed"
    assert result.error["code"] == "network-permission-denied"

def test_validation_finds_nested_duplicate_activity_ids():
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"activities":[{"id":"a","type":"if","condition":"True","then":[{"id":"a","type":"log"}]}]}
    problems=validate_workflow(workflow)
    assert any(p["code"]=="activity-id-duplicate" for p in problems)

def test_example_project_validates_and_runs():
    project,root=load_project(EXAMPLE)
    result=run_project(project,root)
    assert result.status=="succeeded"
    assert len(result.variables["rows"])==2
    assert result.variables["rows"][0]["name"] == "Ada Lovelace"
    assert (root/"out"/"cleaned.csv").exists()

def test_cancellation_is_checked_between_activities(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[{"id":"a","type":"log","message":"one"},{"id":"b","type":"log","message":"two"}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path,options=RunOptions(cancel=lambda:True))
    assert result.status=="cancelled"
    assert not result.logs

def test_project_workflow_path_cannot_escape(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"../outside.json"}}
    (tmp_path/"project.json").write_text(json.dumps(project))
    with pytest.raises(WorkflowError, match="safe project-relative path|escapes project root"):
        load_project(tmp_path)

def test_reproducible_package_round_trip(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[{"id":"a","type":"log","message":"hello"}]}
    root=tmp_path/"project";root.mkdir();(root/"project.json").write_text(json.dumps(project));(root/"w.json").write_text(json.dumps(workflow));(root/"data.txt").write_text("data");(root/".env").write_text("SECRET=x")
    first=tmp_path/"one.orpa";second=tmp_path/"two.orpa"
    pack_project(root,first);pack_project(root,second)
    assert hashlib.sha256(first.read_bytes()).digest()==hashlib.sha256(second.read_bytes()).digest()
    unpack_project(first,tmp_path/"restored")
    assert (tmp_path/"restored"/"data.txt").read_text()=="data"
    assert not (tmp_path/"restored"/".env").exists()

def test_package_rejects_path_traversal(tmp_path: Path):
    import zipfile
    archive=tmp_path/"unsafe.orpa"
    manifest={"format":"org.openrpa.package","formatVersion":1,"name":"unsafe","files":[{"path":"../outside","sha256":"0"*64,"size":1}]}
    with zipfile.ZipFile(archive,"w") as z:
        z.writestr("package-manifest.json",json.dumps(manifest));z.writestr("../outside","x")
    with pytest.raises(WorkflowError,match="Unsafe package path|Package file list"):
        unpack_project(archive,tmp_path/"restored")

def test_package_rejects_windows_separator_traversal(tmp_path: Path):
    import zipfile
    archive=tmp_path/"unsafe-backslash.orpa"
    manifest={"format":"org.openrpa.package","formatVersion":1,"name":"unsafe","files":[{"path":"..\\outside","sha256":hashlib.sha256(b"x").hexdigest(),"size":1}]}
    with zipfile.ZipFile(archive,"w") as z:
        z.writestr("package-manifest.json",json.dumps(manifest));z.writestr("..\\outside","x")
    with pytest.raises(WorkflowError,match="Unsafe package path"):
        unpack_project(archive,tmp_path/"restored")

def test_workflow_invocation_cycles_are_rejected(tmp_path: Path):
    (tmp_path/"a.json").write_text(json.dumps({"format":"openrpa.workflow","formatVersion":1,"name":"A","variables":[],"arguments":[],"activities":[{"id":"to-b","type":"invoke_workflow","workflow":"B"}]}))
    (tmp_path/"b.json").write_text(json.dumps({"format":"openrpa.workflow","formatVersion":1,"name":"B","variables":[],"arguments":[],"activities":[{"id":"to-a","type":"invoke_workflow","workflow":"A"}]}))
    project={"format":"openrpa.project","formatVersion":1,"name":"cycle","main":"A","workflows":{"A":"a.json","B":"b.json"}}
    with pytest.raises(WorkflowError,match="invocation cycle"):
        run_project(project,tmp_path)

def test_argument_type_is_enforced(tmp_path: Path):
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"typed","variables":[],"arguments":[{"name":"count","type":"integer","direction":"in"}],"activities":[]}
    project={"format":"openrpa.project","formatVersion":1,"name":"typed","main":"M","workflows":{"M":"w.json"}}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path,inputs={"count":"not-an-integer"})
    assert result.status=="failed"
    assert result.error["code"]=="argument-type"

def test_expression_interpreter_rejects_arbitrary_python():
    with pytest.raises(ExpressionError,match="not allowed"):
        evaluate("__import__('os').system('whoami')",{})

def test_run_history_survives_restart_and_marks_interrupted_without_replay(tmp_path: Path):
    from openrpa.history import RunHistory
    store = RunHistory(tmp_path / "history.sqlite3")
    run_id = store.start(str(tmp_path), "test", "Main", "abc123")
    store.event(run_id, {"event":"activity-started", "activityId":"send", "type":"http_request"})
    # A second store models a new application process opening the same ledger.
    reopened = RunHistory(tmp_path / "history.sqlite3")
    record = reopened.get_run(run_id)
    assert record["status"] == "interrupted"
    assert "unknown" in record["recoveryNote"]
    assert record["events"][0]["activityId"] == "send"
    assert "recordedAt" in record["events"][0] and "T" in record["events"][0]["recordedAt"]
    assert reopened.list_runs()[0]["runId"] == run_id

def test_run_history_database_scopes_close_connections(tmp_path: Path):
    import sqlite3
    from openrpa.history import RunHistory

    store = RunHistory(tmp_path / "history.sqlite3")
    with store._db() as db:
        db.execute("INSERT INTO audit_events(created_at,action,details_json) VALUES('now','test','{}')")
    with pytest.raises(sqlite3.ProgrammingError, match="closed database"):
        db.execute("SELECT 1")
    assert store.list_audit(1)[0]["action"] == "test"

def test_history_persists_completed_failed_result(tmp_path: Path):
    from openrpa.history import RunHistory
    store = RunHistory(tmp_path / "history.sqlite3")
    run_id = store.start(str(tmp_path), "test", "Main", "digest")
    result = {"status":"failed", "activities":{"write":{"status":"failed"}}, "logs":[]}
    store.finish(run_id, "failed", result, {"code":"file-permission-denied", "message":"blocked"})
    reopened = RunHistory(tmp_path / "history.sqlite3")
    assert reopened.get_run(run_id)["result"] == result
    assert reopened.get_run(run_id)["error"]["code"] == "file-permission-denied"

def test_history_schema_upgrade_backs_up_before_migration(tmp_path: Path):
    import sqlite3
    from openrpa.history import RunHistory
    path=tmp_path/"history.sqlite3"
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE schema_version(version INTEGER NOT NULL)")
        db.execute("INSERT INTO schema_version VALUES(0)")
    RunHistory(path)
    assert (tmp_path/"history.sqlite3.v0.bak").is_file()
    with sqlite3.connect(path) as db:
        assert db.execute("SELECT version FROM schema_version").fetchone()[0]==1

def test_http_egress_blocks_local_private_redirect_and_nonstandard_port(monkeypatch):
    from openrpa.network import EgressDenied, _target, _resolve_public
    import openrpa.network as network
    _target("http://127.0.0.1/admin")  # URL parsing alone is not the network boundary.
    with pytest.raises(EgressDenied, match="non-public"):
        _resolve_public("127.0.0.1", 80)
    with pytest.raises(EgressDenied, match="standard"):
        _target("https://example.org:8443/")
    with pytest.raises(EgressDenied, match="Credentials"):
        _target("https://user:pass@example.org/")
    monkeypatch.setattr(network.socket, "getaddrinfo", lambda *args, **kwargs: [(None,None,None,None,("10.0.0.7",443))])
    with pytest.raises(EgressDenied, match="non-public"):
        _resolve_public("internal.example", 443)

def test_http_request_never_connects_to_private_dns_answer(monkeypatch):
    import openrpa.network as network
    from openrpa.network import EgressDenied
    monkeypatch.setattr(network.socket, "getaddrinfo", lambda *args, **kwargs: [(None,None,None,None,("192.168.1.5",443))])
    monkeypatch.setattr(network.socket, "create_connection", lambda *args, **kwargs: pytest.fail("must not connect"))
    with pytest.raises(EgressDenied, match="non-public"):
        network.request("https://buyer.example/", "GET", {}, None, 3, 1024)

def test_http_client_returns_redirect_without_following_it(monkeypatch):
    import openrpa.network as network
    monkeypatch.setattr(network,"_resolve_public",lambda host,port:"93.184.216.34")
    calls=[]
    class Response:
        status=302
        def getheaders(self): return [("Location","http://127.0.0.1/admin")]
        def read(self,size): return b"redirect"
    class Connection:
        def __init__(self,host,port,address,timeout): calls.append((host,port,address))
        def request(self,*args,**kwargs): calls.append("request")
        def getresponse(self): return Response()
        def close(self): pass
    monkeypatch.setattr(network,"_PinnedHTTPS",Connection)
    response=network.request("https://safe.example/","GET",{},None,5,1024)
    assert response["status"]==302 and response["redirectBlocked"] is True
    assert calls==[("safe.example",443,"93.184.216.34"),"request"]

def test_http_activity_with_network_permission_still_blocks_loopback(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[{"id":"http","type":"http_request","url":"http://127.0.0.1/secret"}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path,options=RunOptions(allow_network=True))
    assert result.status=="failed"
    assert result.error["code"]=="http-egress-denied"

def test_engine_log_events_are_bounded(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    activities=[{"id":f"a{i}","type":"log","message":"x"*1000} for i in range(20)]
    (tmp_path/"w.json").write_text(json.dumps({"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":activities}))
    events=[]
    result=run_project(project,tmp_path,options=RunOptions(max_log_events=3,max_log_event_chars=256,event=events.append))
    assert result.status=="succeeded"
    assert len(result.logs)==3
    assert len(events)==3

def test_nested_loop_activity_quota_stops_excessive_execution(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[
        {"id":"loop","type":"for_each","items":"[1,2,3]","itemName":"item","body":[{"id":"log","type":"log","message":"x"}]}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path,options=RunOptions(max_activity_executions=2))
    assert result.status=="failed" and result.error["code"]=="activity-quota"

def test_regex_backtracking_has_a_time_quota(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"x","main":"M","workflows":{"M":"w.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"x","variables":[],"arguments":[],"activities":[
        {"id":"regex","type":"regex_match","text":"a"*10_000+"!","pattern":"(a+)+$","result":"match"}]}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path)
    assert result.status=="failed" and result.error["code"]=="regex-timeout"

def test_keyring_secret_never_returns_in_logs_variables_or_activity_states(tmp_path: Path, monkeypatch):
    import sys, types
    from openrpa.engine import run_project
    class WinVaultKeyring: pass
    WinVaultKeyring.__module__="keyring.backends.Windows"
    fake=types.ModuleType("keyring")
    fake.get_keyring=lambda:WinVaultKeyring()
    fake.get_password=lambda service,username:"very-secret-value"
    monkeypatch.setitem(sys.modules,"keyring",fake)
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"secrets","variables":[],"arguments":[],"activities":[
        {"id":"get","type":"credential_get","service":"test","username":"alice","result":"token"},
        {"id":"log","type":"log","message":"token={{token}}"}]}
    project={"format":"openrpa.project","formatVersion":1,"name":"secret","main":"M","workflows":{"M":"w.json"}}
    (tmp_path/"w.json").write_text(json.dumps(workflow))
    result=run_project(project,tmp_path)
    rendered=json.dumps(result.__dict__)
    assert result.status=="succeeded"
    assert "very-secret-value" not in rendered
    assert "[REDACTED]" in rendered

def test_plaintext_keyring_backend_is_rejected(monkeypatch):
    import sys,types
    from openrpa.credentials import get_safe_keyring
    class PlainKeyring: pass
    PlainKeyring.__module__="keyrings.alt.file"
    fake=types.ModuleType("keyring");fake.get_keyring=lambda:PlainKeyring()
    monkeypatch.setitem(sys.modules,"keyring",fake)
    with pytest.raises(RuntimeError,match="non-approved keyring backend"):
        get_safe_keyring()

def test_state_machine_runs_guards_then_unconditional_fallback(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"State machine","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Route","executionModel":"state_machine","variables":[{"name":"approved","type":"boolean","value":False},{"name":"route","type":"string","value":""}],"arguments":[],"activities":[],"stateMachine":{"initial":"review","states":[
        {"id":"review","label":"Review request","activities":[{"id":"review-log","type":"log","message":"reviewing"}]},
        {"id":"accept","label":"Accepted","final":True,"activities":[{"id":"accept-route","type":"assign","target":"route","value":"accepted"}]},
        {"id":"reject","label":"Rejected","final":True,"activities":[{"id":"reject-route","type":"assign","target":"route","value":"rejected"}]}
    ],"transitions":[{"from":"review","to":"accept","condition":"approved"},{"from":"review","to":"reject"}]}}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["route"]=="rejected"
    assert result.activities["review"]["status"]=="succeeded"
    assert result.activities["reject-route"]["status"]=="succeeded"

def test_state_machine_rejects_unreachable_state_and_bad_fallback_order():
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Bad states","executionModel":"state_machine","variables":[],"arguments":[],"activities":[],"stateMachine":{"initial":"a","states":[{"id":"a"},{"id":"b","final":True},{"id":"unused","final":True}],"transitions":[{"from":"a","to":"b"},{"from":"a","to":"unused","condition":"True"}]}}
    codes={problem["code"] for problem in validate_workflow(workflow)}
    assert "state-machine-fallback-order" in codes
    assert "state-machine-unreachable" not in codes  # every edge is reachable, regardless of guards

def test_state_machine_reports_non_boolean_guard_as_failed_state(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Bad guard","main":"Main","workflows":{"Main":"main.json"}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Bad guard","executionModel":"state_machine","variables":[],"arguments":[],"activities":[],"stateMachine":{"initial":"a","states":[{"id":"a"},{"id":"b","final":True}],"transitions":[{"from":"a","to":"b","condition":"1"}]}}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed"
    assert result.error["code"]=="state-machine-condition-type"
    assert result.activities["a"]["status"]=="failed"

def test_state_machine_cycles_obey_activity_visit_quota(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Cycle","main":"Main","workflows":{"Main":"main.json"},"limits":{"maxActivityExecutions":3}}
    workflow={"format":"openrpa.workflow","formatVersion":1,"name":"Cycle","executionModel":"state_machine","variables":[],"arguments":[],"activities":[],"stateMachine":{"initial":"again","states":[{"id":"again","activities":[]},{"id":"done","final":True,"activities":[]}],"transitions":[{"from":"again","to":"again","condition":"True"},{"from":"again","to":"done"}]}}
    (tmp_path/"main.json").write_text(json.dumps(workflow),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="failed"
    assert result.error["code"]=="activity-quota"

def test_invoke_workflow_runs_state_machine_and_maps_outputs(tmp_path: Path):
    project={"format":"openrpa.project","formatVersion":1,"name":"Invoke state","main":"Main","workflows":{"Main":"main.json","Child":"child.json"}}
    parent={"format":"openrpa.workflow","formatVersion":1,"name":"Main","variables":[],"arguments":[],"activities":[{"id":"invoke","type":"invoke_workflow","workflow":"Child","arguments":{},"resultTarget":"childResults"}]}
    child={"format":"openrpa.workflow","formatVersion":1,"name":"Child","executionModel":"state_machine","variables":[{"name":"answer","type":"string","value":"ready"}],"arguments":[],"activities":[],"stateMachine":{"initial":"complete","states":[{"id":"complete","final":True,"activities":[]}],"transitions":[]}}
    (tmp_path/"main.json").write_text(json.dumps(parent),encoding="utf-8")
    (tmp_path/"child.json").write_text(json.dumps(child),encoding="utf-8")
    result=run_project(project,tmp_path)
    assert result.status=="succeeded"
    assert result.variables["childResults"]["answer"]=="ready"
